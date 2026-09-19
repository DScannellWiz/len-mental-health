import argparse
import hashlib
import json
import math
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from types import MappingProxyType
from xml.sax.saxutils import escape
from tkinter import (
    BOTH,
    BOTTOM,
    END,
    LEFT,
    RIGHT,
    TOP,
    Button,
    Canvas,
    Checkbutton,
    Entry,
    Frame,
    IntVar,
    Label,
    LabelFrame,
    Menu,
    StringVar,
    Text,
    Tk,
    Toplevel,
    filedialog,
    messagebox,
    ttk,
)

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except Exception:
    load_workbook = None
    Alignment = None
    Font = None
    PatternFill = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
except Exception:
    colors = None
    inch = 72
    Image = None
    KeepTogether = None
    PageBreak = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None
    Table = None
    TableStyle = None

try:
    from PIL import Image as PILImage
    from PIL import ImageDraw, ImageFont
except Exception:
    PILImage = None


APP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIR.parents[1] if APP_DIR.parent.name == "src" else APP_DIR.parent
DATA_DIR = PROJECT_ROOT / "data"
REPORTS_DIR = PROJECT_ROOT / "reports"
ICON_PATH = PROJECT_ROOT / "packaging" / "assets" / "PHQ9_Tracker.ico"
APPLICATION_NAME = "Len"


def default_db_path() -> Path:
    explicit = os.environ.get("PHQ9_TRACKER_DB_PATH")
    if explicit:
        return Path(explicit)
    if getattr(sys, "frozen", False):
        if os.environ.get("PHQ9_TRACKER_PORTABLE") == "1":
            return Path(sys.executable).resolve().parent / "phq9_tracker.sqlite"
        local_app_data = os.environ.get("LOCALAPPDATA")
        if local_app_data:
            return Path(local_app_data) / "PHQ9Tracker" / "phq9_tracker.sqlite"
    return DATA_DIR / "phq9_tracker.sqlite"


def generated_output_dir(create: bool = False) -> Path:
    """Return the one local folder used for both clinician outputs."""
    if getattr(sys, "frozen", False):
        if os.environ.get("PHQ9_TRACKER_PORTABLE") == "1":
            base_dir = Path(sys.executable).resolve().parent
        else:
            local_app_data = os.environ.get("LOCALAPPDATA")
            base_dir = Path(local_app_data) / "PHQ9Tracker" if local_app_data else Path(sys.executable).resolve().parent
        output_dir = base_dir / "reports"
    else:
        output_dir = REPORTS_DIR
    if create:
        output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir


def next_available_output_path(filename: str) -> Path:
    """Choose a non-destructive path in the common output folder."""
    output_dir = generated_output_dir(create=True)
    candidate = output_dir / filename
    suffix_number = 2
    while candidate.exists():
        candidate = output_dir / f"{Path(filename).stem}_{suffix_number}{Path(filename).suffix}"
        suffix_number += 1
    return candidate


def open_in_default_application(path: Path) -> None:
    """Open a local file or folder with its Windows default application."""
    startfile = getattr(os, "startfile", None)
    if startfile is None:
        raise OSError("Windows could not open this location automatically.")
    startfile(str(path))


def offer_to_open_generated_file(path: Path, output_name: str) -> None:
    """Offer to open a successfully saved output without changing save success."""
    if not messagebox.askyesno(
        f"{output_name} saved",
        f"Saved successfully to:\n\n{path}\n\nOpen it now?",
    ):
        return
    try:
        open_in_default_application(path)
    except OSError as exc:
        messagebox.showwarning(
            f"{output_name} saved",
            f"The file was saved successfully, but Windows could not open it automatically.\n\nSaved path:\n{path}\n\n{exc}",
        )


DB_PATH = default_db_path()
BUNDLED_PYTHON = Path(os.environ["PHQ9_TRACKER_BUNDLED_PYTHON"]) if os.environ.get("PHQ9_TRACKER_BUNDLED_PYTHON") else None
DISCLAIMER = "This report is for discussion with a licensed clinician and is not a diagnosis."
NON_DIAGNOSTIC_OUTPUT_NOTICE = (
    "Questionnaire outputs describe recorded responses and deterministic calculations only. "
    "They do not provide a diagnosis, treatment recommendation, or emergency monitoring."
)
UNIVERSAL_SAFETY_MESSAGE = (
    "If you feel unsafe or may act on thoughts of self-harm, contact local emergency services "
    "or a crisis service now. Len does not monitor responses or provide emergency help."
)
DAILY_SCORE_LABEL = "Daily Severity Score"
FREQUENCY_SCORE_LABEL = "14-Day Symptom Frequency Score"
ANALYSIS_WORKBOOK_SCHEMA_VERSION = "1.3"
COMPACT_SPINBOX_PADDING = (2, 0)
UNANSWERED_RESPONSE = "Select a response"
SCORING_EXPLANATION = f"""{APPLICATION_NAME} calculates two related but different measurements.

Daily Severity Score
Each PHQ-9 or GAD-7 item is rated from 0 (symptom not present) to 3 (high symptom severity). The item responses are summed for that date. This answers: How severe were the reported symptoms on this particular day?

14-Day Symptom Frequency Score
For the 14 calendar days ending on the selected date, each response greater than 0 counts as one symptom-present day. For each item, 0 days = 0 points, 1-6 days = 1 point, 7-11 days = 2 points, and 12-14 days = 3 points. The converted item scores are summed. This answers: How consistently were these symptoms present during the last 14 days?

Important distinction
A daily response of 1 and a daily response of 3 each count as one symptom-present day in the 14-day calculation, although they contribute differently to the Daily Severity Score. For example, an item present on 6 days receives 1 frequency point; an item present on 12 days receives 3 frequency points.

Data coverage
The application displays how many daily entries were available. Calendar days without an entry are treated as days with no recorded symptom-present response; that is not proof that the symptom was absent.

Mindful check-ins
Daily check-ins should encourage mindful reflection rather than rapid completion. The application intentionally requires users to consider each symptom individually to reduce habitual responses and improve the quality of the recorded data."""

TREATMENT_EVENT_TYPES = [
    "Therapy",
    "Ketamine infusion",
    "Psychiatry appointment",
    "Medication change",
    "Primary care appointment",
    "Exercise",
    "Support group",
]

PHQ9_ITEM_LABELS = [
    "Little interest or pleasure in doing things",
    "Feeling down, depressed, or hopeless",
    "Sleep changes",
    "Low energy",
    "Appetite changes",
    "Feeling bad about yourself",
    "Trouble concentrating",
    "Moving/speaking slowly or being restless",
    "Thoughts of self-harm",
]

GAD7_ITEM_LABELS = [
    "Feeling nervous, anxious, or on edge",
    "Not being able to stop or control worrying",
    "Worrying too much about different things",
    "Trouble relaxing",
    "Being so restless that it is hard to sit still",
    "Becoming easily annoyed or irritable",
    "Feeling afraid as if something awful might happen",
]

ITEM_LABELS = PHQ9_ITEM_LABELS


DEFINITION_ID_PATTERN = re.compile(r"^[a-z][a-z0-9]*(?:[._-][a-z0-9]+)*$")
INTERPRETATION_POLICIES = frozenset({"validated_builtin", "descriptive_only", "raw_only"})
DEFINITION_ORIGINS = frozenset({"builtin", "custom"})
QUESTIONNAIRE_BEHAVIORS = frozenset()
QUESTION_BEHAVIORS = frozenset({"phq9_item9_context"})


@dataclass(frozen=True)
class ResponseOption:
    value: str | int
    label: str
    score: int | float | None = None


@dataclass(frozen=True)
class QuestionDefinition:
    question_id: str
    prompt: str
    options: tuple[ResponseOption, ...]
    response_type: str = "single_choice"
    required: bool = True
    behavior_ids: tuple[str, ...] = ()
    report_label: str | None = None


@dataclass(frozen=True)
class ScoringStrategyDescriptor:
    """Application-owned scoring capability referenced by a stable definition ID."""

    strategy_id: str
    implementation_id: str | None
    deterministic: bool


@dataclass(frozen=True)
class ProfileStrategyDescriptor:
    """Application-owned profile capability referenced by a stable definition ID."""

    strategy_id: str
    implementation_id: str | None


SCORING_STRATEGY_DESCRIPTORS = MappingProxyType(
    {
        "sum": ScoringStrategyDescriptor("sum", "sum_option_scores_v1", True),
        "subscale_sum": ScoringStrategyDescriptor("subscale_sum", None, True),
        "none": ScoringStrategyDescriptor("none", "no_total", True),
    }
)
PROFILE_STRATEGY_DESCRIPTORS = MappingProxyType(
    {
        "symptom_presence_14d": ProfileStrategyDescriptor(
            "symptom_presence_14d", "symptom_presence_thresholds_14d_v1"
        ),
        "item_history": ProfileStrategyDescriptor("item_history", None),
        "daily_total": ProfileStrategyDescriptor("daily_total", None),
        "subscale_trend": ProfileStrategyDescriptor("subscale_trend", None),
        "none": ProfileStrategyDescriptor("none", "no_profile"),
    }
)
SCORING_RULES = frozenset(SCORING_STRATEGY_DESCRIPTORS)
PROFILE_RULES = frozenset(PROFILE_STRATEGY_DESCRIPTORS)


@dataclass(frozen=True)
class QuestionnaireDefinition:
    questionnaire_id: str
    definition_version: int
    display_name: str
    short_name: str
    description: str
    timeframe_text: str
    source_citation: str
    source_url: str
    rights_status: str
    rights_source_url: str
    required_notice: str
    origin: str
    items: tuple[QuestionDefinition, ...]
    scoring_rule: str
    score_min: int | float | None
    score_max: int | float | None
    profile_rule: str
    interpretation_policy: str
    behavior_ids: tuple[str, ...] = ()
    active: bool = True

    @property
    def assessment_id(self) -> str:
        """Compatibility name retained for existing storage and API callers."""
        return self.questionnaire_id

    @property
    def item_count(self) -> int:
        return len(self.items)

    @property
    def item_labels(self) -> list[str]:
        """Return a copy so legacy callers cannot mutate the frozen definition."""
        return [item.report_label or item.prompt for item in self.items]

    @property
    def max_score(self) -> int | float | None:
        return self.score_max

    @property
    def source(self) -> str:
        return self.source_citation

    @property
    def redistribution_status(self) -> str:
        return self.rights_status

    @property
    def has_item9_safety_context(self) -> bool:
        return any("phq9_item9_context" in item.behavior_ids for item in self.items)


# Assessment terminology remains an import-level compatibility alias. New code
# should use the questionnaire vocabulary above.
AssessmentDefinition = QuestionnaireDefinition


def _validate_identifier(value: str, field_name: str) -> None:
    if not isinstance(value, str) or not DEFINITION_ID_PATTERN.fullmatch(value):
        raise ValueError(f"{field_name} must be a stable lowercase identifier.")


def validate_questionnaire_definition(definition: QuestionnaireDefinition) -> QuestionnaireDefinition:
    _validate_identifier(definition.questionnaire_id, "questionnaire_id")
    if isinstance(definition.definition_version, bool) or not isinstance(definition.definition_version, int) or definition.definition_version < 1:
        raise ValueError("definition_version must be a positive integer.")
    for field_name in ("display_name", "short_name", "description", "timeframe_text", "source_citation"):
        if not isinstance(getattr(definition, field_name), str) or not getattr(definition, field_name).strip():
            raise ValueError(f"{field_name} is required.")
    if definition.origin not in DEFINITION_ORIGINS:
        raise ValueError(f"Unknown definition origin: {definition.origin}")
    if definition.scoring_rule not in SCORING_RULES:
        raise ValueError(f"Unknown scoring rule: {definition.scoring_rule}")
    if definition.profile_rule not in PROFILE_RULES:
        raise ValueError(f"Unknown profile rule: {definition.profile_rule}")
    if definition.interpretation_policy not in INTERPRETATION_POLICIES:
        raise ValueError(f"Unknown interpretation policy: {definition.interpretation_policy}")
    if definition.origin == "custom" and definition.interpretation_policy == "validated_builtin":
        raise ValueError("Custom questionnaires cannot use validated_builtin interpretation.")
    if definition.origin == "builtin":
        for field_name in ("source_url", "rights_status", "rights_source_url"):
            if not isinstance(getattr(definition, field_name), str) or not getattr(definition, field_name).strip():
                raise ValueError(f"Built-in questionnaires require {field_name}.")
    if not isinstance(definition.required_notice, str):
        raise ValueError("required_notice must be text.")
    if not isinstance(definition.active, bool):
        raise ValueError("active must be true or false.")
    if not isinstance(definition.behavior_ids, tuple) or not all(isinstance(value, str) for value in definition.behavior_ids):
        raise ValueError("Questionnaire behavior IDs must be an immutable tuple of strings.")
    unknown_behaviors = set(definition.behavior_ids) - QUESTIONNAIRE_BEHAVIORS
    if unknown_behaviors:
        raise ValueError(f"Unknown questionnaire behavior: {sorted(unknown_behaviors)[0]}")
    if not definition.items:
        raise ValueError("Questionnaire definitions require at least one question.")

    question_ids: set[str] = set()
    calculated_min = 0.0
    calculated_max = 0.0
    for item in definition.items:
        _validate_identifier(item.question_id, "question_id")
        if item.question_id in question_ids:
            raise ValueError(f"Duplicate question_id: {item.question_id}")
        question_ids.add(item.question_id)
        if not isinstance(item.prompt, str) or not item.prompt.strip():
            raise ValueError(f"Question {item.question_id} requires a prompt.")
        if item.response_type != "single_choice":
            raise ValueError(f"Unsupported response type: {item.response_type}")
        if not isinstance(item.required, bool):
            raise ValueError(f"Question {item.question_id} required must be true or false.")
        if item.report_label is not None and (not isinstance(item.report_label, str) or not item.report_label.strip()):
            raise ValueError(f"Question {item.question_id} has an invalid report label.")
        if not isinstance(item.behavior_ids, tuple) or not all(isinstance(value, str) for value in item.behavior_ids):
            raise ValueError(f"Question {item.question_id} behavior IDs must be an immutable tuple of strings.")
        unknown_item_behaviors = set(item.behavior_ids) - QUESTION_BEHAVIORS
        if unknown_item_behaviors:
            raise ValueError(f"Unknown question behavior: {sorted(unknown_item_behaviors)[0]}")
        if "phq9_item9_context" in item.behavior_ids and not (
            definition.questionnaire_id == "phq9" and item.question_id == "phq9.item9"
        ):
            raise ValueError("phq9_item9_context is restricted to phq9.item9.")
        if not item.options:
            raise ValueError(f"Question {item.question_id} requires response options.")

        option_values: set[str | int] = set()
        option_scores: list[float] = []
        for option in item.options:
            if isinstance(option.value, bool) or not isinstance(option.value, (str, int)):
                raise ValueError(f"Question {item.question_id} has an invalid option value.")
            if option.value in option_values:
                raise ValueError(f"Question {item.question_id} has duplicate option values.")
            option_values.add(option.value)
            if not isinstance(option.label, str) or not option.label.strip():
                raise ValueError(f"Question {item.question_id} has an option without a label.")
            if option.score is not None:
                if isinstance(option.score, bool) or not isinstance(option.score, (int, float)) or not math.isfinite(option.score):
                    raise ValueError(f"Question {item.question_id} has an invalid option score.")
                option_scores.append(float(option.score))

        if definition.scoring_rule == "sum":
            if len(option_scores) != len(item.options):
                raise ValueError("Sum-scored questionnaires require a numeric score for every option.")
            calculated_min += min(option_scores)
            calculated_max += max(option_scores)

    if definition.scoring_rule == "sum":
        if (
            definition.score_min is None
            or definition.score_max is None
            or isinstance(definition.score_min, bool)
            or isinstance(definition.score_max, bool)
            or not isinstance(definition.score_min, (int, float))
            or not isinstance(definition.score_max, (int, float))
            or not math.isfinite(definition.score_min)
            or not math.isfinite(definition.score_max)
        ):
            raise ValueError("Sum-scored questionnaires require score bounds.")
        if not math.isclose(float(definition.score_min), calculated_min) or not math.isclose(float(definition.score_max), calculated_max):
            raise ValueError("Score bounds do not match the declared response options.")
    elif definition.score_min is not None or definition.score_max is not None:
        raise ValueError("Nonscored questionnaires cannot declare score bounds.")
    return definition


def questionnaire_definition_data(definition: QuestionnaireDefinition) -> dict[str, object]:
    """Return the stable JSON-compatible representation used by future snapshots."""
    validate_questionnaire_definition(definition)
    return {
        "questionnaire_id": definition.questionnaire_id,
        "definition_version": definition.definition_version,
        "display_name": definition.display_name,
        "short_name": definition.short_name,
        "description": definition.description,
        "timeframe_text": definition.timeframe_text,
        "source_citation": definition.source_citation,
        "source_url": definition.source_url,
        "rights_status": definition.rights_status,
        "rights_source_url": definition.rights_source_url,
        "required_notice": definition.required_notice,
        "origin": definition.origin,
        "items": [
            {
                "question_id": item.question_id,
                "prompt": item.prompt,
                "response_type": item.response_type,
                "options": [
                    {"value": option.value, "label": option.label, "score": option.score}
                    for option in item.options
                ],
                "required": item.required,
                "behavior_ids": list(item.behavior_ids),
                "report_label": item.report_label,
            }
            for item in definition.items
        ],
        "scoring_rule": definition.scoring_rule,
        "score_min": definition.score_min,
        "score_max": definition.score_max,
        "profile_rule": definition.profile_rule,
        "interpretation_policy": definition.interpretation_policy,
        "behavior_ids": list(definition.behavior_ids),
        "active": definition.active,
    }


def serialize_questionnaire_definition(definition: QuestionnaireDefinition) -> str:
    return json.dumps(questionnaire_definition_data(definition), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def deserialize_questionnaire_definition(payload: str) -> QuestionnaireDefinition:
    """Parse definition-only JSON; validation rejects unsafe or unsupported capabilities."""
    try:
        data = json.loads(payload)
        if not isinstance(data, dict):
            raise TypeError("Definition JSON must contain an object.")
        items = tuple(
            QuestionDefinition(
                question_id=item["question_id"],
                prompt=item["prompt"],
                response_type=item["response_type"],
                options=tuple(
                    ResponseOption(value=option["value"], label=option["label"], score=option["score"])
                    for option in item["options"]
                ),
                required=item["required"],
                behavior_ids=tuple(item["behavior_ids"]),
                report_label=item["report_label"],
            )
            for item in data["items"]
        )
        definition = QuestionnaireDefinition(
            questionnaire_id=data["questionnaire_id"],
            definition_version=data["definition_version"],
            display_name=data["display_name"],
            short_name=data["short_name"],
            description=data["description"],
            timeframe_text=data["timeframe_text"],
            source_citation=data["source_citation"],
            source_url=data["source_url"],
            rights_status=data["rights_status"],
            rights_source_url=data["rights_source_url"],
            required_notice=data["required_notice"],
            origin=data["origin"],
            items=items,
            scoring_rule=data["scoring_rule"],
            score_min=data["score_min"],
            score_max=data["score_max"],
            profile_rule=data["profile_rule"],
            interpretation_policy=data["interpretation_policy"],
            behavior_ids=tuple(data["behavior_ids"]),
            active=data["active"],
        )
    except (KeyError, TypeError, json.JSONDecodeError) as exc:
        raise ValueError("Invalid questionnaire definition JSON.") from exc
    return validate_questionnaire_definition(definition)


LEGACY_FREQUENCY_RESPONSE_OPTIONS = tuple(
    ResponseOption(value=value, label=label, score=value)
    for value, label in enumerate(("Not at all", "Several days", "More than half the days", "Nearly every day"))
)
DAILY_SEVERITY_RESPONSE_OPTIONS = tuple(
    ResponseOption(value=value, label=label, score=value)
    for value, label in enumerate(("Not present", "Mild", "Moderate", "High"))
)
DAILY_SEVERITY_INSTRUCTION = "For today, rate how severe each symptom was from 0 (not present) to 3 (high)."


def _builtin_questions(
    questionnaire_id: str,
    labels: list[str],
    response_options: tuple[ResponseOption, ...],
) -> tuple[QuestionDefinition, ...]:
    return tuple(
        QuestionDefinition(
            question_id=f"{questionnaire_id}.item{index}",
            prompt=label,
            options=response_options,
            behavior_ids=("phq9_item9_context",) if questionnaire_id == "phq9" and index == 9 else (),
            report_label=label,
        )
        for index, label in enumerate(labels, start=1)
    )


PHQ_SCREENER_SOURCE_URL = "https://www.phqscreeners.com/select-screener"


def _builtin_definition(
    questionnaire_id: str,
    definition_version: int,
    display_name: str,
    description: str,
    item_labels: list[str],
    score_max: int,
    response_options: tuple[ResponseOption, ...],
    timeframe_text: str,
) -> QuestionnaireDefinition:
    return QuestionnaireDefinition(
        questionnaire_id=questionnaire_id,
        definition_version=definition_version,
        display_name=display_name,
        short_name=display_name,
        description=description,
        timeframe_text=timeframe_text,
        source_citation=description,
        source_url=PHQ_SCREENER_SOURCE_URL,
        rights_status="Redistributable built-in",
        rights_source_url=PHQ_SCREENER_SOURCE_URL,
        required_notice="",
        origin="builtin",
        items=_builtin_questions(questionnaire_id, item_labels, response_options),
        scoring_rule="sum",
        score_min=0,
        score_max=score_max,
        profile_rule="symptom_presence_14d",
        interpretation_policy="validated_builtin",
    )


LEGACY_BUILTIN_DEFINITIONS = {
    "phq9": _builtin_definition(
        questionnaire_id="phq9",
        definition_version=1,
        display_name="PHQ-9",
        description="Patient Health Questionnaire-9",
        item_labels=PHQ9_ITEM_LABELS,
        score_max=27,
        response_options=LEGACY_FREQUENCY_RESPONSE_OPTIONS,
        timeframe_text="Record responses for the selected check-in date.",
    ),
    "gad7": _builtin_definition(
        questionnaire_id="gad7",
        definition_version=1,
        display_name="GAD-7",
        description="Generalized Anxiety Disorder-7",
        item_labels=GAD7_ITEM_LABELS,
        score_max=21,
        response_options=LEGACY_FREQUENCY_RESPONSE_OPTIONS,
        timeframe_text="Record responses for the selected check-in date.",
    ),
}
ASSESSMENTS = {
    questionnaire_id: _builtin_definition(
        questionnaire_id=questionnaire_id,
        definition_version=2,
        display_name=legacy.display_name,
        description=legacy.description,
        item_labels=legacy.item_labels,
        score_max=int(legacy.score_max),
        response_options=DAILY_SEVERITY_RESPONSE_OPTIONS,
        timeframe_text=DAILY_SEVERITY_INSTRUCTION,
    )
    for questionnaire_id, legacy in LEGACY_BUILTIN_DEFINITIONS.items()
}
BUILTIN_DEFINITION_SNAPSHOTS = (
    *LEGACY_BUILTIN_DEFINITIONS.values(),
    *ASSESSMENTS.values(),
)
for _definition in BUILTIN_DEFINITION_SNAPSHOTS:
    validate_questionnaire_definition(_definition)

ASSESSMENT_ORDER = ["phq9", "gad7"]
QUESTIONNAIRES = ASSESSMENTS
QUESTIONNAIRE_ORDER = ASSESSMENT_ORDER


@dataclass(frozen=True)
class SchemaMigration:
    """One application-owned, checksummed SQLite schema migration."""

    migration_id: str
    statements: tuple[str, ...]

    @property
    def checksum(self) -> str:
        payload = f"{self.migration_id}\n" + "\n".join(statement.strip() for statement in self.statements)
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()


SCHEMA_MIGRATION_LEDGER_STATEMENTS = (
    """
    CREATE TABLE IF NOT EXISTS schema_migrations (
        migration_id TEXT PRIMARY KEY,
        application_id TEXT NOT NULL CHECK(application_id = 'len'),
        checksum_sha256 TEXT NOT NULL CHECK(length(checksum_sha256) = 64),
        applied_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    )
    """,
    """
    CREATE TRIGGER IF NOT EXISTS schema_migrations_no_update
    BEFORE UPDATE ON schema_migrations
    BEGIN
        SELECT RAISE(ABORT, 'schema migration records are immutable');
    END
    """,
    """
    CREATE TRIGGER IF NOT EXISTS schema_migrations_no_delete
    BEFORE DELETE ON schema_migrations
    BEGIN
        SELECT RAISE(ABORT, 'schema migration records are immutable');
    END
    """,
)

QUESTIONNAIRE_SNAPSHOT_MIGRATION = SchemaMigration(
    migration_id="len.014.001.questionnaire_definition_snapshots",
    statements=(
        """
        CREATE TABLE questionnaire_definition_snapshots (
            questionnaire_id TEXT NOT NULL,
            definition_version INTEGER NOT NULL CHECK(definition_version > 0),
            definition_json TEXT NOT NULL CHECK(length(definition_json) > 0),
            definition_sha256 TEXT NOT NULL CHECK(length(definition_sha256) = 64),
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (questionnaire_id, definition_version)
        )
        """,
        """
        CREATE TRIGGER questionnaire_definition_snapshots_no_update
        BEFORE UPDATE ON questionnaire_definition_snapshots
        BEGIN
            SELECT RAISE(ABORT, 'questionnaire definition snapshots are immutable');
        END
        """,
        """
        CREATE TRIGGER questionnaire_definition_snapshots_no_delete
        BEFORE DELETE ON questionnaire_definition_snapshots
        BEGIN
            SELECT RAISE(ABORT, 'questionnaire definition snapshots are immutable');
        END
        """,
    ),
)

NORMALIZED_QUESTIONNAIRE_STORAGE_MIGRATION = SchemaMigration(
    migration_id="len.014.002.normalized_questionnaire_storage",
    statements=(
        """
        CREATE TABLE questionnaire_submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            questionnaire_id TEXT NOT NULL,
            definition_version INTEGER NOT NULL CHECK(definition_version > 0),
            entry_date TEXT NOT NULL,
            total_score NUMERIC,
            severity TEXT NOT NULL,
            notes TEXT,
            note_tag TEXT,
            source TEXT,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            UNIQUE (questionnaire_id, entry_date),
            FOREIGN KEY (questionnaire_id, definition_version)
                REFERENCES questionnaire_definition_snapshots (questionnaire_id, definition_version)
                ON UPDATE RESTRICT ON DELETE RESTRICT
        )
        """,
        """
        CREATE TABLE questionnaire_responses (
            submission_id INTEGER NOT NULL,
            question_id TEXT NOT NULL,
            response_order INTEGER NOT NULL CHECK(response_order > 0),
            response_value_json TEXT NOT NULL,
            response_score NUMERIC,
            PRIMARY KEY (submission_id, question_id),
            UNIQUE (submission_id, response_order),
            FOREIGN KEY (submission_id) REFERENCES questionnaire_submissions (id)
                ON UPDATE RESTRICT ON DELETE CASCADE
        )
        """,
    ),
)

SCHEMA_MIGRATIONS = (
    QUESTIONNAIRE_SNAPSHOT_MIGRATION,
    NORMALIZED_QUESTIONNAIRE_STORAGE_MIGRATION,
)


def _store_questionnaire_definition_snapshots(
    conn: sqlite3.Connection,
    definitions: tuple[QuestionnaireDefinition, ...],
) -> None:
    for definition in definitions:
        payload = serialize_questionnaire_definition(definition)
        payload_hash = hashlib.sha256(payload.encode("utf-8")).hexdigest()
        existing = conn.execute(
            """
            SELECT definition_json, definition_sha256
            FROM questionnaire_definition_snapshots
            WHERE questionnaire_id = ? AND definition_version = ?
            """,
            (definition.questionnaire_id, definition.definition_version),
        ).fetchone()
        if existing is None:
            conn.execute(
                """
                INSERT INTO questionnaire_definition_snapshots (
                    questionnaire_id, definition_version, definition_json, definition_sha256
                )
                VALUES (?, ?, ?, ?)
                """,
                (definition.questionnaire_id, definition.definition_version, payload, payload_hash),
            )
        elif existing != (payload, payload_hash):
            raise RuntimeError(
                "Questionnaire definition version conflict for "
                f"{definition.questionnaire_id} v{definition.definition_version}."
            )


def _backfill_normalized_questionnaire_entries(conn: sqlite3.Connection) -> None:
    legacy_rows = conn.execute(
        """
        SELECT
            assessment_id, entry_date,
            item1, item2, item3, item4, item5, item6, item7, item8, item9,
            total, severity, notes, note_tag, source, created_at, updated_at
        FROM assessment_entries
        ORDER BY assessment_id, entry_date
        """
    ).fetchall()
    for row in legacy_rows:
        questionnaire_id, entry_date = row[0], row[1]
        try:
            legacy_definition = LEGACY_BUILTIN_DEFINITIONS[questionnaire_id]
            active_definition = QUESTIONNAIRES[questionnaire_id]
        except KeyError as exc:
            raise RuntimeError(f"Cannot backfill unknown questionnaire: {questionnaire_id}") from exc

        existing_submission = conn.execute(
            """
            SELECT definition_version
            FROM questionnaire_submissions
            WHERE questionnaire_id = ? AND entry_date = ?
            """,
            (questionnaire_id, entry_date),
        ).fetchone()
        if existing_submission is None or existing_submission[0] == legacy_definition.definition_version:
            definition = legacy_definition
        elif existing_submission[0] == active_definition.definition_version:
            definition = active_definition
        else:
            raise RuntimeError(
                "Cannot reconcile unsupported questionnaire definition version: "
                f"{questionnaire_id} v{existing_submission[0]}."
            )

        stored_items = list(row[2:11])
        responses = stored_items[: definition.item_count]
        if any(value is None for value in responses):
            raise RuntimeError(f"Legacy submission has unanswered required items: {questionnaire_id} {entry_date}")
        if any(value is not None for value in stored_items[definition.item_count :]):
            raise RuntimeError(f"Legacy submission has unexpected extra items: {questionnaire_id} {entry_date}")

        recorded_total, recorded_severity = row[11], row[12]
        calculated_total = calculate_questionnaire_total(definition, responses)
        if recorded_total != calculated_total:
            raise RuntimeError(f"Legacy total does not reconcile: {questionnaire_id} {entry_date}")
        expected_severity = severity_for_assessment(questionnaire_id, recorded_total)
        if recorded_severity != expected_severity:
            raise RuntimeError(f"Legacy severity does not reconcile: {questionnaire_id} {entry_date}")

        submission_values = (
            questionnaire_id,
            definition.definition_version,
            entry_date,
            recorded_total,
            recorded_severity,
            row[13],
            row[14],
            row[15],
            row[16],
            row[17],
        )
        conn.execute(
            """
            INSERT OR IGNORE INTO questionnaire_submissions (
                questionnaire_id, definition_version, entry_date, total_score, severity,
                notes, note_tag, source, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            submission_values,
        )
        submission = conn.execute(
            """
            SELECT
                id, questionnaire_id, definition_version, entry_date, total_score, severity,
                notes, note_tag, source, created_at, updated_at
            FROM questionnaire_submissions
            WHERE questionnaire_id = ? AND entry_date = ?
            """,
            (questionnaire_id, entry_date),
        ).fetchone()
        if submission is None or submission[1:] != submission_values:
            raise RuntimeError(f"Normalized submission does not reconcile: {questionnaire_id} {entry_date}")

        expected_responses = []
        for response_order, (question, response) in enumerate(zip(definition.items, responses), start=1):
            options = [option for option in question.options if option.value == response]
            if len(options) != 1:
                raise RuntimeError(f"Legacy response is not in its definition: {question.question_id} {entry_date}")
            response_values = (
                submission[0],
                question.question_id,
                response_order,
                json.dumps(response, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
                options[0].score,
            )
            expected_responses.append(response_values[1:])
            conn.execute(
                """
                INSERT OR IGNORE INTO questionnaire_responses (
                    submission_id, question_id, response_order, response_value_json, response_score
                ) VALUES (?, ?, ?, ?, ?)
                """,
                response_values,
            )

        stored_responses = conn.execute(
            """
            SELECT question_id, response_order, response_value_json, response_score
            FROM questionnaire_responses
            WHERE submission_id = ?
            ORDER BY response_order
            """,
            (submission[0],),
        ).fetchall()
        if stored_responses != expected_responses:
            raise RuntimeError(f"Normalized responses do not reconcile: {questionnaire_id} {entry_date}")


def apply_schema_migrations(
    conn: sqlite3.Connection,
    migrations: tuple[SchemaMigration, ...] = SCHEMA_MIGRATIONS,
    definitions: tuple[QuestionnaireDefinition, ...] | None = None,
) -> None:
    """Apply known migrations and exact definition snapshots as one transaction."""
    if conn.in_transaction:
        raise RuntimeError("Schema migrations require a connection without an active transaction.")
    definitions = BUILTIN_DEFINITION_SNAPSHOTS if definitions is None else definitions
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("BEGIN IMMEDIATE")
    try:
        for statement in SCHEMA_MIGRATION_LEDGER_STATEMENTS:
            conn.execute(statement)

        known_migration_ids = {migration.migration_id for migration in migrations}
        recorded_rows = conn.execute(
            "SELECT migration_id, application_id, checksum_sha256 FROM schema_migrations"
        ).fetchall()
        unknown_ids = sorted(row[0] for row in recorded_rows if row[0] not in known_migration_ids)
        if unknown_ids:
            raise RuntimeError(f"Database contains an unsupported schema migration: {unknown_ids[0]}")

        recorded = {row[0]: (row[1], row[2]) for row in recorded_rows}
        for migration in migrations:
            existing = recorded.get(migration.migration_id)
            if existing is not None:
                if existing != ("len", migration.checksum):
                    raise RuntimeError(f"Schema migration checksum mismatch: {migration.migration_id}")
                continue
            for statement in migration.statements:
                conn.execute(statement)
            conn.execute(
                """
                INSERT INTO schema_migrations (migration_id, application_id, checksum_sha256)
                VALUES (?, 'len', ?)
                """,
                (migration.migration_id, migration.checksum),
            )

        if QUESTIONNAIRE_SNAPSHOT_MIGRATION.migration_id in known_migration_ids:
            _store_questionnaire_definition_snapshots(conn, definitions)
        if NORMALIZED_QUESTIONNAIRE_STORAGE_MIGRATION.migration_id in known_migration_ids:
            _backfill_normalized_questionnaire_entries(conn)
        conn.commit()
    except Exception:
        conn.rollback()
        raise


def load_questionnaire_definition_snapshot(
    questionnaire_id: str,
    definition_version: int,
    db_path: Path | None = None,
) -> QuestionnaireDefinition:
    """Load and verify the immutable definition needed to interpret historical data."""
    with closing(sqlite3.connect(db_path or DB_PATH)) as conn:
        row = conn.execute(
            """
            SELECT definition_json, definition_sha256
            FROM questionnaire_definition_snapshots
            WHERE questionnaire_id = ? AND definition_version = ?
            """,
            (questionnaire_id, definition_version),
        ).fetchone()
    if row is None:
        raise KeyError(f"No questionnaire definition snapshot for {questionnaire_id} v{definition_version}.")
    payload, recorded_hash = row
    actual_hash = hashlib.sha256(payload.encode("utf-8")).hexdigest()
    if actual_hash != recorded_hash:
        raise RuntimeError(f"Questionnaire definition snapshot hash mismatch for {questionnaire_id} v{definition_version}.")
    definition = deserialize_questionnaire_definition(payload)
    if (definition.questionnaire_id, definition.definition_version) != (questionnaire_id, definition_version):
        raise RuntimeError(f"Questionnaire definition snapshot identity mismatch for {questionnaire_id} v{definition_version}.")
    return definition


def normalize_questionnaire_selection(questionnaire_ids=None) -> list[str]:
    selected = list(QUESTIONNAIRE_ORDER if questionnaire_ids is None else questionnaire_ids)
    unknown = [questionnaire_id for questionnaire_id in selected if questionnaire_id not in QUESTIONNAIRES]
    if unknown:
        raise ValueError(f"Unknown questionnaire: {unknown[0]}")
    if not selected:
        raise ValueError("Select at least one questionnaire.")
    return list(dict.fromkeys(selected))


def normalize_report_date_range(start: str, end: str) -> tuple[str, str]:
    """Validate and normalize the report range before questionnaire selection."""
    try:
        start_date = datetime.strptime(start, "%Y-%m-%d").date()
        end_date = datetime.strptime(end, "%Y-%m-%d").date()
    except (TypeError, ValueError) as exc:
        raise ValueError("Enter report dates as YYYY-MM-DD.") from exc
    if start_date > end_date:
        raise ValueError("The report start date must be on or before the end date.")
    return start_date.isoformat(), end_date.isoformat()


def eligible_report_questionnaires(start: str, end: str) -> list[str]:
    """Return registry questionnaires that have records in the validated range."""
    start, end = normalize_report_date_range(start, end)
    return [
        questionnaire_id
        for questionnaire_id in QUESTIONNAIRE_ORDER
        if fetch_assessment_entries(questionnaire_id, start, end)
    ]


def resolve_scoring_strategy(definition: QuestionnaireDefinition) -> ScoringStrategyDescriptor:
    """Resolve a definition's data-only ID to an application-owned descriptor."""
    try:
        return SCORING_STRATEGY_DESCRIPTORS[definition.scoring_rule]
    except KeyError as exc:
        raise ValueError(f"Unknown scoring rule: {definition.scoring_rule}") from exc


def resolve_profile_strategy(definition: QuestionnaireDefinition) -> ProfileStrategyDescriptor:
    """Resolve a definition's data-only ID to an application-owned descriptor."""
    try:
        return PROFILE_STRATEGY_DESCRIPTORS[definition.profile_rule]
    except KeyError as exc:
        raise ValueError(f"Unknown profile rule: {definition.profile_rule}") from exc


def calculate_questionnaire_total(
    definition: QuestionnaireDefinition,
    responses: list[str | int | None] | tuple[str | int | None, ...],
) -> int | float | None:
    """Calculate a deterministic total without executing definition-provided logic."""
    validate_questionnaire_definition(definition)
    strategy = resolve_scoring_strategy(definition)
    if strategy.implementation_id == "no_total":
        return None
    if strategy.implementation_id != "sum_option_scores_v1":
        raise ValueError(f"Scoring strategy is not implemented: {strategy.strategy_id}")
    if len(responses) != definition.item_count:
        raise ValueError(f"{definition.display_name} requires {definition.item_count} responses.")

    scores: list[int | float] = []
    for item, response in zip(definition.items, responses):
        matching_options = [option for option in item.options if option.value == response]
        if len(matching_options) != 1 or matching_options[0].score is None:
            raise ValueError(f"Unsupported response for {item.question_id}: {response!r}")
        scores.append(matching_options[0].score)
    return sum(scores)


def questionnaire_response_display(option: ResponseOption) -> str:
    return f"{option.label} [{option.value}]"


def questionnaire_response_choices(question: QuestionDefinition) -> tuple[str, ...]:
    return (UNANSWERED_RESPONSE, *(questionnaire_response_display(option) for option in question.options))


def parse_questionnaire_response(question: QuestionDefinition, displayed_value: str) -> str | int | None:
    if displayed_value == UNANSWERED_RESPONSE:
        return None
    for option in question.options:
        if displayed_value == questionnaire_response_display(option):
            return option.value
    raise ValueError(f"Unknown response for {question.question_id}.")


def display_questionnaire_response(question: QuestionDefinition, response: str | int) -> str:
    for option in question.options:
        if option.value == response:
            return questionnaire_response_display(option)
    raise ValueError(f"Unknown response for {question.question_id}.")


def user_facing_response_label(
    definition: QuestionnaireDefinition,
    option: ResponseOption,
) -> str:
    """Avoid presenting the known-regressed v1 labels as daily-severity meaning."""
    if definition.questionnaire_id in LEGACY_BUILTIN_DEFINITIONS and definition.definition_version == 1:
        return str(option.value)
    return option.label


def questionnaire_completion_status(entry_date: str) -> dict[str, str]:
    completed = fetch_day_data(entry_date)["assessments"]
    return {
        questionnaire_id: "Complete" if questionnaire_id in completed else "Not completed"
        for questionnaire_id in QUESTIONNAIRE_ORDER
    }

@dataclass
class EntryRow:
    id: int
    entry_date: str
    items: list[int]
    total: int
    severity: str
    notes: str
    note_tag: str = ""


@dataclass
class AssessmentEntryRow:
    id: int
    assessment_id: str
    entry_date: str
    items: list[int]
    total: int
    severity: str
    notes: str
    note_tag: str = ""
    definition_version: int | None = None


@dataclass
class FourteenDayScore:
    item_counts: list[int]
    item_scores: list[int]
    total_score: int
    severity: str
    start_date: str | None
    end_date: str | None
    entries_included: int
    calendar_days: int = 14


@dataclass(frozen=True)
class PeriodComparison:
    assessment_id: str
    current: FourteenDayScore
    previous: FourteenDayScore
    current_start: str
    current_end: str
    previous_start: str
    previous_end: str

    @property
    def delta(self) -> int:
        return self.current.total_score - self.previous.total_score

    @property
    def has_comparable_data(self) -> bool:
        return self.current.entries_included > 0 and self.previous.entries_included > 0


@dataclass(frozen=True)
class QuestionnaireTrendSeries:
    """One explicitly compatible total-score series for neutral charting."""

    definition: QuestionnaireDefinition
    entries: tuple[AssessmentEntryRow, ...]
    definition_versions: tuple[int, ...]

    @property
    def label(self) -> str:
        if len(self.definition_versions) > 1:
            return self.definition.display_name
        return f"{self.definition.display_name} v{self.definition.definition_version}"


@dataclass(frozen=True)
class TreatmentCycle:
    """A descriptive window anchored to a recorded ketamine infusion."""

    label: str
    start_date: str
    end_date: str
    entries: list[EntryRow | AssessmentEntryRow]


@dataclass(frozen=True)
class ChartPoint:
    """One discoverable score point rendered on a Review chart."""

    x: float
    y: float
    entry_date: str
    series_name: str
    score: int
    color: str


@dataclass(frozen=True)
class Item9Context:
    """Recency and coverage facts for PHQ-9 item 9 discussion language."""

    recent_start: str
    recent_end: str
    recent_checkins: int
    recent_above_zero: int
    historical_above_zero: int
    latest_historical_date: str | None

    @property
    def has_recent_above_zero(self) -> bool:
        return self.recent_above_zero > 0

    @property
    def has_historical_above_zero(self) -> bool:
        return self.historical_above_zero > 0


def severity_for_score(score: int) -> str:
    if score <= 4:
        return "Minimal"
    if score <= 9:
        return "Mild"
    if score <= 14:
        return "Moderate"
    if score <= 19:
        return "Moderately severe"
    return "Severe"


def gad7_severity_for_score(score: int) -> str:
    if score <= 4:
        return "Minimal"
    if score <= 9:
        return "Mild"
    if score <= 14:
        return "Moderate"
    return "Severe"


def severity_for_assessment(assessment_id: str, score: int) -> str:
    if assessment_id == "gad7":
        return gad7_severity_for_score(score)
    return severity_for_score(score)


def convert_14_day_count_to_item_score(days_present: int) -> int:
    if days_present <= 0:
        return 0
    if days_present <= 6:
        return 1
    if days_present <= 11:
        return 2
    return 3


def _require_symptom_presence_profile(definition: QuestionnaireDefinition) -> ProfileStrategyDescriptor:
    validate_questionnaire_definition(definition)
    strategy = resolve_profile_strategy(definition)
    if strategy.implementation_id != "symptom_presence_thresholds_14d_v1":
        raise ValueError(f"Profile strategy is not implemented for 14-day scoring: {strategy.strategy_id}")
    return strategy


def _profile_response_score(question: QuestionDefinition, response: object) -> int | float:
    matches = [option for option in question.options if option.value == response]
    if len(matches) != 1:
        raise ValueError(f"Unsupported response for {question.question_id}: {response!r}")
    score = matches[0].score
    if score is None:
        raise ValueError(f"Profile strategy requires numeric option scores: {question.question_id}")
    return score


def _profile_response_is_present(question: QuestionDefinition, response: object) -> bool:
    return _profile_response_score(question, response) > 0


def _calculate_symptom_presence_profile(
    definition: QuestionnaireDefinition,
    entries: list[EntryRow | AssessmentEntryRow],
    window_start: str,
    window_end: str,
    item_count: int | None = None,
) -> FourteenDayScore:
    _require_symptom_presence_profile(definition)
    start_dt = datetime.fromisoformat(window_start).date()
    end_dt = datetime.fromisoformat(window_end).date()
    if end_dt < start_dt:
        raise ValueError("Window end date must not be before its start date.")
    window_entries = [
        row
        for row in entries
        if start_dt <= datetime.fromisoformat(row.entry_date).date() <= end_dt
    ]
    resolved_item_count = definition.item_count if item_count is None else item_count
    item_counts = [
        sum(
            1
            for row in window_entries
            if idx < len(row.items) and _profile_response_is_present(definition.items[idx], row.items[idx])
        )
        for idx in range(resolved_item_count)
    ]
    item_scores = [convert_14_day_count_to_item_score(count) for count in item_counts]
    total_score = sum(item_scores)
    severity = ""
    if definition.interpretation_policy == "validated_builtin":
        severity = severity_for_assessment(definition.questionnaire_id, total_score)
    return FourteenDayScore(
        item_counts=item_counts,
        item_scores=item_scores,
        total_score=total_score,
        severity=severity,
        start_date=window_start,
        end_date=window_end,
        entries_included=len(window_entries),
        calendar_days=(end_dt - start_dt).days + 1,
    )


def calculate_questionnaire_profile_for_window(
    definition: QuestionnaireDefinition,
    entries: list[EntryRow | AssessmentEntryRow],
    window_start: str,
    window_end: str,
) -> FourteenDayScore:
    """Resolve and execute the application-owned profile strategy for a definition."""
    return _calculate_symptom_presence_profile(definition, entries, window_start, window_end)


def calculate_14_day_symptom_frequency_score(
    entries: list[EntryRow | AssessmentEntryRow],
    item_count: int | None = None,
    assessment_id: str = "phq9",
) -> FourteenDayScore:
    """Calculate the clinical-style 14-day symptom-frequency score.

    The window is the most recent 14 calendar days ending on the most recent
    entry date. For each item, each day with an entry and item score > 0 counts
    as one symptom-present day. Missing calendar days are treated as no recorded
    symptom-present day for the count, and entries_included reports how many
    actual entry rows were available in the window.
    """
    definition = ASSESSMENTS[assessment_id]
    _require_symptom_presence_profile(definition)
    item_count = item_count or (len(entries[0].items) if entries else definition.item_count)
    if not entries:
        return FourteenDayScore(
            [0] * item_count,
            [0] * item_count,
            0,
            severity_for_assessment(assessment_id, 0),
            None,
            None,
            0,
        )
    sorted_entries = sorted(entries, key=lambda row: row.entry_date)
    end_dt = datetime.fromisoformat(sorted_entries[-1].entry_date).date()
    start_dt = end_dt - timedelta(days=13)
    return _calculate_symptom_presence_profile(
        definition,
        sorted_entries,
        start_dt.isoformat(),
        end_dt.isoformat(),
        item_count,
    )


def calculate_symptom_frequency_score_for_window(
    entries: list[EntryRow | AssessmentEntryRow],
    window_start: str,
    window_end: str,
    item_count: int,
    assessment_id: str,
) -> FourteenDayScore:
    """Calculate a symptom-frequency score for an explicit calendar window."""
    return _calculate_symptom_presence_profile(
        ASSESSMENTS[assessment_id],
        entries,
        window_start,
        window_end,
        item_count,
    )


def questionnaire_total_trend_omission_reason(definition: QuestionnaireDefinition) -> str | None:
    """Return why a definition cannot produce a total-score trend, or None."""
    validate_questionnaire_definition(definition)
    strategy = resolve_scoring_strategy(definition)
    if strategy.implementation_id == "no_total":
        return f"{definition.display_name} does not define a total score, so no total-score trend is shown."
    if strategy.implementation_id != "sum_option_scores_v1":
        return f"{definition.display_name} does not have an implemented total-score trend capability."
    if definition.score_min is None or definition.score_max is None or definition.score_max <= definition.score_min:
        return f"{definition.display_name} does not declare usable score bounds for a total-score trend."
    return None


def questionnaire_profile_omission_reason(definition: QuestionnaireDefinition) -> str | None:
    """Return why a definition cannot produce a 14-day item profile, or None."""
    validate_questionnaire_definition(definition)
    strategy = resolve_profile_strategy(definition)
    if strategy.implementation_id == "no_profile":
        return f"{definition.display_name} does not define a 14-day item profile."
    if strategy.implementation_id != "symptom_presence_thresholds_14d_v1":
        return f"{definition.display_name} does not have an implemented 14-day item profile capability."
    return None


def _definition_groups_for_entries(
    questionnaire_id: str,
    entries: list[AssessmentEntryRow],
) -> list[tuple[QuestionnaireDefinition, list[AssessmentEntryRow]]]:
    """Group rows by their immutable definition without blending versions."""
    current_definition = QUESTIONNAIRES[questionnaire_id]
    grouped: dict[int, list[AssessmentEntryRow]] = {}
    definitions: dict[int, QuestionnaireDefinition] = {}
    for entry in entries:
        if entry.assessment_id != questionnaire_id:
            raise ValueError(f"Entry questionnaire does not match {questionnaire_id}: {entry.assessment_id}")
        version = entry.definition_version or current_definition.definition_version
        grouped.setdefault(version, []).append(entry)
        if version not in definitions:
            definitions[version] = (
                current_definition
                if version == current_definition.definition_version
                else load_questionnaire_definition_snapshot(questionnaire_id, version)
            )
    if not grouped:
        return [(current_definition, [])]
    return [(definitions[version], grouped[version]) for version in sorted(grouped)]


ANALYTICAL_SERIES_IDS = MappingProxyType({
    (questionnaire_id, version): f"{questionnaire_id}:daily_severity_v1"
    for questionnaire_id in ("phq9", "gad7")
    for version in (1, 2)
})


def analytical_series_id(questionnaire_id: str, definition_version: int) -> str:
    return ANALYTICAL_SERIES_IDS.get(
        (questionnaire_id, definition_version), f"{questionnaire_id}:v{definition_version}"
    )


def _entry_analytical_series_id(entry: AssessmentEntryRow) -> str:
    version = entry.definition_version or QUESTIONNAIRES[entry.assessment_id].definition_version
    return analytical_series_id(entry.assessment_id, version)


def _analytical_signature(definition: QuestionnaireDefinition) -> tuple:
    """Fields that must agree before score and item profiles can share a series."""
    return (
        definition.scoring_rule, definition.score_min, definition.score_max,
        definition.profile_rule, definition.interpretation_policy, definition.behavior_ids,
        tuple((item.question_id, item.behavior_ids,
               tuple((option.value, option.score) for option in item.options))
              for item in definition.items),
    )


def _analytical_groups_for_entries(
    questionnaire_id: str,
    entries: list[AssessmentEntryRow],
) -> list[tuple[str, QuestionnaireDefinition, list[AssessmentEntryRow], tuple[int, ...]]]:
    """Join only approved versions; unknown versions remain separate by default."""
    grouped: dict[str, list[tuple[QuestionnaireDefinition, list[AssessmentEntryRow]]]] = {}
    for definition, definition_entries in _definition_groups_for_entries(questionnaire_id, entries):
        series_id = analytical_series_id(questionnaire_id, definition.definition_version)
        grouped.setdefault(series_id, []).append((definition, definition_entries))
    result = []
    for series_id, members in grouped.items():
        if len({_analytical_signature(definition) for definition, _ in members}) != 1:
            raise ValueError(f"Incompatible definitions assigned to analytical series {series_id}.")
        definition = max((definition for definition, _ in members), key=lambda item: item.definition_version)
        versions = tuple(sorted(member.definition_version for member, _ in members))
        combined = sorted(
            (entry for _, member_entries in members for entry in member_entries),
            key=lambda entry: (entry.entry_date, entry.id),
        )
        result.append((series_id, definition, combined, versions))
    return sorted(result, key=lambda item: max((entry.entry_date for entry in item[2]), default=""))


def build_questionnaire_trend_series(
    questionnaire_id: str,
    entries: list[AssessmentEntryRow] | None = None,
) -> list[QuestionnaireTrendSeries]:
    """Build approved compatible neutral total-score series."""
    if questionnaire_id not in QUESTIONNAIRES:
        raise ValueError(f"Unknown questionnaire: {questionnaire_id}")
    resolved_entries = fetch_assessment_entries(questionnaire_id) if entries is None else entries
    series = []
    for _, definition, definition_entries, versions in _analytical_groups_for_entries(questionnaire_id, resolved_entries):
        if questionnaire_total_trend_omission_reason(definition) is None:
            series.append(QuestionnaireTrendSeries(definition, tuple(definition_entries), versions))
    return series


def _latest_definition_group(
    questionnaire_id: str,
    entries: list[AssessmentEntryRow],
    end_date: str,
) -> tuple[QuestionnaireDefinition, list[AssessmentEntryRow], bool]:
    """Select the latest compatible series without blending unlike scoring semantics."""
    groups = _analytical_groups_for_entries(questionnaire_id, entries)
    eligible = [
        (definition, [entry for entry in group_entries if entry.entry_date <= end_date])
        for _, definition, group_entries, _ in groups
    ]
    eligible = [(definition, group_entries) for definition, group_entries in eligible if group_entries]
    if not eligible:
        return QUESTIONNAIRES[questionnaire_id], [], len(groups) > 1
    definition, group_entries = max(eligible, key=lambda item: max(entry.entry_date for entry in item[1]))
    return definition, group_entries, len(groups) > 1


def build_14_day_item_profile(end_date: str, questionnaire_ids=None) -> list[dict[str, object]]:
    """Build one derived current-window record per selected questionnaire item."""
    window_end = datetime.fromisoformat(end_date).date()
    window_start = (window_end - timedelta(days=13)).isoformat()
    records: list[dict[str, object]] = []
    for assessment_id in normalize_questionnaire_selection(questionnaire_ids):
        entries = fetch_assessment_entries(assessment_id, window_start, end_date)
        analytical_groups = _analytical_groups_for_entries(assessment_id, entries)
        multiple_series = len(analytical_groups) > 1
        for series_id, definition, definition_entries, versions in analytical_groups:
            if questionnaire_profile_omission_reason(definition) is not None:
                continue
            score = calculate_questionnaire_profile_for_window(
                definition,
                definition_entries,
                window_start,
                end_date,
            )
            version_component = f":{series_id}" if multiple_series else ""
            for item_index, item_label in enumerate(definition.item_labels):
                records.append(
                    {
                        "profile_record_id": f"profile:{end_date}:{assessment_id}{version_component}:item:{item_index + 1}",
                        "window_start": window_start,
                        "window_end": end_date,
                        "assessment_id": assessment_id,
                        "assessment_name": definition.display_name,
                        "definition_version": versions[0] if len(versions) == 1 else None,
                        "definition_versions": "|".join(f"v{version}" for version in versions),
                        "analytical_series_id": series_id,
                        "item_number": item_index + 1,
                        "item_label": item_label,
                        "symptom_present_days": score.item_counts[item_index],
                        "recorded_day_coverage": score.entries_included,
                        "calendar_days": score.calendar_days,
                        "frequency_score": score.item_scores[item_index],
                    }
                )
    return records


def compare_recent_14_day_periods(
    entries: list[EntryRow | AssessmentEntryRow],
    assessment_id: str,
    end_date: str,
    definition: QuestionnaireDefinition | None = None,
) -> PeriodComparison:
    """Compare the latest 14 calendar days with the immediately prior 14 days."""
    definition = definition or ASSESSMENTS[assessment_id]
    current_end = datetime.fromisoformat(end_date).date()
    current_start = current_end - timedelta(days=13)
    previous_end = current_start - timedelta(days=1)
    previous_start = previous_end - timedelta(days=13)
    current = calculate_questionnaire_profile_for_window(
        definition, entries, current_start.isoformat(), current_end.isoformat()
    )
    previous = calculate_questionnaire_profile_for_window(
        definition, entries, previous_start.isoformat(), previous_end.isoformat()
    )
    return PeriodComparison(
        assessment_id=assessment_id,
        current=current,
        previous=previous,
        current_start=current_start.isoformat(),
        current_end=current_end.isoformat(),
        previous_start=previous_start.isoformat(),
        previous_end=previous_end.isoformat(),
    )


def comparison_trend_text(comparison: PeriodComparison, include_arrow: bool = False) -> str:
    """Return neutral, non-diagnostic wording for a two-period comparison."""
    if not comparison.has_comparable_data:
        return "Not enough data for comparison"
    if comparison.delta < 0:
        return f"{'Down - ' if include_arrow else ''}Lower by {abs(comparison.delta)} points"
    if comparison.delta > 0:
        return f"{'Up - ' if include_arrow else ''}Higher by {comparison.delta} points"
    return f"{'Right - ' if include_arrow else ''}No score change"


def shift_calendar_date(value: str, days: int, today: date | None = None) -> str:
    """Move an ISO date by whole calendar days without allowing future dates."""
    parsed = parse_date(value)
    if not parsed:
        raise ValueError("Enter the date as YYYY-MM-DD.")
    shifted = datetime.fromisoformat(parsed).date() + timedelta(days=days)
    if shifted > (today or date.today()):
        raise ValueError("Future check-ins are not available.")
    return shifted.isoformat()


def validate_nonfuture_date(value: str, today: date | None = None) -> str:
    """Normalize an ISO date and reject dates after the local current day."""
    parsed = parse_date(value)
    if not parsed:
        raise ValueError("Enter the date as YYYY-MM-DD.")
    if datetime.fromisoformat(parsed).date() > (today or date.today()):
        raise ValueError("Future dates are not available.")
    return parsed


def history_record_count(day_data: dict) -> int:
    """Count the independently manageable records represented by one day."""
    return len(day_data["assessments"]) + len(day_data["events"]) + int(bool(day_data["notes"] or day_data["note_tag"]))


def history_status_text(entry_date: str, day_data: dict) -> str:
    """Return calm History status text without implying that a record failed to load."""
    record_count = history_record_count(day_data)
    if record_count == 0:
        return (
            f"No existing records for {entry_date}. History can edit only dates that already contain "
            "a check-in, note, or treatment event."
        )
    noun = "record" if record_count == 1 else "records"
    return f"Showing {record_count} existing {noun} for {entry_date}. Edits apply only to this loaded date."


def responsive_canvas_dimension(actual: int, configured: int) -> int:
    """Use live canvas dimensions after layout, including when the window shrinks."""
    return actual if actual > 1 else configured


def chart_vertical_positions(title_bottom: int) -> tuple[int, int]:
    """Keep the legend below the rendered title and the plot below the legend."""
    legend_top = title_bottom + 8
    plot_top = legend_top + 24
    return legend_top, plot_top


def chart_point_callout_text(point: ChartPoint) -> str:
    """Return the exact date and score shown by chart point callouts."""
    return f"{point.entry_date}  |  {point.series_name}: {point.score}"


def nearest_chart_point(
    points: list[ChartPoint],
    x: float,
    y: float,
    max_distance: float = 12,
) -> ChartPoint | None:
    """Find a chart point only when the pointer is within a forgiving hit target."""
    if not points:
        return None
    nearest = min(points, key=lambda point: (point.x - x) ** 2 + (point.y - y) ** 2)
    distance_squared = (nearest.x - x) ** 2 + (nearest.y - y) ** 2
    return nearest if distance_squared <= max_distance**2 else None


def entries_for_window(entries, start_date: str, end_date: str):
    return [row for row in entries if start_date <= row.entry_date <= end_date]


def item9_context(entries: list[AssessmentEntryRow], end_date: str) -> Item9Context:
    """Separate recent item 9 responses from older selected-history context."""
    recent_end = datetime.fromisoformat(end_date).date()
    recent_start = recent_end - timedelta(days=13)
    recent_start_text = recent_start.isoformat()
    recent_end_text = recent_end.isoformat()
    recent_entries = entries_for_window(entries, recent_start_text, recent_end_text)
    historical_positive_entries = [
        row for row in entries if row.entry_date < recent_start_text and row.items[8] > 0
    ]
    return Item9Context(
        recent_start=recent_start_text,
        recent_end=recent_end_text,
        recent_checkins=len(recent_entries),
        recent_above_zero=sum(row.items[8] > 0 for row in recent_entries),
        historical_above_zero=len(historical_positive_entries),
        latest_historical_date=(
            max(row.entry_date for row in historical_positive_entries)
            if historical_positive_entries
            else None
        ),
    )


def item9_context_summary(context: Item9Context) -> str | None:
    """Describe item 9 recency without turning historical data into a current-risk claim."""
    coverage = f"Recent check-in coverage was {context.recent_checkins} of 14 calendar days"
    missing = "days without a check-in are missing information"
    if context.has_recent_above_zero:
        noun = "check-in" if context.recent_above_zero == 1 else "check-ins"
        return (
            f"An above-zero PHQ-9 item 9 response was recorded on {context.recent_above_zero} {noun} "
            f"in the most recent 14-day window ({context.recent_start} to {context.recent_end}). "
            f"{coverage}; {missing}. This is a factual discussion point, not an assessment of current safety."
        )
    if context.has_historical_above_zero:
        noun = "check-in" if context.historical_above_zero == 1 else "check-ins"
        return (
            f"An above-zero PHQ-9 item 9 response was recorded earlier in the selected history on "
            f"{context.historical_above_zero} {noun}, most recently on {context.latest_historical_date}. "
            f"No above-zero item 9 response was recorded among {context.recent_checkins} PHQ-9 check-ins in the "
            f"most recent 14-day window ({context.recent_start} to {context.recent_end}). {coverage}; {missing}. "
            "This older information is neutral historical context and does not establish current risk."
        )
    return None


def item9_conversation_starter(context: Item9Context) -> str | None:
    """Return a recency-aware prompt only when item 9 was above zero in selected history."""
    if context.has_recent_above_zero:
        return (
            f"PHQ-9 item 9 was above zero on {context.recent_above_zero} of {context.recent_checkins} recorded "
            f"check-ins in the most recent 14-day window. Recent check-in coverage was "
            f"{context.recent_checkins} of 14 calendar days; missing days are missing information. "
            "Would it be useful to discuss when this was recorded and what support, if any, would be useful now?"
        )
    if context.has_historical_above_zero:
        return (
            f"An above-zero PHQ-9 item 9 response was recorded earlier in the selected history, most recently on "
            f"{context.latest_historical_date}. It was not recorded above zero among the "
            f"{context.recent_checkins} check-ins in the most recent 14-day window; recent check-in coverage was "
            f"{context.recent_checkins} of 14 calendar days. This is historical context and does not indicate "
            "current risk; would it be useful to discuss what has changed since then?"
        )
    return None


def overall_pattern_summary(
    phq_entries: list[AssessmentEntryRow],
    gad_entries: list[AssessmentEntryRow],
    end_date: str,
    questionnaire_ids=None,
    entry_sets: dict[str, list[AssessmentEntryRow]] | None = None,
) -> str:
    """Describe adjacent 14-day patterns without diagnosis or causal language."""
    phrases = []
    coverage = []
    entry_sets = entry_sets or {"phq9": phq_entries, "gad7": gad_entries}
    for assessment_id in normalize_questionnaire_selection(questionnaire_ids):
        definition, entries, has_multiple_versions = _latest_definition_group(
            assessment_id, entry_sets[assessment_id], end_date
        )
        comparison = compare_recent_14_day_periods(entries, assessment_id, end_date, definition)
        label = definition.display_name
        if has_multiple_versions:
            label = f"{label} v{definition.definition_version}"
        coverage.append(f"{label} {comparison.current.entries_included}/14")
        if not comparison.has_comparable_data:
            phrases.append(f"{label} does not yet have recorded check-ins in both comparison periods")
        elif comparison.delta < 0:
            phrases.append(f"{label} symptom-frequency scores were lower than in the preceding 14 days")
        elif comparison.delta > 0:
            phrases.append(f"{label} symptom-frequency scores were higher than in the preceding 14 days")
        else:
            phrases.append(f"{label} symptom-frequency scores were unchanged from the preceding 14 days")
    return f"{' '.join(f'{phrase}.' for phrase in phrases)} Current-period coverage: {', '.join(coverage)} recorded check-ins."


def symptom_highlights(
    assessment_id: str,
    entries: list[AssessmentEntryRow],
    end_date: str,
    limit: int = 2,
) -> list[str]:
    """Return the largest recorded item-frequency changes using neutral wording."""
    definition, entries, has_multiple_versions = _latest_definition_group(assessment_id, entries, end_date)
    comparison = compare_recent_14_day_periods(entries, assessment_id, end_date, definition)
    display_name = definition.display_name
    if has_multiple_versions:
        display_name = f"{display_name} v{definition.definition_version}"
    current_entries = entries_for_window(entries, comparison.current_start, comparison.current_end)
    previous_entries = entries_for_window(entries, comparison.previous_start, comparison.previous_end)
    if not current_entries:
        return [f"No {display_name} check-ins were recorded in the current 14-day period."]

    current_counts = [sum(row.items[idx] > 0 for row in current_entries) for idx in range(definition.item_count)]
    if not previous_entries:
        ranked = sorted(range(definition.item_count), key=lambda idx: (-current_counts[idx], idx))
        return [
            f"Responses related to {definition.item_labels[idx].lower()} were recorded on {current_counts[idx]} of {len(current_entries)} {display_name} check-ins in this period."
            for idx in ranked[:limit]
            if current_counts[idx] > 0
        ] or [f"No {display_name} symptoms were recorded as present in the current period."]

    previous_counts = [sum(row.items[idx] > 0 for row in previous_entries) for idx in range(definition.item_count)]
    ranked = sorted(
        range(definition.item_count),
        key=lambda idx: (-abs(current_counts[idx] - previous_counts[idx]), idx),
    )
    highlights = []
    for idx in ranked:
        delta = current_counts[idx] - previous_counts[idx]
        if delta == 0:
            continue
        direction = "more often" if delta > 0 else "less often"
        highlights.append(
            f"Responses related to {definition.item_labels[idx].lower()} were recorded {direction}: {current_counts[idx]} of {len(current_entries)} check-ins, compared with {previous_counts[idx]} of {len(previous_entries)} previously."
        )
        if len(highlights) == limit:
            break
    return highlights or [f"Recorded {display_name} symptom frequencies were similar across the two periods."]


def treatment_cycles(
    entries: list[EntryRow | AssessmentEntryRow],
    events: list[tuple[int, str, str, str]],
    end_date: str,
) -> list[TreatmentCycle]:
    """Return current and previous windows between recorded ketamine infusions."""
    anchors = sorted(
        {event_date for _, event_date, event_type, _ in events if normalize_event_type(event_type) == "Ketamine" and event_date <= end_date}
    )
    if not anchors:
        return []
    cycles = []
    latest = anchors[-1]
    cycles.append(TreatmentCycle("Current cycle", latest, end_date, entries_for_window(entries, latest, end_date)))
    if len(anchors) > 1:
        previous = anchors[-2]
        previous_end = (datetime.fromisoformat(latest).date() - timedelta(days=1)).isoformat()
        cycles.append(TreatmentCycle("Previous cycle", previous, previous_end, entries_for_window(entries, previous, previous_end)))
    return cycles


def treatment_cycle_observation(cycle: TreatmentCycle, assessment_name: str = "PHQ-9") -> str:
    if not cycle.entries:
        return f"{cycle.label} ({cycle.start_date} to {cycle.end_date}): no {assessment_name} check-ins were recorded."
    totals = [row.total for row in cycle.entries]
    return (
        f"{cycle.label} ({cycle.start_date} to {cycle.end_date}): {len(totals)} recorded {assessment_name} check-ins; "
        f"scores ranged from {min(totals)} to {max(totals)} with an average of {sum(totals) / len(totals):.1f}."
    )


def parse_date(value) -> str | None:
    if value is None or value == "":
        return None
    if hasattr(value, "date"):
        return value.date().isoformat()
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        return None


def init_db(db_path: Path | None = None) -> None:
    db_path = db_path or DB_PATH
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(db_path)) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS phq9_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                entry_date TEXT NOT NULL UNIQUE,
                item1 INTEGER NOT NULL,
                item2 INTEGER NOT NULL,
                item3 INTEGER NOT NULL,
                item4 INTEGER NOT NULL,
                item5 INTEGER NOT NULL,
                item6 INTEGER NOT NULL,
                item7 INTEGER NOT NULL,
                item8 INTEGER NOT NULL,
                item9 INTEGER NOT NULL,
                total INTEGER NOT NULL,
                severity TEXT NOT NULL,
                notes TEXT DEFAULT '',
                note_tag TEXT DEFAULT '',
                source TEXT DEFAULT 'manual',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        existing_columns = {row[1] for row in conn.execute("PRAGMA table_info(phq9_entries)").fetchall()}
        if "note_tag" not in existing_columns:
            conn.execute("ALTER TABLE phq9_entries ADD COLUMN note_tag TEXT DEFAULT ''")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS assessment_entries (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                assessment_id TEXT NOT NULL,
                entry_date TEXT NOT NULL,
                item1 INTEGER NOT NULL,
                item2 INTEGER NOT NULL,
                item3 INTEGER NOT NULL,
                item4 INTEGER NOT NULL,
                item5 INTEGER NOT NULL,
                item6 INTEGER NOT NULL,
                item7 INTEGER NOT NULL,
                item8 INTEGER,
                item9 INTEGER,
                total INTEGER NOT NULL,
                severity TEXT NOT NULL,
                notes TEXT DEFAULT '',
                note_tag TEXT DEFAULT '',
                source TEXT DEFAULT 'manual',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(assessment_id, entry_date)
            )
            """
        )
        conn.execute(
            """
            INSERT OR IGNORE INTO assessment_entries (
                assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                total, severity, notes, note_tag, source, created_at, updated_at
            )
            SELECT 'phq9', entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                   total, severity, COALESCE(notes, ''), COALESCE(note_tag, ''), source, created_at, updated_at
            FROM phq9_entries
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS treatment_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_date TEXT NOT NULL,
                event_type TEXT NOT NULL,
                description TEXT DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
            """
        )
        conn.commit()
        apply_schema_migrations(conn)


def _write_normalized_submission(
    conn: sqlite3.Connection,
    definition: QuestionnaireDefinition,
    entry_date: str,
    items: list[int],
) -> None:
    legacy = conn.execute(
        """
        SELECT total, severity, notes, note_tag, source, created_at, updated_at
        FROM assessment_entries
        WHERE assessment_id = ? AND entry_date = ?
        """,
        (definition.questionnaire_id, entry_date),
    ).fetchone()
    if legacy is None:
        raise RuntimeError(f"Compatibility row is missing for {definition.questionnaire_id} {entry_date}.")
    conn.execute(
        """
        INSERT INTO questionnaire_submissions (
            questionnaire_id, definition_version, entry_date, total_score, severity,
            notes, note_tag, source, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(questionnaire_id, entry_date) DO UPDATE SET
            definition_version=excluded.definition_version,
            total_score=excluded.total_score,
            severity=excluded.severity,
            notes=excluded.notes,
            note_tag=excluded.note_tag,
            source=excluded.source,
            created_at=excluded.created_at,
            updated_at=excluded.updated_at
        """,
        (definition.questionnaire_id, definition.definition_version, entry_date, *legacy),
    )
    submission_id = conn.execute(
        "SELECT id FROM questionnaire_submissions WHERE questionnaire_id = ? AND entry_date = ?",
        (definition.questionnaire_id, entry_date),
    ).fetchone()[0]
    conn.execute("DELETE FROM questionnaire_responses WHERE submission_id = ?", (submission_id,))
    for response_order, (question, response) in enumerate(zip(definition.items, items), start=1):
        option = next((option for option in question.options if option.value == response), None)
        if option is None:
            raise ValueError(f"Unsupported response for {question.question_id}: {response!r}")
        conn.execute(
            """
            INSERT INTO questionnaire_responses (
                submission_id, question_id, response_order, response_value_json, response_score
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (
                submission_id,
                question.question_id,
                response_order,
                json.dumps(response, ensure_ascii=False, sort_keys=True, separators=(",", ":")),
                option.score,
            ),
        )


def _upsert_assessment_entry_in_connection(
    conn: sqlite3.Connection,
    assessment_id: str,
    entry_date: str,
    items: list[int],
    notes: str,
    source: str,
    note_tag: str,
) -> None:
    definition = ASSESSMENTS[assessment_id]
    if len(items) != definition.item_count:
        raise ValueError(f"{definition.display_name} requires {definition.item_count} items.")
    if any(score < 0 or score > 3 for score in items):
        raise ValueError(f"{definition.display_name} item scores must be 0, 1, 2, or 3.")
    padded_items = [*items, *([None] * (9 - len(items)))]
    total = calculate_questionnaire_total(definition, items)
    if total is None:
        raise ValueError(f"{definition.display_name} does not define a total score.")
    severity = severity_for_assessment(assessment_id, total)
    conn.execute(
        """
        INSERT INTO assessment_entries (
            assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
            total, severity, notes, note_tag, source
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(assessment_id, entry_date) DO UPDATE SET
            item1=excluded.item1,
            item2=excluded.item2,
            item3=excluded.item3,
            item4=excluded.item4,
            item5=excluded.item5,
            item6=excluded.item6,
            item7=excluded.item7,
            item8=excluded.item8,
            item9=excluded.item9,
            total=excluded.total,
            severity=excluded.severity,
            notes=CASE WHEN excluded.notes != '' THEN excluded.notes ELSE assessment_entries.notes END,
            note_tag=CASE WHEN excluded.note_tag != '' THEN excluded.note_tag ELSE assessment_entries.note_tag END,
            source=excluded.source,
            updated_at=CURRENT_TIMESTAMP
        """,
        (assessment_id, entry_date, *padded_items, total, severity, notes or "", note_tag or "", source),
    )
    _write_normalized_submission(conn, definition, entry_date, items)


def upsert_assessment_entry(
    assessment_id: str,
    entry_date: str,
    items: list[int],
    notes: str = "",
    source: str = "manual",
    note_tag: str = "",
) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("PRAGMA foreign_keys = ON")
        _upsert_assessment_entry_in_connection(conn, assessment_id, entry_date, items, notes, source, note_tag)


def _upsert_phq9_entry_in_connection(
    conn: sqlite3.Connection,
    entry_date: str,
    items: list[int],
    notes: str,
    source: str,
    note_tag: str,
) -> None:
    total = calculate_questionnaire_total(QUESTIONNAIRES["phq9"], items)
    if total is None:
        raise ValueError("PHQ-9 does not define a total score.")
    severity = severity_for_score(total)
    conn.execute(
        """
        INSERT INTO phq9_entries (
            entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
            total, severity, notes, note_tag, source
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(entry_date) DO UPDATE SET
            item1=excluded.item1,
            item2=excluded.item2,
            item3=excluded.item3,
            item4=excluded.item4,
            item5=excluded.item5,
            item6=excluded.item6,
            item7=excluded.item7,
            item8=excluded.item8,
            item9=excluded.item9,
            total=excluded.total,
            severity=excluded.severity,
            notes=CASE WHEN excluded.notes != '' THEN excluded.notes ELSE phq9_entries.notes END,
            note_tag=CASE WHEN excluded.note_tag != '' THEN excluded.note_tag ELSE phq9_entries.note_tag END,
            source=excluded.source,
            updated_at=CURRENT_TIMESTAMP
        """,
        (entry_date, *items, total, severity, notes or "", note_tag or "", source),
    )


def upsert_questionnaire_entries(
    entry_date: str,
    responses_by_questionnaire: dict[str, list[int]],
    notes: str = "",
    source: str = "manual",
    note_tag: str = "",
) -> None:
    if not responses_by_questionnaire:
        raise ValueError("Select at least one questionnaire to record.")
    unknown = set(responses_by_questionnaire) - set(QUESTIONNAIRES)
    if unknown:
        raise ValueError(f"Unknown questionnaire: {sorted(unknown)[0]}")
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("PRAGMA foreign_keys = ON")
        for questionnaire_id in QUESTIONNAIRE_ORDER:
            if questionnaire_id not in responses_by_questionnaire:
                continue
            items = responses_by_questionnaire[questionnaire_id]
            if questionnaire_id == "phq9":
                _upsert_phq9_entry_in_connection(conn, entry_date, items, notes, source, note_tag)
            _upsert_assessment_entry_in_connection(
                conn,
                questionnaire_id,
                entry_date,
                items,
                notes,
                source,
                note_tag,
            )


def upsert_entry(entry_date: str, items: list[int], notes: str = "", source: str = "manual", note_tag: str = "") -> None:
    upsert_questionnaire_entries(
        entry_date,
        {"phq9": items},
        notes=notes,
        source=source,
        note_tag=note_tag,
    )


def fetch_assessment_entries(
    assessment_id: str,
    start: str | None = None,
    end: str | None = None,
    limit: int | None = None,
) -> list[AssessmentEntryRow]:
    """Prefer verified normalized records, with a fallback for unmigrated databases."""
    definition = ASSESSMENTS[assessment_id]
    with closing(sqlite3.connect(DB_PATH)) as conn:
        has_normalized_storage = conn.execute(
            """
            SELECT COUNT(*)
            FROM sqlite_master
            WHERE type = 'table' AND name IN ('questionnaire_submissions', 'questionnaire_responses')
            """
        ).fetchone()[0] == 2
        if has_normalized_storage:
            normalized_rows = _fetch_normalized_assessment_entries(conn, assessment_id, start, end, None)
            legacy_rows = _fetch_legacy_assessment_entries(assessment_id, definition, start, end, None)
            normalized_dates = {row.entry_date for row in normalized_rows}
            rows = sorted(
                [*normalized_rows, *(row for row in legacy_rows if row.entry_date not in normalized_dates)],
                key=lambda row: row.entry_date,
            )
            return rows[: int(limit)] if limit else rows
    return _fetch_legacy_assessment_entries(assessment_id, definition, start, end, limit)


def _fetch_normalized_assessment_entries(
    conn: sqlite3.Connection,
    questionnaire_id: str,
    start: str | None,
    end: str | None,
    limit: int | None,
) -> list[AssessmentEntryRow]:
    clauses = ["submission.questionnaire_id = ?"]
    params: list[object] = [questionnaire_id]
    if start:
        clauses.append("submission.entry_date >= ?")
        params.append(start)
    if end:
        clauses.append("submission.entry_date <= ?")
        params.append(end)
    limit_sql = f"LIMIT {int(limit)}" if limit else ""
    submissions = conn.execute(
        f"""
        SELECT
            submission.id, legacy.id, submission.questionnaire_id, submission.definition_version,
            submission.entry_date, submission.total_score, submission.severity,
            COALESCE(submission.notes, ''), COALESCE(submission.note_tag, ''),
            snapshot.definition_json, snapshot.definition_sha256
        FROM questionnaire_submissions AS submission
        JOIN questionnaire_definition_snapshots AS snapshot
          ON snapshot.questionnaire_id = submission.questionnaire_id
         AND snapshot.definition_version = submission.definition_version
        LEFT JOIN assessment_entries AS legacy
          ON legacy.assessment_id = submission.questionnaire_id
         AND legacy.entry_date = submission.entry_date
        WHERE {" AND ".join(clauses)}
        ORDER BY submission.entry_date ASC
        {limit_sql}
        """,
        params,
    ).fetchall()

    results = []
    for submission in submissions:
        submission_id, legacy_id = submission[0], submission[1]
        if legacy_id is None:
            raise RuntimeError(
                f"Normalized submission is missing its compatibility row: {questionnaire_id} {submission[4]}"
            )
        payload, recorded_hash = submission[9], submission[10]
        if hashlib.sha256(payload.encode("utf-8")).hexdigest() != recorded_hash:
            raise RuntimeError(
                f"Questionnaire definition snapshot hash mismatch for {questionnaire_id} v{submission[3]}."
            )
        snapshot = deserialize_questionnaire_definition(payload)
        if (snapshot.questionnaire_id, snapshot.definition_version) != (questionnaire_id, submission[3]):
            raise RuntimeError(
                f"Questionnaire definition snapshot identity mismatch for {questionnaire_id} v{submission[3]}."
            )
        response_rows = conn.execute(
            """
            SELECT question_id, response_order, response_value_json, response_score
            FROM questionnaire_responses
            WHERE submission_id = ?
            ORDER BY response_order
            """,
            (submission_id,),
        ).fetchall()
        if len(response_rows) != snapshot.item_count:
            raise RuntimeError(f"Normalized response count does not reconcile: {questionnaire_id} {submission[4]}")

        responses = []
        for response_order, (question, response_row) in enumerate(zip(snapshot.items, response_rows), start=1):
            if (response_row[0], response_row[1]) != (question.question_id, response_order):
                raise RuntimeError(f"Normalized response identity does not reconcile: {questionnaire_id} {submission[4]}")
            try:
                response = json.loads(response_row[2])
            except json.JSONDecodeError as exc:
                raise RuntimeError(
                    f"Normalized response JSON is invalid: {questionnaire_id} {submission[4]}"
                ) from exc
            options = [option for option in question.options if option.value == response]
            if len(options) != 1 or options[0].score != response_row[3]:
                raise RuntimeError(f"Normalized response does not match its definition: {question.question_id}")
            responses.append(response)

        calculated_total = calculate_questionnaire_total(snapshot, responses)
        if calculated_total != submission[5]:
            raise RuntimeError(f"Normalized total does not reconcile: {questionnaire_id} {submission[4]}")
        if snapshot.interpretation_policy == "validated_builtin":
            expected_severity = severity_for_assessment(questionnaire_id, submission[5])
            if expected_severity != submission[6]:
                raise RuntimeError(f"Normalized severity does not reconcile: {questionnaire_id} {submission[4]}")
        results.append(
            AssessmentEntryRow(
                id=legacy_id,
                assessment_id=questionnaire_id,
                entry_date=submission[4],
                items=responses,
                total=submission[5],
                severity=submission[6],
                notes=submission[7],
                note_tag=submission[8],
                definition_version=submission[3],
            )
        )
    return results


def _fetch_legacy_assessment_entries(
    assessment_id: str,
    definition: QuestionnaireDefinition,
    start: str | None,
    end: str | None,
    limit: int | None,
) -> list[AssessmentEntryRow]:
    clauses = ["assessment_id = ?"]
    params = [assessment_id]
    if start:
        clauses.append("entry_date >= ?")
        params.append(start)
    if end:
        clauses.append("entry_date <= ?")
        params.append(end)
    limit_sql = f"LIMIT {int(limit)}" if limit else ""
    sql = f"""
        SELECT id, assessment_id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
               total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
        FROM assessment_entries
        WHERE {" AND ".join(clauses)}
        ORDER BY entry_date ASC
        {limit_sql}
    """
    with closing(sqlite3.connect(DB_PATH)) as conn:
        rows = conn.execute(sql, params).fetchall()
    return [
        AssessmentEntryRow(
            id=row[0],
            assessment_id=row[1],
            entry_date=row[2],
            items=[int(v) for v in row[3 : 3 + definition.item_count]],
            total=int(row[12]),
            severity=row[13],
            notes=row[14] or "",
            note_tag=row[15] or "",
        )
        for row in rows
    ]


def fetch_entries(start: str | None = None, end: str | None = None, limit: int | None = None) -> list[EntryRow]:
    clauses = []
    params = []
    if start:
        clauses.append("entry_date >= ?")
        params.append(start)
    if end:
        clauses.append("entry_date <= ?")
        params.append(end)
    where = "WHERE " + " AND ".join(clauses) if clauses else ""
    limit_sql = f"LIMIT {int(limit)}" if limit else ""
    sql = f"""
        SELECT id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
               total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
        FROM phq9_entries
        {where}
        ORDER BY entry_date ASC
        {limit_sql}
    """
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        rows = conn.execute(sql, params).fetchall()
    return [
        EntryRow(
            id=row[0],
            entry_date=row[1],
            items=[int(v) for v in row[2:11]],
            total=int(row[11]),
            severity=row[12],
            notes=row[13] or "",
            note_tag=row[14] or "",
        )
        for row in rows
    ]


def fetch_recent_entries(days: int = 14) -> list[EntryRow]:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        rows = conn.execute(
            """
            SELECT id, entry_date, item1, item2, item3, item4, item5, item6, item7, item8, item9,
                   total, severity, COALESCE(notes, ''), COALESCE(note_tag, '')
            FROM phq9_entries
            ORDER BY entry_date DESC
            LIMIT ?
            """,
            (days,),
        ).fetchall()
    rows.reverse()
    return [
        EntryRow(row[0], row[1], [int(v) for v in row[2:11]], int(row[11]), row[12], row[13] or "", row[14] or "")
        for row in rows
    ]


def fetch_events(start: str | None = None, end: str | None = None) -> list[tuple[int, str, str, str]]:
    clauses = []
    params = []
    if start:
        clauses.append("event_date >= ?")
        params.append(start)
    if end:
        clauses.append("event_date <= ?")
        params.append(end)
    where = "WHERE " + " AND ".join(clauses) if clauses else ""
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        return conn.execute(
            f"""
            SELECT id, event_date, event_type, COALESCE(description, '')
            FROM treatment_events
            {where}
            ORDER BY event_date ASC, id ASC
            """,
            params,
        ).fetchall()


def fetch_day_data(entry_date: str) -> dict[str, object]:
    """Return all editable records for one calendar date."""
    assessments = {
        assessment_id: rows[0]
        for assessment_id in ASSESSMENT_ORDER
        if (rows := fetch_assessment_entries(assessment_id, entry_date, entry_date))
    }
    notes = ""
    note_tag = ""
    for assessment_id in ASSESSMENT_ORDER:
        row = assessments.get(assessment_id)
        if row and (row.notes or row.note_tag):
            notes = row.notes
            note_tag = row.note_tag
            break
    return {
        "assessments": assessments,
        "notes": notes,
        "note_tag": note_tag,
        "events": fetch_events(entry_date, entry_date),
    }


def update_assessment_entry(entry_id: int, assessment_id: str, items: list[int]) -> None:
    """Update an assessment in place so its stable record ID is preserved."""
    definition = ASSESSMENTS[assessment_id]
    if len(items) != definition.item_count or any(score < 0 or score > 3 for score in items):
        raise ValueError(f"{definition.display_name} requires {definition.item_count} item scores from 0 to 3.")
    padded_items = [*items, *([None] * (9 - len(items)))]
    total = calculate_questionnaire_total(definition, items)
    if total is None:
        raise ValueError(f"{definition.display_name} does not define a total score.")
    severity = severity_for_assessment(assessment_id, total)
    assignments = ", ".join([*(f"item{i} = ?" for i in range(1, 10)), "total = ?", "severity = ?", "updated_at = CURRENT_TIMESTAMP"])
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("PRAGMA foreign_keys = ON")
        cursor = conn.execute(
            f"UPDATE assessment_entries SET {assignments} WHERE id = ? AND assessment_id = ?",
            (*padded_items, total, severity, entry_id, assessment_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("The selected assessment no longer exists.")
        if assessment_id == "phq9":
            entry_date_row = conn.execute("SELECT entry_date FROM assessment_entries WHERE id = ?", (entry_id,)).fetchone()
            if entry_date_row:
                legacy_assignments = ", ".join([*(f"item{i} = ?" for i in range(1, 10)), "total = ?", "severity = ?", "updated_at = CURRENT_TIMESTAMP"])
                conn.execute(
                    f"UPDATE phq9_entries SET {legacy_assignments} WHERE entry_date = ?",
                    (*items, total, severity, entry_date_row[0]),
                )
        entry_date_row = conn.execute(
            "SELECT entry_date FROM assessment_entries WHERE id = ? AND assessment_id = ?",
            (entry_id, assessment_id),
        ).fetchone()
        if entry_date_row:
            _write_normalized_submission(conn, definition, entry_date_row[0], items)


def update_daily_note(entry_date: str, notes: str, note_tag: str = "") -> None:
    """Synchronize the single day-level note across assessment and legacy rows."""
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute(
            "UPDATE assessment_entries SET notes = ?, note_tag = ?, updated_at = CURRENT_TIMESTAMP WHERE entry_date = ?",
            (notes, note_tag, entry_date),
        )
        conn.execute(
            "UPDATE phq9_entries SET notes = ?, note_tag = ?, updated_at = CURRENT_TIMESTAMP WHERE entry_date = ?",
            (notes, note_tag, entry_date),
        )
        conn.execute(
            "UPDATE questionnaire_submissions SET notes = ?, note_tag = ?, updated_at = CURRENT_TIMESTAMP WHERE entry_date = ?",
            (notes, note_tag, entry_date),
        )


def delete_daily_note(entry_date: str) -> None:
    update_daily_note(entry_date, "", "")


def delete_assessment_entry(entry_id: int, assessment_id: str) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("PRAGMA foreign_keys = ON")
        row = conn.execute(
            "SELECT entry_date FROM assessment_entries WHERE id = ? AND assessment_id = ?",
            (entry_id, assessment_id),
        ).fetchone()
        if not row:
            return
        conn.execute(
            "DELETE FROM questionnaire_submissions WHERE questionnaire_id = ? AND entry_date = ?",
            (assessment_id, row[0]),
        )
        conn.execute("DELETE FROM assessment_entries WHERE id = ? AND assessment_id = ?", (entry_id, assessment_id))
        if assessment_id == "phq9":
            conn.execute("DELETE FROM phq9_entries WHERE entry_date = ?", (row[0],))


def update_event(event_id: int, event_date: str, event_type: str, description: str = "") -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        cursor = conn.execute(
            "UPDATE treatment_events SET event_date = ?, event_type = ?, description = ? WHERE id = ?",
            (event_date, event_type, description, event_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("The selected treatment event no longer exists.")


def delete_event(event_id: int) -> None:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        conn.execute("DELETE FROM treatment_events WHERE id = ?", (event_id,))


def import_spreadsheet(path: str) -> int:
    if load_workbook is None:
        raise RuntimeError("The openpyxl package is required for Excel import.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"date", *(f"item {i}" for i in range(1, 10))}
    target = None
    headers = None
    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            continue
        names = {str(c).strip().lower() for c in first_row if c is not None}
        if required.issubset(names):
            target = sheet
            headers = [str(c).strip() if c is not None else "" for c in first_row]
            break
    if target is None or headers is None:
        raise ValueError("No sheet was found with Date and Item 1 through Item 9 columns.")

    col_index = {name.lower(): idx for idx, name in enumerate(headers)}
    count = 0
    blank_streak = 0
    for row in target.iter_rows(min_row=2, values_only=True):
        raw_date = row[col_index["date"]] if col_index["date"] < len(row) else None
        entry_date = parse_date(raw_date)
        if not entry_date:
            blank_streak += 1
            if blank_streak >= 50:
                break
            continue
        blank_streak = 0
        items = []
        for i in range(1, 10):
            idx = col_index[f"item {i}"]
            value = row[idx] if idx < len(row) else None
            if value is None or str(value).strip().lower() == "nan":
                items = []
                break
            score = int(value)
            if score < 0 or score > 3:
                raise ValueError(f"Item {i} on {entry_date} has value {score}; expected 0-3.")
            items.append(score)
        if len(items) != 9:
            continue
        notes = ""
        if "notes" in col_index and col_index["notes"] < len(row) and row[col_index["notes"]]:
            notes = str(row[col_index["notes"]])
        note_tag = ""
        if "note tag" in col_index and col_index["note tag"] < len(row) and row[col_index["note tag"]]:
            note_tag = str(row[col_index["note tag"]])
        upsert_entry(
            entry_date,
            items,
            notes=notes if notes.lower() != "nan" else "",
            source="spreadsheet",
            note_tag=note_tag if note_tag.lower() != "nan" else "",
        )
        for event_name in ("Ketamine", "Therapy"):
            key = event_name.lower()
            if key in col_index:
                value = row[col_index[key]] if col_index[key] < len(row) else None
                if bool(value) and str(value).lower() not in ("nan", "false", "0"):
                    add_event(entry_date, event_name, f"{event_name} marked in imported spreadsheet", dedupe=True)
        count += 1
    workbook.close()
    return count


def add_event(event_date: str, event_type: str, description: str = "", dedupe: bool = True) -> int:
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        if dedupe:
            found = conn.execute(
                """
                SELECT id FROM treatment_events
                WHERE event_date = ? AND event_type = ? AND description = ?
                """,
                (event_date, event_type, description),
            ).fetchone()
            if found:
                return int(found[0])
        cursor = conn.execute(
            "INSERT INTO treatment_events (event_date, event_type, description) VALUES (?, ?, ?)",
            (event_date, event_type, description),
        )
        conn.commit()
        return int(cursor.lastrowid)


def upsert_daily_event(event_date: str, event_type: str, description: str = "") -> int:
    """Update the day's existing event tag or insert it once.

    The explicit event-management workflow can still add another legitimate event.
    """
    with closing(sqlite3.connect(DB_PATH)) as conn, conn:
        found = conn.execute(
            "SELECT id FROM treatment_events WHERE event_date = ? AND event_type = ? ORDER BY id LIMIT 1",
            (event_date, event_type),
        ).fetchone()
        if found:
            if description:
                conn.execute("UPDATE treatment_events SET description = ? WHERE id = ?", (description, found[0]))
            return int(found[0])
        cursor = conn.execute(
            "INSERT INTO treatment_events (event_date, event_type, description) VALUES (?, ?, ?)",
            (event_date, event_type, description),
        )
        return int(cursor.lastrowid)


def _analysis_workbook_data(questionnaire_ids=None) -> dict[str, list[dict[str, object]]]:
    """Build normalized, analysis-ready records without changing the database schema."""
    selected_ids = normalize_questionnaire_selection(questionnaire_ids)
    with closing(sqlite3.connect(DB_PATH)) as conn:
        conn.row_factory = sqlite3.Row
        compatibility_rows = conn.execute(
            """
            SELECT id, assessment_id, entry_date, item1, item2, item3, item4, item5,
                   item6, item7, item8, item9, total, severity, COALESCE(notes, '') AS notes,
                   COALESCE(note_tag, '') AS note_tag, source, created_at, updated_at
            FROM assessment_entries
            ORDER BY entry_date, assessment_id, id
            """
        ).fetchall()
        compatibility_rows = [row for row in compatibility_rows if row["assessment_id"] in selected_ids]
        events = conn.execute(
            """
            SELECT id, event_date, event_type, COALESCE(description, '') AS description, created_at
            FROM treatment_events
            ORDER BY event_date, id
            """
        ).fetchall()

    compatibility_by_key = {
        (row["assessment_id"], row["entry_date"]): row for row in compatibility_rows
    }
    assessments = [
        entry
        for questionnaire_id in selected_ids
        for entry in fetch_assessment_entries(questionnaire_id)
    ]
    assessments.sort(key=lambda entry: (entry.entry_date, entry.assessment_id, entry.id))

    daily_assessments: list[dict[str, object]] = []
    item_responses: list[dict[str, object]] = []
    notes_by_date: dict[str, dict[str, object]] = {}
    assessment_ids_by_date: dict[str, list[str]] = {}
    assessment_summary_by_date: dict[str, dict[str, object]] = {}

    definition_versions: dict[str, set[int]] = {questionnaire_id: set() for questionnaire_id in selected_ids}
    for row in assessments:
        definition_version = row.definition_version or QUESTIONNAIRES[row.assessment_id].definition_version
        definition = (
            QUESTIONNAIRES[row.assessment_id]
            if definition_version == QUESTIONNAIRES[row.assessment_id].definition_version
            else load_questionnaire_definition_snapshot(row.assessment_id, definition_version)
        )
        definition_versions[row.assessment_id].add(definition_version)
        compatibility = compatibility_by_key[(row.assessment_id, row.entry_date)]
        assessment_record_id = f"assessment:{row.id}"
        daily_record_id = f"day:{row.entry_date}"
        daily_assessments.append(
            {
                "assessment_record_id": assessment_record_id,
                "assessment_entry_id": row.id,
                "daily_record_id": daily_record_id,
                "entry_date": row.entry_date,
                "assessment_id": row.assessment_id,
                "assessment_name": definition.display_name,
                "definition_version": definition_version,
                "interpretation_policy": definition.interpretation_policy,
                "daily_severity_score": row.total,
                "severity_category": row.severity,
                "source": compatibility["source"],
                "created_at": compatibility["created_at"],
                "updated_at": compatibility["updated_at"],
            }
        )
        assessment_ids_by_date.setdefault(row.entry_date, []).append(assessment_record_id)
        assessment_summary_by_date.setdefault(row.entry_date, {})[f"{row.assessment_id}_daily_severity_score"] = row.total
        assessment_summary_by_date[row.entry_date][f"{row.assessment_id}_severity_category"] = row.severity
        for item_number, (question, response_value) in enumerate(zip(definition.items, row.items), start=1):
            response_option = next(option for option in question.options if option.value == response_value)
            item_responses.append(
                {
                    "item_response_record_id": f"{assessment_record_id}:item:{item_number}",
                    "assessment_record_id": assessment_record_id,
                    "assessment_entry_id": row.id,
                    "daily_record_id": daily_record_id,
                    "entry_date": row.entry_date,
                    "assessment_id": row.assessment_id,
                    "definition_version": definition_version,
                    "question_id": question.question_id,
                    "item_number": item_number,
                    "item_label": question.report_label or question.prompt,
                    "response_value": response_value,
                    "response_label": user_facing_response_label(definition, response_option),
                    "response_score": response_option.score if response_option.score is not None else "",
                    "symptom_present": (
                        bool(response_option.score > 0) if response_option.score is not None else ""
                    ),
                }
            )
        if (row.notes or row.note_tag) and row.entry_date not in notes_by_date:
            notes_by_date[row.entry_date] = {
                "note_record_id": f"note:{row.entry_date}",
                "daily_record_id": daily_record_id,
                "entry_date": row.entry_date,
                "note_tag": row.note_tag,
                "note_text": row.notes,
                "created_at": compatibility["created_at"],
                "updated_at": compatibility["updated_at"],
            }

    treatment_events = []
    event_ids_by_date: dict[str, list[str]] = {}
    event_types_by_date: dict[str, list[str]] = {}
    for row in events:
        event_record_id = f"event:{row['id']}"
        normalized_type = normalize_event_type(row["event_type"])
        treatment_events.append(
            {
                "treatment_event_record_id": event_record_id,
                "treatment_event_id": row["id"],
                "daily_record_id": f"day:{row['event_date']}",
                "event_date": row["event_date"],
                "event_type": row["event_type"],
                "normalized_event_type": normalized_type,
                "description": row["description"],
                "created_at": row["created_at"],
            }
        )
        event_ids_by_date.setdefault(row["event_date"], []).append(event_record_id)
        event_types_by_date.setdefault(row["event_date"], []).append(normalized_type)

    ketamine_anchors: list[tuple[str, str]] = []
    for row in treatment_events:
        if row["normalized_event_type"] == "Ketamine" and not any(anchor[0] == row["event_date"] for anchor in ketamine_anchors):
            ketamine_anchors.append((str(row["event_date"]), str(row["treatment_event_record_id"])))
    all_dates = sorted(set(assessment_ids_by_date) | set(notes_by_date) | set(event_ids_by_date))
    last_recorded_date = all_dates[-1] if all_dates else None
    treatment_cycle_rows = []
    for index, (anchor_date, anchor_event_id) in enumerate(ketamine_anchors):
        next_anchor_date = ketamine_anchors[index + 1][0] if index + 1 < len(ketamine_anchors) else None
        cycle_end = (
            (datetime.fromisoformat(next_anchor_date).date() - timedelta(days=1)).isoformat()
            if next_anchor_date
            else last_recorded_date or anchor_date
        )
        treatment_cycle_rows.append(
            {
                "treatment_cycle_record_id": f"cycle:{anchor_date}",
                "anchor_treatment_event_record_id": anchor_event_id,
                "cycle_number": index + 1,
                "cycle_start_date": anchor_date,
                "cycle_end_date": cycle_end,
                "is_current_cycle": index == len(ketamine_anchors) - 1,
            }
        )

    daily_summary = []
    for entry_date in all_dates:
        summary = assessment_summary_by_date.get(entry_date, {})
        event_types = event_types_by_date.get(entry_date, [])
        daily_summary.append(
            {
                "daily_record_id": f"day:{entry_date}",
                "entry_date": entry_date,
                "assessment_record_ids": "|".join(assessment_ids_by_date.get(entry_date, [])),
                "phq9_daily_severity_score": summary.get("phq9_daily_severity_score", ""),
                "phq9_severity_category": summary.get("phq9_severity_category", ""),
                "gad7_daily_severity_score": summary.get("gad7_daily_severity_score", ""),
                "gad7_severity_category": summary.get("gad7_severity_category", ""),
                "note_record_id": notes_by_date.get(entry_date, {}).get("note_record_id", ""),
                "treatment_event_record_ids": "|".join(event_ids_by_date.get(entry_date, [])),
                "treatment_event_count": len(event_ids_by_date.get(entry_date, [])),
                "ketamine_recorded": "Ketamine" in event_types,
                "therapy_recorded": "Therapy" in event_types,
                "medication_change_recorded": any(event_type.startswith("Medication") for event_type in event_types),
            }
        )

    metadata = [
        {"metadata_key": "workbook_schema_version", "metadata_value": ANALYSIS_WORKBOOK_SCHEMA_VERSION, "description": "Version of this normalized export layout."},
        {"metadata_key": "generated_at", "metadata_value": datetime.now().astimezone().isoformat(timespec="seconds"), "description": "Local generation timestamp in ISO 8601 format."},
        {"metadata_key": "application", "metadata_value": APPLICATION_NAME, "description": "Application that generated the workbook."},
        {"metadata_key": "questionnaires_included", "metadata_value": "|".join(selected_ids), "description": "Questionnaires selected for this workbook."},
        {
            "metadata_key": "questionnaire_definition_versions",
            "metadata_value": "|".join(
                f"{questionnaire_id}:v{version}"
                for questionnaire_id in selected_ids
                for version in sorted(definition_versions[questionnaire_id])
            ),
            "description": "Exact stored questionnaire definition versions represented in this workbook.",
        },
        {"metadata_key": "safety_message", "metadata_value": UNIVERSAL_SAFETY_MESSAGE, "description": "Universal owner-approved safety message."},
        {"metadata_key": "non_diagnostic_notice", "metadata_value": NON_DIAGNOSTIC_OUTPUT_NOTICE, "description": "Interpretation boundary for all workbook content."},
        {"metadata_key": "date_format", "metadata_value": "YYYY-MM-DD", "description": "Calendar-date format used in all date fields."},
        {"metadata_key": "daily_relationship", "metadata_value": "daily_record_id", "description": "Join records across worksheets by the deterministic day:YYYY-MM-DD identifier."},
        {"metadata_key": "assessment_relationship", "metadata_value": "assessment_record_id", "description": "Join Item Responses to Daily Assessments by assessment_record_id."},
        {"metadata_key": "event_relationship", "metadata_value": "treatment_event_record_id", "description": "Treatment Events retain the stable SQLite event ID as event:<id>."},
        {"metadata_key": "cycle_relationship", "metadata_value": "anchor_treatment_event_record_id", "description": "Treatment Cycles reference the first recorded ketamine event on each anchor date."},
        {"metadata_key": "daily_severity_score", "metadata_value": "sum of item responses on one check-in", "description": "Measures recorded symptom severity for a single assessment date."},
        {"metadata_key": "14_day_frequency_score", "metadata_value": "derived by calendar window", "description": "The 14-Day Item Profile contains one current-window derived record per assessment item."},
        {"metadata_key": "missing_checkins", "metadata_value": "missing information", "description": "A day without a check-in is not evidence that symptoms were absent."},
        {"metadata_key": "daily_summary_authority", "metadata_value": "derived", "description": "Daily Summary is a convenience view; normalized worksheets remain authoritative."},
        {"metadata_key": "privacy_notice", "metadata_value": "local sensitive data", "description": "This workbook may contain PHI. Store and share it deliberately."},
    ]

    profile_end = max((row.entry_date for row in assessments), default=date.today().isoformat())
    return {
        "Daily Assessments": daily_assessments,
        "Item Responses": item_responses,
        "Notes": list(notes_by_date.values()),
        "Treatment Events": treatment_events,
        "Treatment Cycles": treatment_cycle_rows,
        "Metadata": metadata,
        "Daily Summary": daily_summary,
        "14-Day Item Profile": build_14_day_item_profile(profile_end, selected_ids),
    }


def export_analysis_workbook(path: str, questionnaire_ids=None) -> None:
    """Write the separate normalized XLSX export used for external analysis."""
    if not path.lower().endswith(".xlsx"):
        raise ValueError("The analysis-ready export must be saved as an .xlsx workbook.")
    if pd is None:
        raise RuntimeError("Analysis-ready Excel export requires pandas/openpyxl.")
    sheet_columns = {
        "Daily Assessments": ["assessment_record_id", "assessment_entry_id", "daily_record_id", "entry_date", "assessment_id", "assessment_name", "daily_severity_score", "severity_category", "source", "created_at", "updated_at", "definition_version", "interpretation_policy"],
        "Item Responses": ["item_response_record_id", "assessment_record_id", "assessment_entry_id", "daily_record_id", "entry_date", "assessment_id", "item_number", "item_label", "response_value", "symptom_present", "definition_version", "question_id", "response_label", "response_score"],
        "Notes": ["note_record_id", "daily_record_id", "entry_date", "note_tag", "note_text", "created_at", "updated_at"],
        "Treatment Events": ["treatment_event_record_id", "treatment_event_id", "daily_record_id", "event_date", "event_type", "normalized_event_type", "description", "created_at"],
        "Treatment Cycles": ["treatment_cycle_record_id", "anchor_treatment_event_record_id", "cycle_number", "cycle_start_date", "cycle_end_date", "is_current_cycle"],
        "Metadata": ["metadata_key", "metadata_value", "description"],
        "Daily Summary": ["daily_record_id", "entry_date", "assessment_record_ids", "phq9_daily_severity_score", "phq9_severity_category", "gad7_daily_severity_score", "gad7_severity_category", "note_record_id", "treatment_event_record_ids", "treatment_event_count", "ketamine_recorded", "therapy_recorded", "medication_change_recorded"],
        "14-Day Item Profile": ["profile_record_id", "window_start", "window_end", "assessment_id", "assessment_name", "item_number", "item_label", "symptom_present_days", "recorded_day_coverage", "calendar_days", "frequency_score", "definition_version", "definition_versions", "analytical_series_id"],
    }
    workbook_data = _analysis_workbook_data(questionnaire_ids)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, columns in sheet_columns.items():
            frame = pd.DataFrame(workbook_data[sheet_name], columns=columns)
            frame.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="172033")
                cell.fill = PatternFill(fill_type="solid", fgColor="E8EEF5")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for column_index, column_name in enumerate(columns, start=1):
                values = [column_name, *("" if value is None else str(value) for value in frame[column_name].tolist())]
                width = min(max(len(value) for value in values) + 2, 48)
                column_letter = worksheet.cell(row=1, column=column_index).column_letter
                worksheet.column_dimensions[column_letter].width = width
                for cell in worksheet[column_letter][1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)


def python_supports_modules(python_path: Path, modules: tuple[str, ...]) -> bool:
    if not python_path.exists():
        return False
    imports = "; ".join(f"import {module}" for module in modules)
    try:
        result = subprocess.run(
            [str(python_path), "-c", imports],
            capture_output=True,
            text=True,
            timeout=15,
        )
        return result.returncode == 0
    except (OSError, subprocess.SubprocessError):
        return False


def find_helper_python(required_modules: tuple[str, ...]) -> Path | None:
    """Find a configured or project-local Python that has the requested features."""
    candidates = []
    if BUNDLED_PYTHON is not None:
        candidates.append(BUNDLED_PYTHON)
    candidates.extend(
        [
            PROJECT_ROOT / ".venv" / "Scripts" / "python.exe",
            PROJECT_ROOT / "venv" / "Scripts" / "python.exe",
        ]
    )
    current = Path(sys.executable).resolve()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            continue
        if resolved == current:
            continue
        if python_supports_modules(resolved, required_modules):
            return resolved
    return None


def run_bundled_cli(args: list[str], required_modules: tuple[str, ...]) -> None:
    helper_python = find_helper_python(required_modules)
    if helper_python is None:
        names = ", ".join(required_modules)
        raise RuntimeError(
            f"This action needs Python packages that are not available: {names}. "
            "Install the project requirements in .venv, venv, or the active Python environment, "
            "or set PHQ9_TRACKER_BUNDLED_PYTHON to a compatible python.exe."
        )
    command = [str(helper_python), str(Path(__file__).resolve()), *args]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "Bundled command failed.")


def add_pdf_table(
    story,
    title: str,
    rows: list[list[str]],
    repeat_header: bool = True,
    highlight_flag_col: int | None = None,
    col_widths: list[float] | None = None,
    wrap_columns: set[int] | None = None,
) -> None:
    styles = getSampleStyleSheet()
    story.append(Paragraph(title, styles["Heading2"]))
    if wrap_columns:
        wrapped = []
        for row_idx, row in enumerate(rows):
            wrapped_row = []
            for col_idx, value in enumerate(row):
                if row_idx > 0 and col_idx in wrap_columns:
                    safe_value = escape(str(value)).replace("\n", "<br/>")
                    wrapped_row.append(Paragraph(safe_value, styles["BodyText"]))
                else:
                    wrapped_row.append(str(value))
            wrapped.append(wrapped_row)
        rows = wrapped
    table = Table(rows, repeatRows=1 if repeat_header and len(rows) > 1 else 0, colWidths=col_widths)
    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#172033")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]
    if highlight_flag_col is not None:
        for row_idx, row in enumerate(rows[1:], start=1):
            if highlight_flag_col < len(row) and str(row[highlight_flag_col]).strip():
                style_commands.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#FFF7ED")))
                style_commands.append(("TEXTCOLOR", (highlight_flag_col, row_idx), (highlight_flag_col, row_idx), colors.HexColor("#9A3412")))
    table.setStyle(TableStyle(style_commands)
    )
    story.append(table)
    story.append(Spacer(1, 0.14 * inch))


def normalize_event_type(event_type: str) -> str:
    text = (event_type or "").strip().lower()
    if "ketamine" in text:
        return "Ketamine"
    if "physical therapy" in text:
        return "Physical Therapy"
    if "therapy" in text:
        return "Therapy"
    if "dose increase" in text:
        return "Medication Dose Increase"
    if "dose decrease" in text:
        return "Medication Dose Decrease"
    if "start" in text:
        return "Medication Start"
    if "stop" in text or "discontinued" in text:
        return "Medication Stop"
    return event_type or "Treatment event"


def draw_line_chart(
    path: str,
    entries: list[EntryRow],
    series: list[tuple[str, list[float | int | None], str]],
    title: str,
    y_max: int,
    events: list[tuple[int, str, str, str]] | None = None,
    show_severity: bool = False,
    note_markers: bool = False,
    width: int = 1100,
    height: int = 520,
) -> None:
    if PILImage is None:
        raise RuntimeError("Chart generation requires Pillow.")
    img = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    title_font = ImageFont.load_default()
    left, right, top, bottom = 70, 35, 58, 72
    plot_w = width - left - right
    plot_h = height - top - bottom
    draw.text((left, 18), title, fill="#172033", font=title_font)
    if not entries:
        draw.text((width // 2 - 60, height // 2), "No entries", fill="#64748B", font=font)
        img.save(path)
        return

    def x_for_idx(idx: int) -> float:
        return left + (idx / max(len(entries) - 1, 1)) * plot_w

    def y_for_value(value: float) -> float:
        return top + plot_h - (value / y_max) * plot_h

    if show_severity:
        bands = [
            (0, 4, "#ECFDF5", "Minimal"),
            (5, 9, "#F0FDF4", "Mild"),
            (10, 14, "#FEFCE8", "Moderate"),
            (15, 19, "#FFF7ED", "Moderately Severe"),
            (20, 27, "#FEF2F2", "Severe"),
        ]
        for low, high, color, label in bands:
            y1 = y_for_value(high)
            y2 = y_for_value(low)
            draw.rectangle([left, y1, left + plot_w, y2], fill=color)
            draw.text((left + 6, y1 + 4), label, fill="#475569", font=font)

    for tick in range(0, y_max + 1, 3 if y_max > 5 else 1):
        y = y_for_value(tick)
        draw.line([(left, y), (left + plot_w, y)], fill="#E2E8F0")
        draw.text((18, y - 6), str(tick), fill="#475569", font=font)
    draw.line([(left, top), (left, top + plot_h), (left + plot_w, top + plot_h)], fill="#94A3B8", width=2)

    event_lookup = {}
    for event in events or []:
        event_lookup.setdefault(event[1], []).append(event)
    for idx, row in enumerate(entries):
        x = x_for_idx(idx)
        for event in event_lookup.get(row.entry_date, []):
            kind = normalize_event_type(event[2])
            color = {
                "Ketamine": "#7C3AED",
                "Therapy": "#0F766E",
                "Medication Start": "#2563EB",
                "Medication Stop": "#DC2626",
                "Medication Dose Increase": "#EA580C",
                "Medication Dose Decrease": "#0891B2",
            }.get(kind, "#64748B")
            draw.line([(x, top), (x, top + plot_h)], fill=color, width=2)
            draw.polygon([(x, top - 8), (x - 5, top), (x + 5, top)], fill=color)
        if note_markers and row.notes.strip():
            y = y_for_value(row.total)
            draw.ellipse([x - 5, y - 5, x + 5, y + 5], outline="#111827", width=2)

    for name, values, color in series:
        points = []
        for idx, value in enumerate(values):
            if value is None:
                if len(points) > 1:
                    draw.line(points, fill=color, width=3)
                points = []
                continue
            points.append((x_for_idx(idx), y_for_value(float(value))))
        if len(points) > 1:
            draw.line(points, fill=color, width=3)
        for x, y in points[:: max(1, len(points) // 35)]:
            draw.ellipse([x - 3, y - 3, x + 3, y + 3], fill=color)

    draw.text((left, height - 44), entries[0].entry_date, fill="#475569", font=font)
    draw.text((width - 110, height - 44), entries[-1].entry_date, fill="#475569", font=font)
    legend_x, legend_y = left, height - 24
    for name, _, color in series:
        draw.rectangle([legend_x, legend_y, legend_x + 12, legend_y + 12], fill=color)
        draw.text((legend_x + 18, legend_y - 1), name, fill="#334155", font=font)
        legend_x += 150
    if events:
        draw.text((legend_x, legend_y - 1), "Vertical markers: treatment/medication events", fill="#64748B", font=font)
    if note_markers:
        draw.text((left + 420, 18), "Outlined dots: notes", fill="#111827", font=font)
    img.save(path)


def add_chart(story, image_path: str, caption: str | None = None, width: float = 7.2 * inch) -> None:
    chart_items = []
    if caption:
        chart_items.append(Paragraph(caption, getSampleStyleSheet()["Heading2"]))
    chart_items.append(Image(image_path, width=width, height=width * 0.47))
    chart_items.append(Spacer(1, 0.14 * inch))
    story.append(KeepTogether(chart_items))


def draw_treatment_bar_chart(path: str, title: str, rows: list[list[str]], labels: list[str]) -> None:
    if PILImage is None:
        raise RuntimeError("Chart generation requires Pillow.")
    width, height = 950, 360
    img = PILImage.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.load_default()
    draw.text((50, 18), title, fill="#172033", font=font)
    if not rows:
        draw.text((width // 2 - 80, height // 2), "No events recorded", fill="#64748B", font=font)
        img.save(path)
        return
    values_by_label = []
    for col_idx in range(2, 2 + len(labels)):
        values = []
        for row in rows:
            try:
                values.append(float(row[col_idx]))
            except (ValueError, IndexError):
                pass
        values_by_label.append(sum(values) / len(values) if values else None)
    left, top, bottom = 85, 68, 58
    plot_h = height - top - bottom
    plot_w = width - left - 45
    y_max = max([27, *[v for v in values_by_label if v is not None]])
    for tick in range(0, 28, 3):
        y = top + plot_h - (tick / y_max) * plot_h
        draw.line([(left, y), (left + plot_w, y)], fill="#E2E8F0")
        draw.text((35, y - 6), str(tick), fill="#475569", font=font)
    colors_local = ["#2563EB", "#7C3AED", "#0F766E"]
    bar_gap = plot_w / max(len(labels), 1)
    for idx, label in enumerate(labels):
        value = values_by_label[idx]
        x1 = left + idx * bar_gap + bar_gap * 0.2
        x2 = left + (idx + 1) * bar_gap - bar_gap * 0.2
        if value is not None:
            y = top + plot_h - (value / y_max) * plot_h
            draw.rectangle([x1, y, x2, top + plot_h], fill=colors_local[idx % len(colors_local)])
            draw.text((x1, y - 18), f"{value:.1f}", fill="#172033", font=font)
        draw.text((x1, top + plot_h + 12), label, fill="#334155", font=font)
    draw.line([(left, top), (left, top + plot_h), (left + plot_w, top + plot_h)], fill="#94A3B8", width=2)
    img.save(path)


def on_report_page(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(36, 28, DISCLAIMER)
    canvas.drawRightString(576, 28, f"Page {doc.page}")
    canvas.setFont("Helvetica", 6.5)
    canvas.drawString(36, 17, "Safety: If you feel unsafe or may act on thoughts of self-harm, contact local emergency services or a crisis service now.")
    canvas.drawString(36, 9, "Len does not monitor responses or provide emergency help.")
    canvas.restoreState()


def available_report_date_range(questionnaire_ids=None) -> tuple[str, str]:
    """Return the full assessment-history range used by one-click PDF reports."""
    dates = [
        row.entry_date
        for assessment_id in normalize_questionnaire_selection(questionnaire_ids)
        for row in fetch_assessment_entries(assessment_id)
    ]
    if not dates:
        raise ValueError("No check-ins are available for a clinician report.")
    return min(dates), max(dates)


def generate_report(start: str, end: str, pdf_path: str, questionnaire_ids=None) -> None:
    if colors is None or PILImage is None:
        raise RuntimeError("PDF export requires reportlab and Pillow.")
    start, end = normalize_report_date_range(start, end)
    requested_ids = normalize_questionnaire_selection(questionnaire_ids)
    entries_by_questionnaire = {
        questionnaire_id: fetch_assessment_entries(questionnaire_id, start, end)
        for questionnaire_id in requested_ids
    }
    if not any(entries_by_questionnaire.values()):
        raise ValueError("No entries found in the selected date range.")
    selected_ids = [questionnaire_id for questionnaire_id in requested_ids if entries_by_questionnaire[questionnaire_id]]
    comparison_start = (datetime.fromisoformat(end).date() - timedelta(days=27)).isoformat()
    comparison_entries_by_questionnaire = {
        questionnaire_id: fetch_assessment_entries(questionnaire_id, comparison_start, end)
        for questionnaire_id in selected_ids
    }
    entries = entries_by_questionnaire.get("phq9", [])
    gad_entries = entries_by_questionnaire.get("gad7", [])
    comparison_phq_entries = comparison_entries_by_questionnaire.get("phq9", [])
    comparison_gad_entries = comparison_entries_by_questionnaire.get("gad7", [])
    events = [
        (event_id, event_date, normalize_event_type(event_type), description)
        for event_id, event_date, event_type, description in fetch_events(start, end)
    ]
    phq_definition_groups = _definition_groups_for_entries("phq9", entries) if entries else []
    if entries:
        phq_cycle_definition, phq_cycle_entries, phq_has_multiple_versions = _latest_definition_group(
            "phq9", entries, end
        )
    else:
        phq_cycle_definition = QUESTIONNAIRES["phq9"]
        phq_cycle_entries = []
        phq_has_multiple_versions = False
    cycles = treatment_cycles(phq_cycle_entries, events, end)
    item9_entries = [
        entry
        for definition, definition_entries in phq_definition_groups
        if any(
            question.question_id == "phq9.item9" and "phq9_item9_context" in question.behavior_ids
            for question in definition.items
        )
        for entry in definition_entries
    ]
    highlights = []
    interpreted_profile_ids = []
    for questionnaire_id in selected_ids:
        definition, _, _ = _latest_definition_group(
            questionnaire_id,
            comparison_entries_by_questionnaire[questionnaire_id],
            end,
        )
        if (
            definition.interpretation_policy == "validated_builtin"
            and questionnaire_profile_omission_reason(definition) is None
        ):
            interpreted_profile_ids.append(questionnaire_id)
            highlights.extend(
                symptom_highlights(
                    questionnaire_id,
                    comparison_entries_by_questionnaire[questionnaire_id],
                    end,
                    limit=2,
                )
            )
    item_profile = build_14_day_item_profile(end, selected_ids)
    report_definitions = [
        definition
        for questionnaire_id in selected_ids
        for definition, definition_entries in _definition_groups_for_entries(
            questionnaire_id, entries_by_questionnaire[questionnaire_id]
        )
        if definition_entries
    ]

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    chart_dir = Path(tempfile.mkdtemp(prefix="phq9_report_charts_"))
    chart_colors = ("#2563EB", "#0F766E", "#7C3AED", "#B45309")
    report_charts = []
    for questionnaire_index, questionnaire_id in enumerate(selected_ids):
        for series in build_questionnaire_trend_series(
            questionnaire_id, entries_by_questionnaire[questionnaire_id]
        ):
            chart_path = chart_dir / (
                f"{questionnaire_id}_v{series.definition.definition_version}_recent.png"
            )
            recent_entries = list(series.entries[-90:])
            draw_line_chart(
                str(chart_path),
                recent_entries,
                [
                    (
                        series.label,
                        [row.total for row in recent_entries],
                        chart_colors[questionnaire_index % len(chart_colors)],
                    )
                ],
                f"{series.label} Recorded Score Trend",
                int(series.definition.score_max or 1),
                events=events,
                width=1100,
                height=360,
            )
            report_charts.append((series, chart_path))

    story = [
        Paragraph(f"{APPLICATION_NAME} Clinician Discussion Report", styles["Title"]),
        Paragraph(f"Date range: {start} to {end}", styles["Normal"]),
        Paragraph(DISCLAIMER, styles["BodyText"]),
        Paragraph(NON_DIAGNOSTIC_OUTPUT_NOTICE, styles["BodyText"]),
        Paragraph(UNIVERSAL_SAFETY_MESSAGE, styles["BodyText"]),
        Spacer(1, 0.18 * inch),
        Paragraph("How to Read This Report", styles["Heading1"]),
    ]

    definition_rows = [["Questionnaire", "Definition", "Output policy"]]
    for definition in report_definitions:
        output_policy = (
            "Validated built-in interpretation"
            if definition.interpretation_policy == "validated_builtin"
            else "Descriptive/raw output only; no automatic clinical interpretation"
        )
        definition_rows.append(
            [definition.display_name, f"v{definition.definition_version}", output_policy]
        )
    add_pdf_table(
        story,
        "Questionnaire Definitions",
        definition_rows,
        col_widths=[1.35 * inch, 0.8 * inch, 4.55 * inch],
        wrap_columns={0, 1, 2},
    )
    if report_charts:
        story.append(
            Paragraph(
                "<b>Daily Severity Score:</b> when a questionnaire definition declares a supported total-score "
                "calculation, its recorded item responses are combined for that check-in. Only explicitly "
                "compatible definition versions share a score trend.",
                styles["BodyText"],
            )
        )
    if item_profile:
        story.extend(
            [
                Paragraph(
                    "<b>14-Day Symptom Frequency Score:</b> for definitions that declare the supported profile, each "
                    "item counts calendar days with a scored response above zero. Counts convert to 0 points for 0 "
                    "days, 1 point for 1-6 days, 2 points for 7-11 days, or 3 points for 12-14 days.",
                    styles["BodyText"],
                ),
                Paragraph(
                    "<b>Coverage and missing check-ins:</b> coverage shows recorded daily check-ins in the calendar "
                    "window. A day without a check-in is missing information, not evidence that a symptom was absent.",
                    styles["BodyText"],
                ),
            ]
        )
    story.extend([Spacer(1, 0.12 * inch), Paragraph("Recorded Period Overview", styles["Heading1"])])
    glance_rows = [["Questionnaire", "Definition", "Check-ins", "Most recent result", "Most recent date"]]
    for questionnaire_id in selected_ids:
        for definition, definition_entries in _definition_groups_for_entries(
            questionnaire_id, entries_by_questionnaire[questionnaire_id]
        ):
            if not definition_entries:
                continue
            recent = definition_entries[-1]
            recent_result = str(recent.total)
            if definition.interpretation_policy == "validated_builtin" and recent.severity:
                recent_result = f"{recent.total} ({recent.severity})"
            glance_rows.append(
                [
                    definition.display_name,
                    f"v{definition.definition_version}",
                    str(len(definition_entries)),
                    recent_result,
                    recent.entry_date,
                ]
            )
    add_pdf_table(
        story,
        "Recorded Check-Ins",
        glance_rows,
        col_widths=[1.25 * inch, 0.7 * inch, 0.7 * inch, 2.45 * inch, 1.15 * inch],
        wrap_columns={0, 1, 2, 3, 4},
    )
    if interpreted_profile_ids:
        summary_story = [Paragraph("Overall pattern", styles["Heading2"])]
        summary_story.append(
            Paragraph(
                overall_pattern_summary(
                    comparison_phq_entries,
                    comparison_gad_entries,
                    end,
                    interpreted_profile_ids,
                    comparison_entries_by_questionnaire,
                ),
                styles["BodyText"],
            )
        )
        summary_story.append(Paragraph("Symptom highlights", styles["Heading2"]))
        for highlight in highlights:
            summary_story.append(Paragraph(f"- {highlight}", styles["BodyText"]))
        summary_story.append(Spacer(1, 0.08 * inch))
        summary_story.append(
            Paragraph(
                "Highlights describe recorded check-ins only. A day without a check-in is missing information, not evidence that a symptom was absent.",
                styles["BodyText"],
            )
        )
        story.append(KeepTogether(summary_story))

    story.append(PageBreak())
    story.append(Paragraph("Recorded Symptom Trends", styles["Heading1"]))
    for series, chart_path in report_charts:
        add_chart(
            story,
            str(chart_path),
            f"{series.label} scores across the selected period",
            width=6.1 * inch,
        )
    for definition in report_definitions:
        omission = questionnaire_total_trend_omission_reason(definition)
        if omission:
            story.append(Paragraph(f"{definition.display_name} v{definition.definition_version}: {omission}", styles["BodyText"]))
    q9_context = item9_context(item9_entries, end)
    q9_summary = item9_context_summary(q9_context)
    if q9_summary:
        story.append(Paragraph("PHQ-9 item 9 context", styles["Heading2"]))
        story.append(
            Paragraph(
                q9_summary,
                styles["BodyText"],
            )
        )

    story.append(PageBreak())
    story.append(Paragraph("Current 14-Day Item Profile", styles["Heading1"]))
    if item_profile:
        story.append(
            Paragraph(
                f"Window: {item_profile[0]['window_start']} to {end}. Present days count recorded scored responses "
                "above zero; coverage is recorded check-ins out of 14 calendar days. Missing days remain missing information.",
                styles["BodyText"],
            )
        )
    for assessment_id in selected_ids:
        assessment_profile_rows = [row for row in item_profile if row["assessment_id"] == assessment_id]
        for series_id in sorted({row["analytical_series_id"] for row in assessment_profile_rows}):
            profile_rows = [
                row for row in assessment_profile_rows if row["analytical_series_id"] == series_id
            ]
            if not profile_rows:
                continue
            definition = next(definition for definition in report_definitions
                              if definition.questionnaire_id == assessment_id
                              and definition.definition_version == max(
                                  int(version[1:]) for version in profile_rows[0]["definition_versions"].split("|")
                              ))
            profile_label = (definition.display_name if "|" in profile_rows[0]["definition_versions"]
                             else f"{definition.display_name} v{definition.definition_version}")
            add_pdf_table(
                story,
                profile_label,
                [["Item", "Present days", "Coverage", "Score (0-3)"]]
                + [
                    [
                        f"{row['item_number']}. {row['item_label']}",
                        str(row["symptom_present_days"]),
                        f"{row['recorded_day_coverage']} of {row['calendar_days']}",
                        str(row["frequency_score"]),
                    ]
                    for row in profile_rows
                ],
                col_widths=[4.05 * inch, 0.9 * inch, 0.85 * inch, 0.9 * inch],
                wrap_columns={0, 1, 2, 3},
            )
        for definition in [
            candidate for candidate in report_definitions if candidate.questionnaire_id == assessment_id
        ]:
            omission = questionnaire_profile_omission_reason(definition)
            if omission:
                story.append(
                    Paragraph(
                        f"{definition.display_name} v{definition.definition_version}: {omission}",
                        styles["BodyText"],
                    )
                )

    story.append(PageBreak())
    story.append(Paragraph("Treatment Context", styles["Heading1"]))
    story.append(
        Paragraph(
            "Cycles are bounded by recorded ketamine infusion dates. These observations describe timing and recorded scores; they do not determine whether treatment caused a change.",
            styles["BodyText"],
        )
    )
    if cycles:
        cycle_definition_label = (
            f"PHQ-9 v{phq_cycle_definition.definition_version}"
            if phq_has_multiple_versions
            else "PHQ-9"
        )
        cycle_rows = [["Cycle", "Dates", f"Recorded {cycle_definition_label} pattern"]]
        for cycle in cycles:
            cycle_rows.append([cycle.label, f"{cycle.start_date} to {cycle.end_date}", treatment_cycle_observation(cycle)])
        add_pdf_table(
            story,
            "Current and Previous Recorded Cycles",
            cycle_rows,
            col_widths=[1.0 * inch, 1.65 * inch, 4.05 * inch],
            wrap_columns={0, 1, 2},
        )
        for cycle in cycles:
            cycle_chart = chart_dir / f"{cycle.label.lower().replace(' ', '_')}.png"
            draw_line_chart(
                str(cycle_chart),
                cycle.entries,
                [(cycle_definition_label, [row.total for row in cycle.entries], "#7C3AED")],
                f"{cycle.label}: Recorded {cycle_definition_label} Scores",
                int(phq_cycle_definition.score_max or 1),
                width=1100,
                height=300,
            )
            add_chart(story, str(cycle_chart), width=6.2 * inch)
    else:
        story.append(Paragraph("No ketamine infusion was recorded in the selected period, so cycle views are not available.", styles["BodyText"]))

    timeline = []
    for _, event_date, event_type, description in events:
        timeline.append((event_date, event_type, description or "Recorded event"))
    seen_notes = set()
    for row in [*entries, *gad_entries]:
        note_key = (row.entry_date, row.notes.strip(), row.note_tag.strip())
        if row.notes.strip() and note_key not in seen_notes:
            seen_notes.add(note_key)
            timeline.append((row.entry_date, f"Note{f' - {row.note_tag}' if row.note_tag else ''}", row.notes.strip()))
    timeline = sorted(timeline, key=lambda item: (item[0], item[1]))

    story.append(PageBreak())
    story.append(Paragraph("Journal and Event Context", styles["Heading1"]))
    if timeline:
        add_pdf_table(
            story,
            "Recorded Context",
            [["Date", "Type", "Recorded context"]]
            + [[event_date, event_type, text] for event_date, event_type, text in timeline],
            col_widths=[0.95 * inch, 1.45 * inch, 4.3 * inch],
            wrap_columns={1, 2},
        )
    else:
        story.append(Paragraph("No notes or treatment events were recorded in the selected period.", styles["BodyText"]))
    story.append(Paragraph("Conversation Starters", styles["Heading2"]))
    prompts = [
        "Do the recorded symptom patterns match what you remember about this period?",
        "Were there particular days, events, or treatment dates that would help explain the recorded context?",
        "Which symptom changes would be most useful to discuss or monitor together next?",
    ]
    q9_prompt = item9_conversation_starter(q9_context)
    if q9_prompt:
        prompts.insert(0, q9_prompt)
    for prompt in prompts[:4]:
        story.append(Paragraph(f"- {prompt}", styles["BodyText"]))
    doc.build(story, onFirstPage=on_report_page, onLaterPages=on_report_page)


class LineChart(Canvas):
    def __init__(self, parent, **kwargs):
        kwargs.setdefault("takefocus", True)
        super().__init__(
            parent,
            bg="#FFFFFF",
            highlightthickness=1,
            highlightbackground="#CBD5E1",
            highlightcolor="#2563EB",
            **kwargs,
        )
        self._series_args = None
        self._chart_points: list[ChartPoint] = []
        self._hovered_point_index: int | None = None
        self._locked_point_index: int | None = None
        self.bind("<Configure>", self._redraw_after_resize)
        self.bind("<Motion>", self._show_hover_callout)
        self.bind("<Leave>", self._hide_hover_callout)
        self.bind("<Button-1>", self._toggle_click_callout)
        self.bind("<Left>", lambda event: self._move_keyboard_callout(-1))
        self.bind("<Right>", lambda event: self._move_keyboard_callout(1))
        self.bind("<Return>", self._lock_keyboard_callout)
        self.bind("<Escape>", self._clear_callout)

    def draw_series(self, entries: list[EntryRow], series: list[tuple[str, list[int], str]], y_max: int, title: str) -> None:
        self._series_args = (entries, series, y_max, title)
        self._render_series(entries, series, y_max, title)

    def _redraw_after_resize(self, _event=None) -> None:
        if self._series_args is not None:
            self._render_series(*self._series_args)

    def _render_series(self, entries: list[EntryRow], series: list[tuple[str, list[int], str]], y_max: int, title: str) -> None:
        locked_index = self._locked_point_index
        self.delete("all")
        self._chart_points = []
        self._hovered_point_index = None
        width = responsive_canvas_dimension(self.winfo_width(), int(self["width"]))
        height = responsive_canvas_dimension(self.winfo_height(), int(self["height"]))
        pad_l, pad_r, pad_b = 46, 18, 36
        title_id = self.create_text(
            12,
            10,
            text=title,
            width=max(120, width - 24),
            anchor="nw",
            justify=LEFT,
            fill="#172033",
            font=("Segoe UI", 11, "bold"),
            tags=("chart-title",),
        )
        title_bounds = self.bbox(title_id) or (12, 10, width - 12, 28)
        legend_top, pad_t = chart_vertical_positions(title_bounds[3])
        if not entries or not series:
            self.create_text(width / 2, max(pad_t + 20, height / 2), text="No entries to display", fill="#64748B")
            return
        plot_w = width - pad_l - pad_r
        plot_h = height - pad_t - pad_b
        self.create_line(pad_l, pad_t, pad_l, pad_t + plot_h, fill="#94A3B8")
        self.create_line(pad_l, pad_t + plot_h, pad_l + plot_w, pad_t + plot_h, fill="#94A3B8")
        for tick in range(0, y_max + 1, max(1, y_max // 5)):
            y = pad_t + plot_h - (tick / y_max) * plot_h
            self.create_line(pad_l - 4, y, pad_l + plot_w, y, fill="#E2E8F0")
            self.create_text(pad_l - 8, y, text=str(tick), anchor="e", fill="#475569", font=("Segoe UI", 8))
        points_count = max(len(entries) - 1, 1)
        for name, values, color in series:
            line_points = []
            series_points = []
            for idx, value in enumerate(values[: len(entries)]):
                x = pad_l + (idx / points_count) * plot_w
                y = pad_t + plot_h - (value / y_max) * plot_h
                line_points.extend([x, y])
                series_points.append(ChartPoint(x, y, entries[idx].entry_date, name, value, color))
            if len(line_points) >= 4:
                self.create_line(*line_points, fill=color, width=2, smooth=True, tags=("chart-line",))
            for point in series_points:
                self._chart_points.append(point)
                self.create_oval(
                    point.x - 4,
                    point.y - 4,
                    point.x + 4,
                    point.y + 4,
                    fill=point.color,
                    outline="#FFFFFF",
                    width=1,
                    tags=("chart-point",),
                )
        labels = [entries[0].entry_date, entries[-1].entry_date] if len(entries) > 1 else [entries[0].entry_date]
        self.create_text(pad_l, height - 14, text=labels[0], anchor="w", fill="#475569", font=("Segoe UI", 8))
        if len(labels) > 1:
            self.create_text(width - pad_r, height - 14, text=labels[1], anchor="e", fill="#475569", font=("Segoe UI", 8))
        legend_x = pad_l + 8
        for name, _, color in series:
            self.create_rectangle(legend_x, legend_top, legend_x + 10, legend_top + 10, fill=color, outline=color, tags=("chart-legend",))
            self.create_text(legend_x + 14, legend_top + 5, text=name, anchor="w", fill="#334155", font=("Segoe UI", 8), tags=("chart-legend",))
            legend_x += max(90, len(name) * 7)

        if locked_index is not None and locked_index < len(self._chart_points):
            self._locked_point_index = locked_index
            self._display_callout(locked_index)
        else:
            self._locked_point_index = None

    def _point_index_at(self, x: float, y: float) -> int | None:
        point = nearest_chart_point(self._chart_points, x, y)
        if point is None:
            return None
        return self._chart_points.index(point)

    def _display_callout(self, point_index: int) -> None:
        self.delete("chart-callout")
        point = self._chart_points[point_index]
        width = responsive_canvas_dimension(self.winfo_width(), int(self["width"]))
        text_x = point.x + 10 if point.x <= width * 0.62 else point.x - 10
        anchor = "sw" if point.x <= width * 0.62 else "se"
        text_y = point.y - 9
        if text_y < 28:
            text_y = point.y + 9
            anchor = "nw" if point.x <= width * 0.62 else "ne"
        text_id = self.create_text(
            text_x,
            text_y,
            text=chart_point_callout_text(point),
            anchor=anchor,
            fill="#172033",
            font=("Segoe UI", 9, "bold"),
            tags=("chart-callout",),
        )
        bounds = self.bbox(text_id)
        if bounds:
            rectangle_id = self.create_rectangle(
                bounds[0] - 6,
                bounds[1] - 4,
                bounds[2] + 6,
                bounds[3] + 4,
                fill="#FFF7D6",
                outline="#A16207",
                width=1,
                tags=("chart-callout",),
            )
            self.tag_lower(rectangle_id, text_id)
        self.create_oval(
            point.x - 7,
            point.y - 7,
            point.x + 7,
            point.y + 7,
            outline="#172033",
            width=2,
            tags=("chart-callout",),
        )

    def _show_hover_callout(self, event) -> None:
        if self._locked_point_index is not None:
            return
        point_index = self._point_index_at(event.x, event.y)
        self.config(cursor="hand2" if point_index is not None else "")
        if point_index is None:
            self._hovered_point_index = None
            self.delete("chart-callout")
            return
        if point_index != self._hovered_point_index:
            self._hovered_point_index = point_index
            self._display_callout(point_index)

    def _hide_hover_callout(self, _event=None) -> None:
        self.config(cursor="")
        if self._locked_point_index is None:
            self._hovered_point_index = None
            self.delete("chart-callout")

    def _toggle_click_callout(self, event) -> str:
        self.focus_set()
        point_index = self._point_index_at(event.x, event.y)
        if point_index is None:
            self._locked_point_index = None
            self._hovered_point_index = None
            self.delete("chart-callout")
            return "break"
        self._locked_point_index = None if point_index == self._locked_point_index else point_index
        self._hovered_point_index = point_index
        if self._locked_point_index is None:
            self.delete("chart-callout")
        else:
            self._display_callout(point_index)
        return "break"

    def _move_keyboard_callout(self, direction: int) -> str:
        if not self._chart_points:
            return "break"
        if self._locked_point_index is not None:
            current = self._locked_point_index
        elif self._hovered_point_index is not None:
            current = self._hovered_point_index
        else:
            current = 0 if direction < 0 else -1
        self._locked_point_index = (current + direction) % len(self._chart_points)
        self._hovered_point_index = self._locked_point_index
        self._display_callout(self._locked_point_index)
        return "break"

    def _lock_keyboard_callout(self, _event=None) -> str:
        if self._chart_points and self._locked_point_index is None:
            self._locked_point_index = self._hovered_point_index or 0
            self._display_callout(self._locked_point_index)
        return "break"

    def _clear_callout(self, _event=None) -> str:
        self._locked_point_index = None
        self._hovered_point_index = None
        self.delete("chart-callout")
        return "break"


PRIMARY_TAB_ORDER = (
    "Today's Check-In",
    "Review",
    "History / Manage Entries",
    "Treatment Events",
    "How Scoring Works",
)

REVIEW_ACTION_LABELS = (
    "Refresh",
    "Import Spreadsheet",
    "Generate PDF",
    "Analysis Workbook",
    "Open Reports Folder",
)

CHART_INTERACTION_HELP = (
    "Chart values: hover or click a point for its date and exact score. "
    "Keyboard: Tab to a chart, use Left/Right to move, Enter to show, and Escape to clear."
)


class PHQ9App(Tk):
    def __init__(self):
        super().__init__()
        self.title(APPLICATION_NAME)
        self.geometry("1180x780")
        self.minsize(980, 680)
        self.configure(bg="#F8FAFC")
        try:
            self.iconbitmap(default=str(ICON_PATH))
        except Exception:
            pass
        init_db()
        self._checkin_snapshot = None
        self._loaded_checkin_date = None
        self._history_snapshot = None
        self._loaded_history_date = None
        self.configure_ui_styles()
        self.create_menu()
        self.create_widgets()
        self.load_checkin_date(confirm_unsaved=False)
        self.load_history_date(confirm_unsaved=False)
        self.refresh_all()

    def configure_ui_styles(self):
        """Apply restrained, readable defaults without introducing a theme dependency."""
        self.option_add("*Font", ("Segoe UI", 10))
        self.option_add("*Button.padX", 10)
        self.option_add("*Button.padY", 5)
        self.option_add("*Text.Font", ("Segoe UI", 10))
        style = ttk.Style(self)
        style.configure(".", font=("Segoe UI", 10))
        style.configure("TNotebook.Tab", padding=(12, 7))
        style.configure("Treeview", rowheight=26)
        style.configure("TSpinbox", padding=COMPACT_SPINBOX_PADDING, arrowsize=14)
        style.configure("TCombobox", padding=3)

    def create_menu(self):
        menu = Menu(self)
        file_menu = Menu(menu, tearoff=0)
        file_menu.add_command(label="Import spreadsheet...", command=self.import_file)
        file_menu.add_command(label="Export analysis workbook...", command=self.choose_analysis_export)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.destroy)
        menu.add_cascade(label="File", menu=file_menu)
        help_menu = Menu(menu, tearoff=0)
        help_menu.add_command(label="How Scoring Works", command=self.show_scoring_help)
        menu.add_cascade(label="Help", menu=help_menu)
        self.config(menu=menu)

    def create_widgets(self):
        header = Frame(self, bg="#172033", padx=16, pady=12)
        header.pack(fill="x")
        Label(header, text=APPLICATION_NAME, fg="white", bg="#172033", font=("Segoe UI", 18, "bold")).pack(side=LEFT)
        Label(
            header,
            text="All data is stored locally in SQLite. No upload or cloud sync is performed by this app.",
            fg="#DDE7F3",
            bg="#172033",
            font=("Segoe UI", 10),
        ).pack(side=RIGHT)

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill=BOTH, expand=True, padx=12, pady=12)
        self.dashboard = Frame(self.notebook, bg="#F8FAFC")
        self.entry_tab = Frame(self.notebook, bg="#F8FAFC")
        self.events_tab = Frame(self.notebook, bg="#F8FAFC")
        self.history_tab = Frame(self.notebook, bg="#F8FAFC")
        self.scoring_tab = Frame(self.notebook, bg="#F8FAFC")
        tabs = {
            "Today's Check-In": self.entry_tab,
            "Review": self.dashboard,
            "History / Manage Entries": self.history_tab,
            "Treatment Events": self.events_tab,
            "How Scoring Works": self.scoring_tab,
        }
        for tab_name in PRIMARY_TAB_ORDER:
            self.notebook.add(tabs[tab_name], text=tab_name)
        self.build_dashboard()
        self.build_entry_tab()
        self.build_history_tab()
        self.build_events_tab()
        self.build_scoring_tab()
        self.notebook.select(self.entry_tab)
        self.bind_all("<F5>", self.refresh_from_shortcut)

    def build_dashboard(self):
        top = Frame(self.dashboard, bg="#F8FAFC")
        top.pack(fill="x", pady=(0, 10))
        Label(top, text="Review", bg="#F8FAFC", fg="#172033", font=("Segoe UI", 16, "bold")).pack(side=LEFT)
        actions = Frame(top, bg="#F8FAFC")
        actions.pack(side=RIGHT)
        action_commands = {
            "Refresh": self.refresh_all,
            "Import Spreadsheet": self.import_file,
            "Generate PDF": self.choose_report,
            "Analysis Workbook": self.choose_analysis_export,
            "Open Reports Folder": self.open_reports_folder,
        }
        for index, label in enumerate(REVIEW_ACTION_LABELS):
            Button(actions, text=label, command=action_commands[label], width=18).pack(
                side=LEFT,
                padx=(0 if index == 0 else 8, 0),
                pady=2,
            )

        summary = Frame(self.dashboard, bg="#FFFFFF", padx=14, pady=12, highlightthickness=1, highlightbackground="#CBD5E1")
        summary.pack(fill="x", pady=(0, 10))
        Label(summary, text="What stands out", bg="#FFFFFF", fg="#172033", font=("Segoe UI", 12, "bold")).pack(anchor="w")
        self.review_summary = Label(summary, text="No check-ins are available yet.", bg="#FFFFFF", fg="#334155", wraplength=1060, justify=LEFT)
        self.review_summary.pack(anchor="w", pady=(6, 6))
        self.review_highlights = []
        for _ in range(4):
            label = Label(summary, text="", bg="#FFFFFF", fg="#334155", wraplength=1040, justify=LEFT)
            label.pack(anchor="w", pady=2)
            self.review_highlights.append(label)
        Label(
            summary,
            text=CHART_INTERACTION_HELP,
            bg="#FFFFFF",
            fg="#475569",
            wraplength=1040,
            justify=LEFT,
            font=("Segoe UI", 9),
        ).pack(anchor="w", pady=(8, 0))

        self.review_notebook = ttk.Notebook(self.dashboard)
        self.review_notebook.pack(fill=BOTH, expand=True)
        overview = Frame(self.review_notebook, bg="#F8FAFC")
        treatment = Frame(self.review_notebook, bg="#F8FAFC")
        long_term = Frame(self.review_notebook, bg="#F8FAFC")
        self.review_notebook.add(overview, text="Recent Trends")
        self.review_notebook.add(treatment, text="Treatment Cycles")
        self.review_notebook.add(long_term, text="Long-Term Trends")

        self.recent_trend_charts = {}
        self.long_term_trend_charts = {}
        for index, questionnaire_id in enumerate(QUESTIONNAIRE_ORDER):
            chart = LineChart(overview, width=520, height=250)
            chart.pack(
                side=LEFT,
                fill=BOTH,
                expand=True,
                padx=(0 if index == 0 else 5, 0 if index == len(QUESTIONNAIRE_ORDER) - 1 else 5),
                pady=6,
            )
            self.recent_trend_charts[questionnaire_id] = chart

        self.cycle_labels = []
        self.cycle_charts = []
        for side in (LEFT, RIGHT):
            cycle_frame = Frame(treatment, bg="#F8FAFC")
            cycle_frame.pack(side=side, fill=BOTH, expand=True, padx=6, pady=6)
            label = Label(cycle_frame, text="", bg="#F8FAFC", fg="#334155", wraplength=500, justify=LEFT)
            label.pack(anchor="w", pady=(0, 5))
            chart = LineChart(cycle_frame, width=500, height=220)
            chart.pack(fill=BOTH, expand=True)
            self.cycle_labels.append(label)
            self.cycle_charts.append(chart)

        for index, questionnaire_id in enumerate(QUESTIONNAIRE_ORDER):
            chart = LineChart(long_term, width=1050, height=220)
            chart.pack(fill=BOTH, expand=True, padx=6, pady=(6 if index == 0 else 3, 6 if index == len(QUESTIONNAIRE_ORDER) - 1 else 3))
            self.long_term_trend_charts[questionnaire_id] = chart

        # Compatibility aliases retained for existing tests and integrations.
        self.phq_recent_chart = self.recent_trend_charts["phq9"]
        self.gad_recent_chart = self.recent_trend_charts["gad7"]
        self.phq_long_chart = self.long_term_trend_charts["phq9"]
        self.gad_long_chart = self.long_term_trend_charts["gad7"]

    def build_entry_tab(self):
        self.entry_canvas = Canvas(self.entry_tab, bg="#F8FAFC", highlightthickness=0)
        self.entry_scrollbar = ttk.Scrollbar(self.entry_tab, orient="vertical", command=self.entry_canvas.yview)
        self.entry_canvas.configure(yscrollcommand=self.entry_scrollbar.set)
        self.entry_scrollbar.pack(side=RIGHT, fill="y")
        self.entry_canvas.pack(side=LEFT, fill=BOTH, expand=True)
        self.entry_scroll_content = Frame(self.entry_canvas, bg="#F8FAFC")
        self.entry_scroll_window = self.entry_canvas.create_window(
            (0, 0),
            window=self.entry_scroll_content,
            anchor="nw",
        )
        self.entry_scroll_content.bind(
            "<Configure>",
            lambda _event: self.entry_canvas.configure(scrollregion=self.entry_canvas.bbox("all")),
        )
        self.entry_canvas.bind(
            "<Configure>",
            lambda event: self.entry_canvas.itemconfigure(self.entry_scroll_window, width=event.width),
        )
        self.bind_all("<MouseWheel>", self.scroll_today_checkin)

        form = LabelFrame(self.entry_scroll_content, text="Today's Check-In", bg="#F8FAFC", padx=12, pady=12)
        form.pack(fill=BOTH, expand=True, padx=4, pady=4)
        Button(form, text="← Previous Day", command=lambda: self.navigate_checkin_date(-1)).grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.date_var = StringVar(value=date.today().isoformat())
        self.checkin_date_entry = Entry(form, textvariable=self.date_var, width=16, justify="center")
        self.checkin_date_entry.grid(row=0, column=1, sticky="w", padx=8, pady=4)
        self.checkin_date_entry.bind("<Return>", self.load_checkin_date_from_shortcut)
        self.checkin_date_entry.bind("<FocusOut>", self.auto_load_checkin_date_if_safe)
        self.next_day_button = Button(form, text="Next Day →", command=lambda: self.navigate_checkin_date(1))
        self.next_day_button.grid(row=0, column=2, sticky="w", padx=8)
        Button(form, text="Load Date", command=self.load_checkin_date).grid(row=0, column=3, sticky="w", padx=8)
        self.checkin_mode = Label(form, text="New daily entry", bg="#F8FAFC", fg="#047857", font=("Segoe UI", 10, "bold"))
        self.checkin_mode.grid(row=0, column=4, sticky="w", padx=8)

        self.assessment_item_vars = {}
        self.questionnaire_selected_vars = {}
        self.questionnaire_status_labels = {}
        self.questionnaire_boxes = {}
        self.active_questionnaire_id = StringVar(value=QUESTIONNAIRE_ORDER[0])
        assessment_area = Frame(form, bg="#F8FAFC")
        assessment_area.grid(row=1, column=0, columnspan=5, sticky="nsew", pady=(8, 4))
        selector = LabelFrame(
            assessment_area,
            text="Questionnaire status and form selection",
            bg="#F8FAFC",
            padx=8,
            pady=6,
        )
        selector.grid(row=0, column=0, sticky="we", pady=(0, 8))
        Label(selector, text="Open one form at a time. Include every completed form you want to save.", bg="#F8FAFC").grid(
            row=0, column=0, columnspan=3, sticky="w", pady=(0, 4)
        )
        for roster_row, assessment_id in enumerate(QUESTIONNAIRE_ORDER, start=1):
            definition = QUESTIONNAIRES[assessment_id]
            selected_var = IntVar(value=1)
            self.questionnaire_selected_vars[assessment_id] = selected_var
            ttk.Radiobutton(
                selector,
                text=definition.display_name,
                value=assessment_id,
                variable=self.active_questionnaire_id,
                command=self.update_questionnaire_visibility,
            ).grid(row=roster_row, column=0, sticky="w", padx=(0, 8), pady=2)
            Checkbutton(
                selector,
                text="Include in save",
                variable=selected_var,
                bg="#F8FAFC",
            ).grid(row=roster_row, column=1, sticky="w", padx=8, pady=2)
            status = Label(selector, text="Not completed", bg="#F8FAFC", fg="#64748B")
            status.grid(row=roster_row, column=2, sticky="w", padx=8, pady=2)
            self.questionnaire_status_labels[assessment_id] = status
        for assessment_id in ASSESSMENT_ORDER:
            definition = ASSESSMENTS[assessment_id]
            box = LabelFrame(assessment_area, text=definition.display_name, bg="#F8FAFC", padx=10, pady=10)
            box.grid(row=1, column=0, sticky="nsew")
            self.questionnaire_boxes[assessment_id] = box
            self.assessment_item_vars[assessment_id] = []
            Label(
                box,
                text=UNIVERSAL_SAFETY_MESSAGE,
                bg="#FFF7ED",
                fg="#9A3412",
                wraplength=940,
                justify=LEFT,
                padx=10,
                pady=8,
            ).grid(row=0, column=0, columnspan=2, sticky="we", pady=(0, 8))
            Label(
                box,
                text=definition.timeframe_text,
                bg="#F8FAFC",
                fg="#172033",
                wraplength=940,
                justify=LEFT,
                font=("Segoe UI", 10, "bold"),
            ).grid(row=1, column=0, columnspan=2, sticky="w", pady=(0, 8))
            for idx, question in enumerate(definition.items, start=1):
                Label(box, text=f"{idx}. {question.report_label or question.prompt}", bg="#F8FAFC", wraplength=650, justify=LEFT).grid(
                    row=idx + 1, column=0, sticky="w", pady=3
                )
                var = StringVar(value=UNANSWERED_RESPONSE)
                self.assessment_item_vars[assessment_id].append(var)
                ttk.Combobox(
                    box,
                    textvariable=var,
                    values=questionnaire_response_choices(question),
                    state="readonly",
                    width=34,
                ).grid(row=idx + 1, column=1, sticky="w", padx=8, pady=3)
        assessment_area.columnconfigure(0, weight=1)
        self.update_questionnaire_visibility()

        Button(form, text="Save Today's Check-In", command=self.save_entry).grid(row=2, column=1, sticky="w", padx=8, pady=(10, 8))
        Label(
            form,
            text="Mindful check-ins intentionally require each symptom to be considered individually; previous responses are never copied or autofilled.",
            bg="#F8FAFC",
            fg="#475569",
            wraplength=900,
            justify=LEFT,
        ).grid(row=3, column=0, columnspan=5, sticky="w", pady=(2, 8))

        self.checkin_details = LabelFrame(
            form,
            text="Anything else to record about today?",
            bg="#F8FAFC",
            padx=10,
            pady=10,
        )
        self.checkin_details.grid(row=4, column=0, columnspan=5, sticky="we", pady=(4, 2))
        Label(
            self.checkin_details,
            text="These details are optional. Your core check-in has already been recorded.",
            bg="#F8FAFC",
            fg="#475569",
        ).grid(row=0, column=0, columnspan=4, sticky="w", pady=(0, 6))

        Label(self.checkin_details, text="Personal note (optional)", bg="#F8FAFC").grid(row=1, column=0, sticky="nw", pady=(4, 2))
        self.notes_box = Text(self.checkin_details, height=4, width=72)
        self.notes_box.grid(row=1, column=1, columnspan=3, sticky="we", padx=8, pady=(4, 2))
        Label(self.checkin_details, text="Optional note tag", bg="#F8FAFC").grid(row=2, column=0, sticky="w", pady=(4, 2))
        self.note_tag = StringVar(value="")
        ttk.Combobox(
            self.checkin_details,
            textvariable=self.note_tag,
            values=["", "Finances", "Work", "Family", "Health", "Sleep", "Relationships", "Other"],
            state="readonly",
            width=20,
        ).grid(row=2, column=1, sticky="w", padx=8, pady=(4, 2))

        events_box = LabelFrame(self.checkin_details, text="Treatment or health-related events", bg="#F8FAFC", padx=8, pady=6)
        events_box.grid(row=3, column=0, columnspan=4, sticky="we", pady=(10, 2))
        self.checkin_event_vars = {}
        for idx, event_type in enumerate(TREATMENT_EVENT_TYPES):
            var = IntVar(value=0)
            self.checkin_event_vars[event_type] = var
            Checkbutton(events_box, text=event_type, variable=var, bg="#F8FAFC").grid(row=idx // 4, column=idx % 4, sticky="w", padx=6)
        self.include_custom_event = IntVar(value=0)
        Checkbutton(events_box, text="Add Custom Event", variable=self.include_custom_event, bg="#F8FAFC").grid(row=2, column=0, sticky="w", padx=6)
        self.custom_event_type = StringVar(value="")
        Entry(events_box, textvariable=self.custom_event_type, width=28).grid(row=2, column=1, sticky="w", padx=6)
        Label(self.checkin_details, text="Event description (optional)", bg="#F8FAFC").grid(row=4, column=0, sticky="nw", pady=(6, 0))
        self.checkin_event_desc = Text(self.checkin_details, height=2, width=72)
        self.checkin_event_desc.grid(row=4, column=1, columnspan=3, sticky="we", padx=8, pady=(6, 0))
        self.save_optional_details_button = Button(
            self.checkin_details,
            text="Save Optional Details",
            command=self.save_optional_details,
        )
        self.save_optional_details_button.grid(row=5, column=1, sticky="w", padx=8, pady=10)
        Button(self.checkin_details, text="Not Right Now", command=self.hide_checkin_details).grid(row=5, column=2, sticky="w", padx=8, pady=10)
        Button(self.checkin_details, text="Add Another Treatment Event", command=lambda: self.open_history_for_date(self.date_var.get())).grid(row=5, column=3, sticky="e", padx=8, pady=10)
        self.checkin_details.columnconfigure(3, weight=1)
        self.checkin_details.grid_remove()
        form.columnconfigure(4, weight=1)

    def build_history_tab(self):
        top = Frame(self.history_tab, bg="#F8FAFC")
        top.pack(fill="x", padx=6, pady=6)
        self.history_previous_button = Button(top, text="← Previous Day", command=lambda: self.navigate_history_date(-1))
        self.history_previous_button.pack(side=LEFT, padx=(0, 8))
        Label(top, text="Manage records for date (YYYY-MM-DD)", bg="#F8FAFC").pack(side=LEFT)
        self.history_date = StringVar(value=date.today().isoformat())
        self.history_date_entry = Entry(top, textvariable=self.history_date, width=16)
        self.history_date_entry.pack(side=LEFT, padx=8)
        self.history_date_entry.bind("<Return>", self.load_history_date_from_shortcut)
        self.history_date_entry.bind("<FocusOut>", self.auto_load_history_date_if_safe)
        Button(top, text="Load Date", command=self.load_history_date).pack(side=LEFT)
        self.history_next_button = Button(top, text="Next Day →", command=lambda: self.navigate_history_date(1))
        self.history_next_button.pack(side=LEFT, padx=8)
        self.history_status = Label(top, text="", bg="#F8FAFC", fg="#334155")
        self.history_status.pack(side=LEFT, padx=12)

        self.history_assessment_vars = {}
        self.history_assessment_ids = {}
        self.history_assessment_buttons = {}
        self.history_assessment_widgets = {}
        assessments_frame = Frame(self.history_tab, bg="#F8FAFC")
        assessments_frame.pack(fill="x", padx=6, pady=4)
        for col_idx, assessment_id in enumerate(ASSESSMENT_ORDER):
            definition = ASSESSMENTS[assessment_id]
            box = LabelFrame(assessments_frame, text=f"{definition.display_name} responses", bg="#F8FAFC", padx=8, pady=8)
            box.grid(row=0, column=col_idx, sticky="nsew", padx=(0, 8))
            self.history_assessment_vars[assessment_id] = []
            self.history_assessment_widgets[assessment_id] = []
            for idx in range(definition.item_count):
                Label(box, text=f"{idx + 1}", bg="#F8FAFC").grid(row=0, column=idx, padx=2)
                var = IntVar(value=0)
                self.history_assessment_vars[assessment_id].append(var)
                spinner = ttk.Spinbox(box, from_=0, to=3, textvariable=var, width=3)
                spinner.grid(row=1, column=idx, padx=2)
                self.history_assessment_widgets[assessment_id].append(spinner)
            update_button = Button(box, text="Update Existing Entry", command=lambda aid=assessment_id: self.save_history_assessment(aid))
            update_button.grid(row=2, column=0, columnspan=4, sticky="w", pady=(8, 0))
            delete_button = Button(box, text="Delete Assessment", command=lambda aid=assessment_id: self.delete_history_assessment(aid))
            delete_button.grid(row=2, column=4, columnspan=5, sticky="e", pady=(8, 0))
            self.history_assessment_buttons[assessment_id] = (update_button, delete_button)
            assessments_frame.columnconfigure(col_idx, weight=1)

        notes_box = LabelFrame(self.history_tab, text="Daily note", bg="#F8FAFC", padx=8, pady=8)
        notes_box.pack(fill="x", padx=6, pady=4)
        self.history_notes = Text(notes_box, height=3, width=80)
        self.history_notes.grid(row=0, column=0, columnspan=3, sticky="we")
        self.history_note_tag = StringVar(value="")
        self.history_note_tag_widget = ttk.Combobox(notes_box, textvariable=self.history_note_tag, values=["", "Finances", "Work", "Family", "Health", "Sleep", "Relationships", "Other"], state="readonly", width=20)
        self.history_note_tag_widget.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.history_note_update_button = Button(notes_box, text="Update Existing Note", command=self.save_history_note)
        self.history_note_update_button.grid(row=1, column=1, padx=8, pady=(6, 0))
        self.history_note_delete_button = Button(notes_box, text="Delete Note", command=self.delete_history_note)
        self.history_note_delete_button.grid(row=1, column=2, pady=(6, 0))
        notes_box.columnconfigure(0, weight=1)

        event_box = LabelFrame(self.history_tab, text="Treatment events", bg="#F8FAFC", padx=8, pady=8)
        event_box.pack(fill=BOTH, expand=True, padx=6, pady=4)
        self.history_events_table = ttk.Treeview(event_box, columns=["Type", "Description"], show="headings", height=6)
        self.history_events_table.heading("Type", text="Type")
        self.history_events_table.heading("Description", text="Description")
        self.history_events_table.column("Type", width=220)
        self.history_events_table.column("Description", width=650)
        self.history_events_table.grid(row=0, column=0, columnspan=4, sticky="nsew")
        self.history_events_table.bind("<<TreeviewSelect>>", self.select_history_event)
        self.history_event_type = StringVar(value="Therapy")
        self.history_event_type_widget = ttk.Combobox(event_box, textvariable=self.history_event_type, values=[*TREATMENT_EVENT_TYPES, "Custom event"], width=28)
        self.history_event_type_widget.grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.history_event_desc = StringVar(value="")
        self.history_event_desc_widget = Entry(event_box, textvariable=self.history_event_desc, width=65)
        self.history_event_desc_widget.grid(row=1, column=1, sticky="we", padx=6, pady=(6, 0))
        self.history_event_add_button = Button(event_box, text="Add Another Treatment Event", command=self.add_history_event)
        self.history_event_add_button.grid(row=1, column=2, padx=4, pady=(6, 0))
        self.history_event_update_button = Button(event_box, text="Update Selected", command=self.update_history_event)
        self.history_event_update_button.grid(row=2, column=2, padx=4, pady=(6, 0))
        self.history_event_delete_button = Button(event_box, text="Delete Selected", command=self.delete_history_event)
        self.history_event_delete_button.grid(row=2, column=3, padx=4, pady=(6, 0))
        event_box.columnconfigure(1, weight=1)
        event_box.rowconfigure(0, weight=1)

    def build_events_tab(self):
        form = LabelFrame(self.events_tab, text="Medication/Treatment Event Marker", bg="#F8FAFC", padx=12, pady=12)
        form.pack(fill="x", padx=4, pady=4)
        self.event_date = StringVar(value=date.today().isoformat())
        self.event_type = StringVar(value="Ketamine infusion")
        Label(form, text="Date (YYYY-MM-DD)", bg="#F8FAFC").grid(row=0, column=0, sticky="w")
        Entry(form, textvariable=self.event_date, width=16).grid(row=0, column=1, sticky="w", padx=8)
        Label(form, text="Type", bg="#F8FAFC").grid(row=0, column=2, sticky="w", padx=(18, 0))
        ttk.Combobox(
            form,
            textvariable=self.event_type,
            values=[*TREATMENT_EVENT_TYPES, "Custom event"],
            width=30,
        ).grid(row=0, column=3, sticky="w", padx=8)
        Label(form, text="Description", bg="#F8FAFC").grid(row=1, column=0, sticky="nw", pady=(8, 0))
        self.event_desc = Text(form, height=4, width=80)
        self.event_desc.grid(row=1, column=1, columnspan=3, sticky="we", padx=8, pady=(8, 0))
        Button(form, text="Add Another Treatment Event", command=self.save_event).grid(row=2, column=1, sticky="w", padx=8, pady=10)

        self.events_table = ttk.Treeview(self.events_tab, columns=["Date", "Type", "Description"], show="headings", height=16)
        for col, width in [("Date", 120), ("Type", 180), ("Description", 620)]:
            self.events_table.heading(col, text=col)
            self.events_table.column(col, width=width, anchor="w")
        self.events_table.pack(fill=BOTH, expand=True, padx=4, pady=10)

    def build_scoring_tab(self):
        Label(self.scoring_tab, text="How Scoring Works", bg="#F8FAFC", fg="#172033", font=("Segoe UI", 16, "bold")).pack(anchor="w", padx=12, pady=(12, 6))
        explanation = Text(self.scoring_tab, wrap="word", bg="#FFFFFF", fg="#172033", font=("Segoe UI", 10), padx=12, pady=12)
        explanation.insert("1.0", SCORING_EXPLANATION)
        explanation.config(state="disabled")
        explanation.pack(fill=BOTH, expand=True, padx=12, pady=(0, 12))

    def show_scoring_help(self):
        self.notebook.select(self.scoring_tab)

    def refresh_from_shortcut(self, _event=None) -> str:
        """Refresh local Review data using the conventional F5 shortcut."""
        self.refresh_all()
        return "break"

    def show_checkin_details(self):
        self.checkin_details.grid()
        self.update_idletasks()
        self.entry_canvas.configure(scrollregion=self.entry_canvas.bbox("all"))
        self.entry_canvas.yview_moveto(1.0)

    def hide_checkin_details(self):
        self.checkin_details.grid_remove()
        self.entry_canvas.yview_moveto(0.0)
        messagebox.showinfo("Today's Check-In", "Today's check-in is recorded.")

    def scroll_today_checkin(self, event):
        """Scroll the Today form only when the pointer is inside that tab."""
        if self.notebook.select() != str(self.entry_tab):
            return None
        widget = self.winfo_containing(event.x_root, event.y_root)
        while widget is not None and widget is not self.entry_canvas:
            widget = getattr(widget, "master", None)
        if widget is not self.entry_canvas:
            return None
        units = -1 if event.delta > 0 else 1
        self.entry_canvas.yview_scroll(units * 3, "units")
        return "break"

    def checkin_state(self) -> tuple:
        return (
            self.active_questionnaire_id.get(),
            tuple((assessment_id, self.questionnaire_selected_vars[assessment_id].get()) for assessment_id in QUESTIONNAIRE_ORDER),
            tuple(tuple(var.get() for var in self.assessment_item_vars[assessment_id]) for assessment_id in ASSESSMENT_ORDER),
            self.notes_box.get("1.0", END).strip(),
            self.note_tag.get().strip(),
            tuple((event_type, var.get()) for event_type, var in self.checkin_event_vars.items()),
            self.include_custom_event.get(),
            self.custom_event_type.get().strip(),
            self.checkin_event_desc.get("1.0", END).strip(),
        )

    def update_questionnaire_visibility(self):
        active_id = self.active_questionnaire_id.get()
        if active_id not in QUESTIONNAIRES:
            active_id = QUESTIONNAIRE_ORDER[0]
            self.active_questionnaire_id.set(active_id)
        for assessment_id in QUESTIONNAIRE_ORDER:
            box = self.questionnaire_boxes[assessment_id]
            if assessment_id == active_id:
                box.grid(row=1, column=0, sticky="nsew")
            else:
                box.grid_remove()

    def has_unsaved_checkin_changes(self) -> bool:
        return self._checkin_snapshot is not None and self.checkin_state() != self._checkin_snapshot

    def confirm_discard_checkin_changes(self) -> bool:
        if not self.has_unsaved_checkin_changes():
            return True
        return messagebox.askyesno(
            "Discard unsaved check-in changes?",
            "You have unsaved check-in changes. Choose Yes to discard them and load another date, "
            "or No to keep editing this date.",
        )

    def navigate_checkin_date(self, days: int):
        if not self.confirm_discard_checkin_changes():
            return
        try:
            target = shift_calendar_date(self.date_var.get(), days)
        except ValueError as exc:
            messagebox.showinfo("Date unavailable", str(exc))
            return
        self.date_var.set(target)
        self.load_checkin_date(confirm_unsaved=False)

    def load_checkin_date_from_shortcut(self, _event=None) -> str:
        self.load_checkin_date()
        return "break"

    def auto_load_checkin_date_if_safe(self, _event=None) -> None:
        """Load a newly entered date on focus change only when no edits could be lost."""
        entry_date = parse_date(self.date_var.get())
        if not entry_date or entry_date == self._loaded_checkin_date:
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            return
        if self.has_unsaved_checkin_changes():
            self.checkin_mode.config(
                text="Date changed — choose Load Date; current unsaved edits are still preserved.",
                fg="#B45309",
            )
            return
        self.load_checkin_date(confirm_unsaved=False)

    def load_checkin_date(self, confirm_unsaved: bool = True):
        if confirm_unsaved and not self.confirm_discard_checkin_changes():
            return
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            messagebox.showerror("Future date", "Future check-ins are not available.")
            return
        self.date_var.set(entry_date)
        day_data = fetch_day_data(entry_date)
        assessments = day_data["assessments"]
        has_data = bool(assessments or day_data["notes"] or day_data["events"])
        active_id = next((assessment_id for assessment_id in QUESTIONNAIRE_ORDER if assessment_id in assessments), QUESTIONNAIRE_ORDER[0])
        self.active_questionnaire_id.set(active_id)
        for assessment_id in ASSESSMENT_ORDER:
            definition = QUESTIONNAIRES[assessment_id]
            row = assessments.get(assessment_id)
            self.questionnaire_selected_vars[assessment_id].set(1 if row or not has_data else 0)
            self.questionnaire_status_labels[assessment_id].config(
                text="Complete" if row else "Not completed",
                fg="#047857" if row else "#64748B",
            )
            for idx, var in enumerate(self.assessment_item_vars[assessment_id]):
                var.set(
                    display_questionnaire_response(definition.items[idx], row.items[idx])
                    if row
                    else UNANSWERED_RESPONSE
                )
        self.update_questionnaire_visibility()
        self.notes_box.delete("1.0", END)
        self.notes_box.insert("1.0", day_data["notes"])
        self.note_tag.set(day_data["note_tag"])
        existing_types = {event[2] for event in day_data["events"]}
        for event_type, var in self.checkin_event_vars.items():
            var.set(1 if event_type in existing_types else 0)
        custom_events = [event for event in day_data["events"] if event[2] not in TREATMENT_EVENT_TYPES]
        self.include_custom_event.set(1 if custom_events else 0)
        self.custom_event_type.set(custom_events[0][2] if custom_events else "")
        self.checkin_event_desc.delete("1.0", END)
        self.checkin_mode.config(
            text=(
                f"Editing existing records for {entry_date}"
                if has_data
                else f"New check-in for {entry_date}"
            ),
            fg="#B45309" if has_data else "#047857",
        )
        self.next_day_button.config(state="disabled" if entry_date == date.today().isoformat() else "normal")
        self._loaded_checkin_date = entry_date
        self._checkin_snapshot = self.checkin_state()

    def open_history_for_date(self, raw_date: str):
        if not self.confirm_discard_history_changes():
            return
        entry_date = parse_date(raw_date)
        if entry_date:
            self.history_date.set(entry_date)
        self.notebook.select(self.history_tab)
        self.load_history_date(confirm_unsaved=False)

    def history_state(self) -> tuple:
        return (
            tuple(
                (assessment_id, tuple(var.get() for var in self.history_assessment_vars[assessment_id]))
                for assessment_id in ASSESSMENT_ORDER
            ),
            self.history_notes.get("1.0", END).strip(),
            self.history_note_tag.get().strip(),
            self.history_event_type.get().strip(),
            self.history_event_desc.get().strip(),
        )

    def has_unsaved_history_changes(self) -> bool:
        return self._history_snapshot is not None and self.history_state() != self._history_snapshot

    def confirm_discard_history_changes(self) -> bool:
        if not self.has_unsaved_history_changes():
            return True
        return messagebox.askyesno(
            "Discard unsaved History changes?",
            "You have unsaved History changes. Choose Yes to discard them and load another date, "
            "or No to keep editing the currently loaded date.",
        )

    def navigate_history_date(self, days: int):
        if not self.confirm_discard_history_changes():
            return
        try:
            target = shift_calendar_date(self.history_date.get(), days)
        except ValueError as exc:
            messagebox.showinfo("Date unavailable", str(exc))
            return
        self.history_date.set(target)
        self.load_history_date(confirm_unsaved=False)

    def load_history_date_from_shortcut(self, _event=None) -> str:
        self.load_history_date()
        return "break"

    def auto_load_history_date_if_safe(self, _event=None) -> None:
        """Load a newly entered History date only when no unsaved edits could be lost."""
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError:
            return
        if entry_date == self._loaded_history_date:
            return
        if self.has_unsaved_history_changes():
            self.history_status.config(
                text="Date changed. Choose Load Date to continue; edits for the currently loaded date are preserved."
            )
            return
        self.load_history_date(confirm_unsaved=False)

    def _loaded_history_action_date(self) -> str | None:
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError as exc:
            messagebox.showerror("Date unavailable", str(exc))
            return None
        if entry_date != self._loaded_history_date:
            messagebox.showinfo(
                "Load the selected date first",
                "The date field differs from the records currently shown. Load the selected date before making changes.",
            )
            return None
        return entry_date

    def load_history_date(self, confirm_unsaved: bool = True):
        if confirm_unsaved and not self.confirm_discard_history_changes():
            return
        try:
            entry_date = validate_nonfuture_date(self.history_date.get())
        except ValueError as exc:
            messagebox.showerror("Date unavailable", str(exc))
            return
        self.history_date.set(entry_date)
        day_data = fetch_day_data(entry_date)
        assessments = day_data["assessments"]
        for assessment_id in ASSESSMENT_ORDER:
            row = assessments.get(assessment_id)
            self.history_assessment_ids[assessment_id] = row.id if row else None
            for idx, var in enumerate(self.history_assessment_vars[assessment_id]):
                var.set(row.items[idx] if row else 0)
        self.history_notes.config(state="normal")
        self.history_notes.delete("1.0", END)
        self.history_notes.insert("1.0", day_data["notes"])
        self.history_note_tag.set(day_data["note_tag"])
        for item in self.history_events_table.get_children():
            self.history_events_table.delete(item)
        for event_id, _event_date, event_type, description in day_data["events"]:
            self.history_events_table.insert("", END, iid=str(event_id), values=[event_type, description])
        self.history_event_type.set("Therapy")
        self.history_event_desc.set("")
        has_records = history_record_count(day_data) > 0
        has_note = bool(day_data["notes"] or day_data["note_tag"])
        for assessment_id, (update_button, delete_button) in self.history_assessment_buttons.items():
            state = "normal" if assessment_id in assessments else "disabled"
            update_button.config(state=state)
            delete_button.config(state=state)
            for spinner in self.history_assessment_widgets[assessment_id]:
                spinner.config(state=state)
        note_edit_state = "normal" if has_records else "disabled"
        self.history_notes.config(state=note_edit_state)
        self.history_note_tag_widget.config(state="readonly" if has_records else "disabled")
        self.history_note_update_button.config(
            state=note_edit_state,
            text="Update Existing Note" if has_note else "Add Daily Note",
        )
        self.history_note_delete_button.config(state="normal" if has_note else "disabled")
        self.history_event_type_widget.config(state="normal" if has_records else "disabled")
        self.history_event_desc_widget.config(state="normal" if has_records else "disabled")
        self.history_event_add_button.config(state="normal" if has_records else "disabled")
        self.history_event_update_button.config(state="normal" if day_data["events"] else "disabled")
        self.history_event_delete_button.config(state="normal" if day_data["events"] else "disabled")
        self.history_status.config(text=history_status_text(entry_date, day_data))
        self.history_next_button.config(state="disabled" if entry_date == date.today().isoformat() else "normal")
        self._loaded_history_date = entry_date
        self._history_snapshot = self.history_state()

    def save_history_assessment(self, assessment_id: str):
        if not self._loaded_history_action_date():
            return
        entry_id = self.history_assessment_ids.get(assessment_id)
        if not entry_id:
            messagebox.showinfo("No existing assessment", "Use Today's Check-In to create a new assessment for this date.")
            return
        items = [int(var.get()) for var in self.history_assessment_vars[assessment_id]]
        try:
            update_assessment_entry(entry_id, assessment_id, items)
            self.refresh_all()
            self.load_history_date(confirm_unsaved=False)
            messagebox.showinfo("Updated", f"Updated the existing {ASSESSMENTS[assessment_id].display_name} assessment without changing its record ID.")
        except Exception as exc:
            messagebox.showerror("Update failed", str(exc))

    def delete_history_assessment(self, assessment_id: str):
        if not self._loaded_history_action_date():
            return
        entry_id = self.history_assessment_ids.get(assessment_id)
        if not entry_id:
            return
        name = ASSESSMENTS[assessment_id].display_name
        if not messagebox.askyesno(
            "Permanently delete assessment?",
            f"Delete the {name} assessment for {self.history_date.get()}? This cannot be undone.",
        ):
            return
        delete_assessment_entry(entry_id, assessment_id)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def save_history_note(self):
        entry_date = self._loaded_history_action_date()
        if not entry_date:
            return
        update_daily_note(entry_date, self.history_notes.get("1.0", END).strip(), self.history_note_tag.get().strip())
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)
        messagebox.showinfo("Daily note updated", f"Updated the daily note for {entry_date} without creating a duplicate.")

    def delete_history_note(self):
        entry_date = self._loaded_history_action_date()
        if not entry_date or not messagebox.askyesno(
            "Permanently delete daily note?",
            f"Delete the daily note for {entry_date}? This cannot be undone.",
        ):
            return
        delete_daily_note(entry_date)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def select_history_event(self, _event=None):
        selected = self.history_events_table.selection()
        if not selected:
            return
        already_unsaved = self.has_unsaved_history_changes()
        values = self.history_events_table.item(selected[0], "values")
        self.history_event_type.set(values[0])
        self.history_event_desc.set(values[1])
        if not already_unsaved:
            self._history_snapshot = self.history_state()

    def add_history_event(self):
        entry_date = self._loaded_history_action_date()
        event_type = self.history_event_type.get().strip()
        if not entry_date or not event_type:
            messagebox.showerror("Missing information", "Enter a valid date and event type.")
            return
        add_event(entry_date, event_type, self.history_event_desc.get().strip(), dedupe=False)
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def update_history_event(self):
        selected = self.history_events_table.selection()
        entry_date = self._loaded_history_action_date()
        if not selected or not entry_date:
            messagebox.showinfo("Select an event", "Select the treatment event to update.")
            return
        update_event(int(selected[0]), entry_date, self.history_event_type.get().strip() or "Treatment event", self.history_event_desc.get().strip())
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def delete_history_event(self):
        if not self._loaded_history_action_date():
            return
        selected = self.history_events_table.selection()
        if not selected or not messagebox.askyesno(
            "Permanently delete treatment event?",
            "Delete the selected treatment event? This cannot be undone.",
        ):
            return
        delete_event(int(selected[0]))
        self.refresh_all()
        self.load_history_date(confirm_unsaved=False)

    def refresh_all(self):
        self.refresh_events()
        entries_by_questionnaire = {
            questionnaire_id: fetch_assessment_entries(questionnaire_id)
            for questionnaire_id in QUESTIONNAIRE_ORDER
        }
        phq_entries = entries_by_questionnaire["phq9"]
        gad_entries = entries_by_questionnaire["gad7"]
        all_dates = sorted(
            {
                row.entry_date
                for entries in entries_by_questionnaire.values()
                for row in entries
            }
        )
        if all_dates:
            latest_date = all_dates[-1]
            interpreted_profile_ids = [
                questionnaire_id
                for questionnaire_id in QUESTIONNAIRE_ORDER
                if QUESTIONNAIRES[questionnaire_id].interpretation_policy == "validated_builtin"
                and questionnaire_profile_omission_reason(QUESTIONNAIRES[questionnaire_id]) is None
            ]
            self.review_summary.config(
                text=overall_pattern_summary(
                    phq_entries,
                    gad_entries,
                    latest_date,
                    interpreted_profile_ids,
                    entries_by_questionnaire,
                )
            )
            highlights = []
            for questionnaire_id in interpreted_profile_ids:
                highlights.extend(
                    symptom_highlights(
                        questionnaire_id,
                        entries_by_questionnaire[questionnaire_id],
                        latest_date,
                        limit=2,
                    )
                )
            for index, label in enumerate(self.review_highlights):
                label.config(text=f"• {highlights[index]}" if index < len(highlights) else "")

            chart_colors = ("#2563EB", "#0F766E", "#7C3AED", "#B45309")
            for index, questionnaire_id in enumerate(QUESTIONNAIRE_ORDER):
                trend_series = build_questionnaire_trend_series(
                    questionnaire_id, entries_by_questionnaire[questionnaire_id]
                )
                definition = QUESTIONNAIRES[questionnaire_id]
                color = chart_colors[index % len(chart_colors)]
                if trend_series:
                    current_series = trend_series[-1]
                    definition = current_series.definition
                    recent_entries = list(current_series.entries[-28:])
                    long_entries = list(current_series.entries[-120:])
                    version_note = (
                        f" v{definition.definition_version}; other scoring series are not combined"
                        if len(trend_series) > 1
                        else ""
                    )
                    y_max = int(definition.score_max or 1)
                    self.recent_trend_charts[questionnaire_id].draw_series(
                        recent_entries,
                        [(definition.display_name, [row.total for row in recent_entries], color)],
                        y_max,
                        f"Recent {definition.display_name} recorded scores{version_note}",
                    )
                    self.long_term_trend_charts[questionnaire_id].draw_series(
                        long_entries,
                        [(definition.display_name, [row.total for row in long_entries], color)],
                        y_max,
                        f"Long-term {definition.display_name} recorded scores (up to 120 entries){version_note}",
                    )
                else:
                    reason = questionnaire_total_trend_omission_reason(definition)
                    self.recent_trend_charts[questionnaire_id].draw_series([], [], 1, reason)
                    self.long_term_trend_charts[questionnaire_id].draw_series([], [], 1, reason)

            _, cycle_entries, _ = _latest_definition_group("phq9", phq_entries, latest_date)
            cycles = treatment_cycles(cycle_entries, fetch_events(end=latest_date), latest_date)
            for index, (label, chart) in enumerate(zip(self.cycle_labels, self.cycle_charts)):
                if index < len(cycles):
                    cycle = cycles[index]
                    label.config(text=treatment_cycle_observation(cycle))
                    chart.draw_series(
                        cycle.entries,
                        [("PHQ-9", [row.total for row in cycle.entries], "#7C3AED")],
                        27,
                        f"{cycle.label}: recorded PHQ-9 scores",
                    )
                else:
                    label.config(text="A second recorded ketamine infusion is needed to show this cycle.")
                    chart.draw_series([], [], 27, "Treatment cycle not available")

        else:
            self.review_summary.config(text="No check-ins are available yet. Today's Check-In is ready when you are.")
            for label in self.review_highlights:
                label.config(text="")
            for questionnaire_id in QUESTIONNAIRE_ORDER:
                definition = QUESTIONNAIRES[questionnaire_id]
                y_max = int(definition.score_max or 1)
                self.recent_trend_charts[questionnaire_id].draw_series(
                    [], [], y_max, f"Recent {definition.display_name} recorded scores"
                )
                self.long_term_trend_charts[questionnaire_id].draw_series(
                    [], [], y_max, f"Long-term {definition.display_name} recorded scores"
                )
            for label, chart in zip(self.cycle_labels, self.cycle_charts):
                label.config(text="No recorded ketamine infusion is available for cycle review.")
                chart.draw_series([], [], 27, "Treatment cycle not available")

    def refresh_recent_table(self):
        for row in self.recent_table.get_children():
            self.recent_table.delete(row)
        phq_entries = fetch_assessment_entries("phq9")
        gad_entries = fetch_assessment_entries("gad7")
        phq_by_date = {row.entry_date: row for row in phq_entries}
        gad_by_date = {row.entry_date: row for row in gad_entries}
        dates = sorted(set(phq_by_date) | set(gad_by_date))[-14:]
        prior_phq = None
        prior_gad = None
        for entry_date in dates:
            phq = phq_by_date.get(entry_date)
            gad = gad_by_date.get(entry_date)
            trend_parts = []
            tag = "neutral"
            if phq and prior_phq and _entry_analytical_series_id(phq) == _entry_analytical_series_id(prior_phq):
                delta = phq.total - prior_phq.total
                if delta:
                    trend_parts.append(f"PHQ {delta:+d}")
                    tag = "worse" if delta > 0 else "better"
            if gad and prior_gad and _entry_analytical_series_id(gad) == _entry_analytical_series_id(prior_gad):
                delta = gad.total - prior_gad.total
                if delta:
                    trend_parts.append(f"GAD {delta:+d}")
                    if tag == "neutral":
                        tag = "worse" if delta > 0 else "better"
            self.recent_table.insert(
                "",
                END,
                values=[
                    entry_date,
                    phq.total if phq else "--",
                    phq.severity if phq else "--",
                    gad.total if gad else "--",
                    gad.severity if gad else "--",
                    ", ".join(trend_parts) or "No change",
                ],
                tags=(tag,),
            )
            prior_phq = phq or prior_phq
            prior_gad = gad or prior_gad

    def refresh_item_table(self):
        for row in self.item_table.get_children():
            self.item_table.delete(row)
        for questionnaire_id in QUESTIONNAIRE_ORDER:
            entries = fetch_assessment_entries(questionnaire_id)
            analytical_groups = _analytical_groups_for_entries(questionnaire_id, entries)
            for _, definition, definition_entries, versions in analytical_groups:
                if not definition_entries or questionnaire_profile_omission_reason(definition) is not None:
                    continue
                end_date = max(row.entry_date for row in definition_entries)
                start_date = (datetime.fromisoformat(end_date).date() - timedelta(days=13)).isoformat()
                score_14_day = calculate_questionnaire_profile_for_window(
                    definition, definition_entries, start_date, end_date
                )
                version_label = (
                    f" v{definition.definition_version}"
                    if len(analytical_groups) > 1
                    else ""
                )
                for idx, question in enumerate(definition.items, start=1):
                    values = [
                        _profile_response_score(question, row.items[idx - 1])
                        for row in definition_entries
                    ]
                    item_score = score_14_day.item_scores[idx - 1]
                    days_present = score_14_day.item_counts[idx - 1]
                    label = question.report_label or question.prompt
                    self.item_table.insert(
                        "",
                        END,
                        values=[
                            f"{definition.display_name}{version_label}",
                            f"Item {idx}: {label}",
                            f"{sum(values) / len(values):.1f}",
                            f"{item_score} ({days_present} days)",
                        ],
                    )

    def refresh_events(self):
        for row in self.events_table.get_children():
            self.events_table.delete(row)
        for event_id, event_date, event_type, desc in fetch_events():
            self.events_table.insert("", END, iid=str(event_id), values=[event_date, event_type, desc])

    def draw_charts(self):
        return

    def import_file(self):
        path = filedialog.askopenfilename(
            title="Import PHQ-9 spreadsheet",
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if not path:
            return
        try:
            if load_workbook is None:
                run_bundled_cli(["--import", path], ("openpyxl",))
                count = "the selected"
            else:
                count = import_spreadsheet(path)
            self.refresh_all()
            messagebox.showinfo("Import complete", f"Imported or updated {count} PHQ-9 entries. Review now shows the refreshed data.")
        except Exception as exc:
            messagebox.showerror("Import failed", str(exc))

    def export_analysis_file(self):
        PHQ9App._export_analysis_file_for(self, QUESTIONNAIRE_ORDER)

    def choose_analysis_export(self):
        self.choose_output_questionnaires("Choose workbook questionnaires", self._export_analysis_file_for)

    def _export_analysis_file_for(self, selected_ids):
        try:
            path = next_available_output_path(f"Len_Analysis_{date.today().isoformat()}.xlsx")
            if pd is None:
                run_bundled_cli(
                    ["--analysis-export", str(path), "--questionnaires", ",".join(selected_ids)],
                    ("pandas", "openpyxl"),
                )
            else:
                if selected_ids == QUESTIONNAIRE_ORDER:
                    export_analysis_workbook(str(path))
                else:
                    export_analysis_workbook(str(path), selected_ids)
            offer_to_open_generated_file(path, "Analysis workbook")
        except Exception as exc:
            messagebox.showerror("Analysis export failed", str(exc))

    def open_reports_folder(self):
        output_dir = None
        try:
            output_dir = generated_output_dir(create=True)
            open_in_default_application(output_dir)
        except OSError as exc:
            location = str(output_dir) if output_dir is not None else "The configured reports folder"
            messagebox.showwarning(
                "Reports folder unavailable",
                f"Windows could not open the reports folder automatically.\n\nFolder path:\n{location}\n\n{exc}",
            )

    def save_entry(self):
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        if datetime.fromisoformat(entry_date).date() > date.today():
            messagebox.showerror("Future date", "Future check-ins are not available.")
            return
        if entry_date != self._loaded_checkin_date:
            messagebox.showinfo(
                "Load the selected date first",
                "The date field differs from the check-in currently shown. Load the selected date before saving "
                "so responses are not applied to the wrong day.",
            )
            return
        existing = fetch_day_data(entry_date)
        selected = [
            assessment_id
            for assessment_id in QUESTIONNAIRE_ORDER
            if self.questionnaire_selected_vars[assessment_id].get()
        ]
        if not selected:
            messagebox.showerror("No questionnaire selected", "Select at least one questionnaire to record.")
            return
        saved = []
        responses_by_questionnaire = {}
        for assessment_id in selected:
            definition = ASSESSMENTS[assessment_id]
            items = [
                parse_questionnaire_response(question, var.get())
                for question, var in zip(definition.items, self.assessment_item_vars[assessment_id])
            ]
            unanswered = next((index for index, response in enumerate(items, start=1) if response is None), None)
            if unanswered is not None:
                self.active_questionnaire_id.set(assessment_id)
                self.update_questionnaire_visibility()
                messagebox.showerror(
                    "Unanswered question",
                    f"Choose a response for {definition.display_name} question {unanswered} before saving.",
                )
                return
            responses_by_questionnaire[assessment_id] = items
            saved.append(definition.display_name)
        upsert_questionnaire_entries(
            entry_date,
            responses_by_questionnaire,
            notes=str(existing["notes"]),
            note_tag=str(existing["note_tag"]),
            source="manual",
        )
        self.refresh_all()
        self.load_checkin_date(confirm_unsaved=False)
        self.show_checkin_details()
        action = "Updated existing" if existing["assessments"] else "Saved new"
        messagebox.showinfo(
            "Check-in recorded",
            f"{action} {', '.join(saved)} check-in for {entry_date}. The core check-in is saved. "
            "Optional details can be added now or skipped.",
        )

    def save_optional_details(self):
        entry_date = parse_date(self.date_var.get())
        if not entry_date:
            messagebox.showerror("Invalid date", "Enter the date as YYYY-MM-DD.")
            return
        custom_type = self.custom_event_type.get().strip()
        if self.include_custom_event.get() and not custom_type:
            messagebox.showerror("Missing custom event", "Enter a name for the custom event.")
            return
        notes = self.notes_box.get("1.0", END).strip()
        tag = self.note_tag.get().strip()
        update_daily_note(entry_date, notes, tag)
        description = self.checkin_event_desc.get("1.0", END).strip()
        for event_type, var in self.checkin_event_vars.items():
            if var.get():
                upsert_daily_event(entry_date, event_type, description)
        if self.include_custom_event.get():
            upsert_daily_event(entry_date, custom_type, description)
        self.checkin_event_desc.delete("1.0", END)
        self.refresh_all()
        self.load_checkin_date(confirm_unsaved=False)
        self.hide_checkin_details()

    def save_event(self):
        event_date = parse_date(self.event_date.get())
        if not event_date:
            messagebox.showerror("Invalid date", "Enter the event date as YYYY-MM-DD.")
            return
        add_event(event_date, self.event_type.get().strip() or "Treatment event", self.event_desc.get("1.0", END).strip(), dedupe=False)
        self.event_desc.delete("1.0", END)
        self.refresh_events()
        messagebox.showinfo("Saved", f"Saved event for {event_date}.")

    def choose_output_questionnaires(self, title, on_confirm, questionnaire_ids=None):
        questionnaire_ids = normalize_questionnaire_selection(questionnaire_ids)
        dialog = Toplevel(self)
        dialog.title(title)
        dialog.transient(self)
        dialog.grab_set()
        Label(dialog, text="Include questionnaires", font=("Segoe UI", 11, "bold")).pack(anchor="w", padx=16, pady=(14, 6))
        variables = {}
        for questionnaire_id in questionnaire_ids:
            variable = IntVar(value=1)
            variables[questionnaire_id] = variable
            Checkbutton(
                dialog,
                text=QUESTIONNAIRES[questionnaire_id].display_name,
                variable=variable,
            ).pack(anchor="w", padx=18, pady=3)

        def confirm():
            selected_ids = [questionnaire_id for questionnaire_id, variable in variables.items() if variable.get()]
            if not selected_ids:
                messagebox.showerror("No questionnaire selected", "Select at least one questionnaire.", parent=dialog)
                return
            dialog.destroy()
            on_confirm(selected_ids)

        controls = Frame(dialog)
        controls.pack(fill="x", padx=12, pady=12)
        Button(controls, text="Continue", command=confirm).pack(side=LEFT, padx=4)
        Button(controls, text="Cancel", command=dialog.destroy).pack(side=LEFT, padx=4)

    def create_report(self):
        PHQ9App._create_report_for(self, QUESTIONNAIRE_ORDER)

    def choose_report(self):
        try:
            default_start, default_end = available_report_date_range()
        except ValueError as exc:
            messagebox.showinfo("Report unavailable", str(exc))
            return

        dialog = Toplevel(self)
        dialog.title("Choose report date range")
        dialog.transient(self)
        dialog.grab_set()
        Label(dialog, text="Choose date range first", font=("Segoe UI", 11, "bold")).pack(
            anchor="w", padx=16, pady=(14, 6)
        )
        start_value = StringVar(value=default_start)
        end_value = StringVar(value=default_end)
        fields = Frame(dialog)
        fields.pack(fill="x", padx=16, pady=4)
        Label(fields, text="Start (YYYY-MM-DD)").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=3)
        Entry(fields, textvariable=start_value, width=14).grid(row=0, column=1, sticky="w", pady=3)
        Label(fields, text="End (YYYY-MM-DD)").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=3)
        Entry(fields, textvariable=end_value, width=14).grid(row=1, column=1, sticky="w", pady=3)

        def confirm_range():
            try:
                start, end = normalize_report_date_range(start_value.get().strip(), end_value.get().strip())
                eligible_ids = eligible_report_questionnaires(start, end)
            except ValueError as exc:
                messagebox.showerror("Invalid report range", str(exc), parent=dialog)
                return
            if not eligible_ids:
                messagebox.showinfo(
                    "Report unavailable",
                    "No check-ins are available in that date range.",
                    parent=dialog,
                )
                return
            dialog.destroy()
            self.choose_output_questionnaires(
                "Choose report questionnaires",
                lambda selected_ids: self._create_report_for(selected_ids, start, end),
                eligible_ids,
            )

        controls = Frame(dialog)
        controls.pack(fill="x", padx=12, pady=12)
        Button(controls, text="Continue", command=confirm_range).pack(side=LEFT, padx=4)
        Button(controls, text="Cancel", command=dialog.destroy).pack(side=LEFT, padx=4)

    def _create_report_for(self, selected_ids, start=None, end=None):
        try:
            if start is None or end is None:
                start, end = (
                    available_report_date_range()
                    if selected_ids == QUESTIONNAIRE_ORDER
                    else available_report_date_range(selected_ids)
                )
            else:
                start, end = normalize_report_date_range(start, end)
                ineligible = [
                    questionnaire_id
                    for questionnaire_id in selected_ids
                    if questionnaire_id not in eligible_report_questionnaires(start, end)
                ]
                if ineligible:
                    raise ValueError(
                        f"No check-ins are available for {QUESTIONNAIRES[ineligible[0]].display_name} in that date range."
                    )
        except ValueError as exc:
            messagebox.showinfo("Report unavailable", str(exc))
            return
        try:
            target = next_available_output_path(f"Len_Report_{start}_to_{end}.pdf")
            if colors is None or PILImage is None:
                run_bundled_cli(
                    ["--report-pdf", str(target), "--questionnaires", ",".join(selected_ids)],
                    ("reportlab", "PIL"),
                )
            else:
                if selected_ids == QUESTIONNAIRE_ORDER:
                    generate_report(start, end, str(target))
                else:
                    generate_report(start, end, str(target), selected_ids)
            offer_to_open_generated_file(target, "PDF report")
        except Exception as exc:
            messagebox.showerror("Report failed", str(exc))


def main():
    parser = argparse.ArgumentParser(description=f"Local {APPLICATION_NAME}")
    parser.add_argument("--import", dest="import_path", help="Import an Excel workbook, then exit unless --launch is also set.")
    parser.add_argument("--launch", action="store_true", help="Launch the GUI after command-line actions.")
    parser.add_argument("--report-pdf", help="Generate a full-history PDF report at this output path.")
    parser.add_argument("--report-start", help="First report date in YYYY-MM-DD format; use with --report-end.")
    parser.add_argument("--report-end", help="Last report date in YYYY-MM-DD format; use with --report-start.")
    parser.add_argument("--analysis-export", help="Export a normalized analysis-ready XLSX workbook, then exit.")
    parser.add_argument("--questionnaires", help="Comma-separated questionnaire IDs to include in generated outputs.")
    args = parser.parse_args()

    if bool(args.report_start) != bool(args.report_end):
        parser.error("--report-start and --report-end must be provided together.")
    if (args.report_start or args.report_end) and not args.report_pdf:
        parser.error("--report-start and --report-end require --report-pdf.")

    init_db()
    selected_ids = normalize_questionnaire_selection(args.questionnaires.split(",") if args.questionnaires else None)
    if args.import_path:
        count = import_spreadsheet(args.import_path)
        print(f"Imported or updated {count} entries.")
    if args.report_pdf:
        pdf_path = Path(args.report_pdf)
        if args.report_start and args.report_end:
            start, end = normalize_report_date_range(args.report_start, args.report_end)
            ineligible = [
                questionnaire_id
                for questionnaire_id in selected_ids
                if questionnaire_id not in eligible_report_questionnaires(start, end)
            ]
            if ineligible:
                raise ValueError(
                    f"No check-ins are available for {QUESTIONNAIRES[ineligible[0]].display_name} in that date range."
                )
        else:
            start, end = available_report_date_range(selected_ids)
        generate_report(start, end, str(pdf_path), selected_ids)
        selected_labels = ", ".join(
            f"{QUESTIONNAIRES[questionnaire_id].display_name} v{QUESTIONNAIRES[questionnaire_id].definition_version}"
            for questionnaire_id in selected_ids
        )
        print(f"Saved report to {pdf_path} ({start} to {end}; {selected_labels})")
    if args.analysis_export:
        export_analysis_workbook(args.analysis_export, selected_ids)
        print(f"Saved analysis workbook to {args.analysis_export}")
    has_cli_action = bool(args.import_path or args.report_pdf or args.analysis_export)
    if args.launch or not has_cli_action:
        PHQ9App().mainloop()


if __name__ == "__main__":
    main()
