import importlib
import sqlite3
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook


project_root = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(project_root / "src"))
app = importlib.import_module("phq9_tracker.app")


class BaselineSecurityRemediationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.previous_db_path = app.DB_PATH
        self.db_path = Path(self.tmp.name) / "baseline-security.sqlite"
        app.DB_PATH = self.db_path
        app.init_db()

    def tearDown(self):
        app.DB_PATH = self.previous_db_path
        self.tmp.cleanup()

    def make_import_workbook(self, name, rows):
        path = Path(self.tmp.name) / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(
            [
                "Date",
                *(f"Item {number}" for number in range(1, 10)),
                "Notes",
                "Note Tag",
                "Ketamine",
                "Therapy",
            ]
        )
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def database_snapshot(self):
        tables = (
            "phq9_entries",
            "assessment_entries",
            "questionnaire_submissions",
            "questionnaire_responses",
            "treatment_events",
        )
        with sqlite3.connect(self.db_path) as conn:
            return {
                table: conn.execute(f"SELECT * FROM {table} ORDER BY rowid").fetchall()
                for table in tables
            }

    def test_import_rolls_back_assessments_notes_and_events_after_late_invalid_row(self):
        app.upsert_entry("2026-01-01", [0] * 9, notes="Original fictional note", note_tag="Original")
        app.add_event("2026-01-01", "Existing fictional event", "Must remain unchanged", dedupe=False)
        before = self.database_snapshot()
        workbook_path = self.make_import_workbook(
            "late-invalid.xlsx",
            [
                ["2026-01-01", *([1] * 9), "Replacement note", "Imported", True, True],
                ["2026-01-02", 4, *([1] * 8), "Invalid later row", "Imported", True, True],
            ],
        )
        workbook = app.load_workbook(workbook_path, read_only=True, data_only=True)
        close_spy = Mock(wraps=workbook.close)
        workbook.close = close_spy

        with patch.object(app, "load_workbook", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "expected 0-3"):
                app.import_spreadsheet(str(workbook_path))

        close_spy.assert_called_once_with()
        self.assertEqual(self.database_snapshot(), before)

    def test_fully_valid_import_commits_assessments_responses_and_deduplicated_events(self):
        workbook_path = self.make_import_workbook(
            "valid.xlsx",
            [
                ["2026-02-01", *([1] * 9), "First fictional note", "Work", True, True],
                ["2026-02-02", *([2] * 9), "Second fictional note", "Sleep", False, True],
            ],
        )

        self.assertEqual(app.import_spreadsheet(str(workbook_path)), 2)
        self.assertEqual(app.import_spreadsheet(str(workbook_path)), 2)

        with sqlite3.connect(self.db_path) as conn:
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM phq9_entries").fetchone()[0], 2)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM assessment_entries").fetchone()[0], 2)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM questionnaire_submissions").fetchone()[0], 2)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM questionnaire_responses").fetchone()[0], 18)
            self.assertEqual(conn.execute("SELECT COUNT(*) FROM treatment_events").fetchone()[0], 3)
            self.assertEqual(
                conn.execute(
                    "SELECT notes, note_tag, source FROM assessment_entries WHERE entry_date = '2026-02-02'"
                ).fetchone(),
                ("Second fictional note", "Sleep", "spreadsheet"),
            )

    def test_analysis_workbook_preserves_formula_like_text_as_literal_strings(self):
        if app.pd is None or app.load_workbook is None:
            self.skipTest("Analysis workbook validation requires pandas and openpyxl.")
        app.upsert_entry("2026-03-01", [1] * 9, notes="=1+1", note_tag="+fictional tag")
        app.add_event("2026-03-01", "-fictional event", "@SUM(A1:A2)", dedupe=False)
        output = Path(self.tmp.name) / "formula-safe.xlsx"

        app.export_analysis_workbook(str(output))

        formula_view = app.load_workbook(output, read_only=False, data_only=False)
        value_view = app.load_workbook(output, read_only=False, data_only=True)
        notes = formula_view["Notes"]
        note_headers = {cell.value: cell.column for cell in notes[1]}
        events = formula_view["Treatment Events"]
        event_headers = {cell.value: cell.column for cell in events[1]}
        expected = [
            (notes.cell(2, note_headers["note_text"]), "=1+1"),
            (notes.cell(2, note_headers["note_tag"]), "+fictional tag"),
            (events.cell(2, event_headers["event_type"]), "-fictional event"),
            (events.cell(2, event_headers["normalized_event_type"]), "-fictional event"),
            (events.cell(2, event_headers["description"]), "@SUM(A1:A2)"),
        ]
        for cell, text in expected:
            self.assertEqual(cell.value, text)
            self.assertEqual(cell.data_type, "s")
        self.assertFalse(
            any(cell.data_type == "f" for sheet in formula_view.worksheets for row in sheet.iter_rows() for cell in row)
        )
        value_notes = value_view["Notes"]
        value_events = value_view["Treatment Events"]
        self.assertEqual(value_notes.cell(2, note_headers["note_text"]).value, "=1+1")
        self.assertEqual(value_notes.cell(2, note_headers["note_tag"]).value, "+fictional tag")
        self.assertEqual(value_events.cell(2, event_headers["event_type"]).value, "-fictional event")
        self.assertEqual(value_events.cell(2, event_headers["description"]).value, "@SUM(A1:A2)")
        formula_view.close()
        value_view.close()

    def test_report_chart_directory_is_removed_after_success(self):
        if app.colors is None or app.PILImage is None:
            self.skipTest("PDF report generation requires reportlab and Pillow.")
        app.upsert_entry("2026-04-01", [1] * 9)
        chart_root = Path(self.tmp.name) / "success-temp"
        chart_root.mkdir()
        output = Path(self.tmp.name) / "success.pdf"

        with patch.object(app.tempfile, "tempdir", str(chart_root)):
            app.generate_report("2026-04-01", "2026-04-01", str(output), ["phq9"])

        self.assertTrue(output.read_bytes().startswith(b"%PDF"))
        self.assertEqual(list(chart_root.glob("phq9_report_charts_*")), [])

    def test_report_chart_directory_is_removed_after_generation_failure(self):
        if app.colors is None or app.PILImage is None:
            self.skipTest("PDF report generation requires reportlab and Pillow.")
        app.upsert_entry("2026-04-02", [1] * 9)
        chart_root = Path(self.tmp.name) / "failure-temp"
        chart_root.mkdir()
        output = Path(self.tmp.name) / "failure.pdf"

        with patch.object(app.tempfile, "tempdir", str(chart_root)), patch.object(
            app.SimpleDocTemplate, "build", side_effect=RuntimeError("controlled fictional build failure")
        ):
            with self.assertRaisesRegex(RuntimeError, "controlled fictional build failure"):
                app.generate_report("2026-04-02", "2026-04-02", str(output), ["phq9"])

        self.assertEqual(list(chart_root.glob("phq9_report_charts_*")), [])

    def test_multiple_analytical_series_keep_distinct_chart_files_and_associations(self):
        if app.colors is None or app.PILImage is None:
            self.skipTest("PDF report generation requires reportlab and Pillow.")
        app.upsert_entry("2026-05-01", [1] * 9)
        current_entry = app.fetch_assessment_entries("phq9")[0]
        future_definition = replace(
            app.QUESTIONNAIRES["phq9"], definition_version=3, display_name="Fictional future PHQ-9"
        )
        future_entry = app.AssessmentEntryRow(
            2,
            "phq9",
            "2026-05-02",
            [2] * 9,
            18,
            "Moderately Severe",
            "",
            definition_version=3,
        )
        series = [
            app.QuestionnaireTrendSeries(app.QUESTIONNAIRES["phq9"], (current_entry,), (2,)),
            app.QuestionnaireTrendSeries(future_definition, (future_entry,), (3,)),
        ]
        associations = []
        chart_root = Path(self.tmp.name) / "series-temp"
        chart_root.mkdir()

        def write_fictional_chart(path, _entries, _series, title, _y_max, **_kwargs):
            Path(path).write_text(title, encoding="utf-8")

        def capture_chart(_story, image_path, caption=None, **_kwargs):
            associations.append((Path(image_path).name, caption, Path(image_path).read_text(encoding="utf-8")))

        with patch.object(app.tempfile, "tempdir", str(chart_root)), patch.object(
            app, "build_questionnaire_trend_series", return_value=series
        ), patch.object(app, "draw_line_chart", side_effect=write_fictional_chart), patch.object(
            app, "add_chart", side_effect=capture_chart
        ):
            app.generate_report(
                "2026-05-01", "2026-05-01", str(Path(self.tmp.name) / "two-series.pdf"), ["phq9"]
            )

        self.assertEqual([item[0] for item in associations], ["phq9_series_1_recent.png", "phq9_series_2_recent.png"])
        self.assertEqual(len({item[0] for item in associations}), 2)
        self.assertIn("PHQ-9 scores across the selected period", associations[0][1])
        self.assertIn("PHQ-9 Recorded Score Trend", associations[0][2])
        self.assertIn("Fictional future PHQ-9", associations[1][1])
        self.assertIn("Fictional future PHQ-9", associations[1][2])
        self.assertNotIn("v2", associations[0][1])
        self.assertEqual(list(chart_root.glob("phq9_report_charts_*")), [])


if __name__ == "__main__":
    unittest.main()
